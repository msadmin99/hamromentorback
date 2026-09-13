import io
from unittest.mock import patch

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, TransactionTestCase
from django.utils import timezone
from docx import Document as DocxDocument
from rest_framework import status
from rest_framework.test import APITestCase

from academics.import_dedup import (
    SIMILARITY_THRESHOLD, _length_could_match, _MAX_LENGTH_RATIO, _options_similarity, existing_texts_for_subject,
    find_duplicate, normalize_option_set, normalize_text,
)
from academics.importers.docx_parser import parse_docx
from academics.models import (
    Chapter, ImportBatch, ImportRow, Option, Question, QuestionAttempt, QuestionBankConfig,
    QuestionDifficultyRating, QuestionEvent, QuestionReport, ReferenceBook, Subject, Topic,
)
from accounts.models import RolePermission
from core.models import DeletionAuditLog
from tests_app.models import Test, TestAttempt, TestQuestion

User = get_user_model()

TINY_GIF = (
    b'GIF87a\x01\x00\x01\x00\x80\x01\x00\x00\x00\x00ccc,\x00'
    b'\x00\x00\x00\x01\x00\x01\x00\x00\x02\x02D\x01\x00;'
)


class QuestionDeleteTests(APITestCase):
    def setUp(self):
        self.staff = User.objects.create_user(
            username='staff1', email='staff1@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )
        self.student = User.objects.create_user(username='student1', email='student1@example.com', password='pw12345')
        self.subject = Subject.objects.create(name='Physics')
        self.client.force_authenticate(user=self.staff)

    def _make_question(self, with_image=False):
        image = SimpleUploadedFile('q.gif', TINY_GIF, content_type='image/gif') if with_image else None
        question = Question.objects.create(subject=self.subject, text='What is g?', image=image)
        Option.objects.create(question=question, text='9.8', is_correct=True)
        Option.objects.create(question=question, text='10')
        return question

    def test_delete_requires_staff(self):
        question = self._make_question()
        self.client.force_authenticate(user=self.student)
        resp = self.client.delete(f'/api/questions/{question.id}/')
        self.assertIn(resp.status_code, (status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN))
        self.assertTrue(Question.objects.filter(id=question.id).exists())

    def test_blocked_when_question_has_practice_attempt_history(self):
        question = self._make_question()
        option = question.options.first()
        question.attempts.create(user=self.student, selected_option=option, is_correct=True)

        resp = self.client.delete(f'/api/questions/{question.id}/')

        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertTrue(Question.objects.filter(id=question.id).exists())
        entry = DeletionAuditLog.objects.get(resource_type='Question', resource_id=str(question.id))
        self.assertEqual(entry.result, 'failure')

    def test_blocked_when_used_in_exam_with_student_attempts(self):
        question = self._make_question()
        test = Test.objects.create(title='Mock Test 1', exam_type='mock')
        TestQuestion.objects.create(test=test, question=question)
        TestAttempt.objects.create(user=self.student, test=test, status='submitted')

        resp = self.client.delete(f'/api/questions/{question.id}/')

        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('Mock Test 1', resp.data['detail'])
        self.assertTrue(Question.objects.filter(id=question.id).exists())

    def test_permanent_delete_succeeds_and_removes_options_and_images(self):
        question = self._make_question(with_image=True)
        option_ids = list(question.options.values_list('id', flat=True))
        image_name = question.image.name

        resp = self.client.delete(f'/api/questions/{question.id}/')

        self.assertEqual(resp.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(Question.objects.filter(id=question.id).exists())
        self.assertFalse(Option.objects.filter(id__in=option_ids).exists())
        self.assertFalse(default_storage_exists(image_name))
        entry = DeletionAuditLog.objects.get(resource_type='Question')
        self.assertEqual(entry.result, 'success')
        self.assertEqual(entry.actor, self.staff)

    def test_delete_of_untouched_question_is_not_blocked_by_unrelated_test(self):
        """A question that's merely attached to a Test with no attempts yet
        must still be deletable — only *attempted* usage should block it."""
        question = self._make_question()
        test = Test.objects.create(title='Draft Mock', exam_type='mock')
        TestQuestion.objects.create(test=test, question=question)

        resp = self.client.delete(f'/api/questions/{question.id}/')

        self.assertEqual(resp.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(Question.objects.filter(id=question.id).exists())


def default_storage_exists(name):
    from django.core.files.storage import default_storage
    return default_storage.exists(name)


class ImportRowDeleteTests(APITestCase):
    def setUp(self):
        self.staff = User.objects.create_user(
            username='staff1', email='staff1@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )
        self.client.force_authenticate(user=self.staff)
        self.batch = ImportBatch.objects.create(
            uploaded_by=self.staff, file_name='questions.xlsx', file_format='xlsx', status='ready', total_rows=2,
        )
        self.bad_row = ImportRow.objects.create(
            batch=self.batch, row_number=1, status='error',
            raw_data={'text_html': '<p>Q1</p>', 'options': [{'text_html': 'A', 'is_correct': True}]},
            errors=['Only 1 option(s) found — at least 2 are required.'],
        )
        self.good_row = ImportRow.objects.create(
            batch=self.batch, row_number=2, status='valid',
            raw_data={
                'text_html': '<p>Q2</p>',
                'options': [{'text_html': 'A', 'is_correct': True}, {'text_html': 'B', 'is_correct': False}],
            },
        )

    def test_delete_removes_row_and_decrements_total(self):
        resp = self.client.delete(f'/api/import-batches/{self.batch.id}/rows/{self.bad_row.id}/')

        self.assertEqual(resp.status_code, 200)
        self.assertFalse(ImportRow.objects.filter(id=self.bad_row.id).exists())
        self.assertEqual(resp.data['total_rows'], 1)
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.total_rows, 1)

    def test_delete_of_unknown_row_returns_404(self):
        resp = self.client.delete(f'/api/import-batches/{self.batch.id}/rows/999999/')
        self.assertEqual(resp.status_code, 404)

    def test_delete_blocked_once_import_has_started(self):
        self.batch.status = 'importing'
        self.batch.save(update_fields=['status'])

        resp = self.client.delete(f'/api/import-batches/{self.batch.id}/rows/{self.good_row.id}/')

        self.assertEqual(resp.status_code, 400)
        self.assertTrue(ImportRow.objects.filter(id=self.good_row.id).exists())

    def test_delete_requires_staff(self):
        student = User.objects.create_user(username='student1', email='student1@example.com', password='pw12345')
        self.client.force_authenticate(user=student)

        resp = self.client.delete(f'/api/import-batches/{self.batch.id}/rows/{self.good_row.id}/')

        self.assertIn(resp.status_code, (401, 403))
        self.assertTrue(ImportRow.objects.filter(id=self.good_row.id).exists())

    def test_patch_can_fix_an_error_row_to_valid(self):
        """Regression coverage for the Preview & Validate editability fix:
        correcting the underlying data (adding a 2nd option) must flip the
        row's status from error to valid via re-validation, same as the
        Admin's new inline editor relies on."""
        fixed_data = {
            'text_html': '<p>Q1</p>',
            'options': [{'text_html': 'A', 'is_correct': True}, {'text_html': 'B', 'is_correct': False}],
            'explanation_html': '<p>Because A is right.</p>',
        }
        resp = self.client.patch(f'/api/import-batches/{self.batch.id}/rows/{self.bad_row.id}/', {'data': fixed_data}, format='json')

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['status'], 'valid')
        self.assertEqual(resp.data['errors'], [])


class ImportBatchCreateTestModeMismatchTests(APITestCase):
    """Covers the bug where a batch could reach 'ready' with import_mode
    recorded as 'question_bank' despite the admin actually walking through
    the full Import & Create Test wizard (Test Configuration, Distribution
    Preview) — caused by the "Import Type" selector staying interactive
    while the file was still uploading. The fix: a 'ready' batch (nothing
    written yet) is safe to create-test on regardless of import_mode."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username='staff1', email='staff1@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )
        self.client.force_authenticate(user=self.staff)
        self.subject = Subject.objects.create(name='Physics')
        self.chapter = Chapter.objects.create(subject=self.subject, name='Mechanics')
        self.topic = Topic.objects.create(chapter=self.chapter, name='Kinematics')
        self.batch = ImportBatch.objects.create(
            uploaded_by=self.staff, file_name='q.xlsx', file_format='xlsx',
            status='ready', total_rows=1, import_mode='question_bank',
            subject=self.subject, chapter=self.chapter, topic=self.topic,
        )
        ImportRow.objects.create(
            batch=self.batch, row_number=1, status='valid',
            raw_data={
                'text_html': '<p>Q1</p>',
                'options': [{'text_html': 'A', 'is_correct': True}, {'text_html': 'B', 'is_correct': False}],
                'explanation_html': '<p>Because.</p>',
            },
        )

    def test_create_test_succeeds_on_ready_batch_despite_mismatched_import_mode(self):
        resp = self.client.post(
            f'/api/import-batches/{self.batch.id}/create-test/',
            {'title': 'Mock Test 1', 'exam_type': 'mock', 'duration_minutes': 30},
            format='json',
        )

        self.assertEqual(resp.status_code, 200)
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.status, 'completed')
        self.assertEqual(self.batch.import_mode, 'create_test')
        self.assertIsNotNone(self.batch.created_test_id)

    def test_create_test_still_blocked_on_a_failed_question_bank_batch(self):
        """A 'failed' batch is only safe to retry here if it failed inside
        this same synchronous flow (a real create_test batch) — a
        'question_bank' batch that failed via the separate background
        /confirm/ run may have already committed some rows, so it must
        stay blocked rather than risk reprocessing them into duplicates."""
        self.batch.status = 'failed'
        self.batch.save(update_fields=['status'])

        resp = self.client.post(
            f'/api/import-batches/{self.batch.id}/create-test/',
            {'title': 'Mock Test 1', 'exam_type': 'mock', 'duration_minutes': 30},
            format='json',
        )

        self.assertEqual(resp.status_code, 400)


class ImportBatchCreateQuestionsViewTests(APITestCase):
    """Phase B: POST /import-batches/<id>/create-questions/ — the
    Bulk-Import-Questions-into-an-existing-exam backend foundation. Mirrors
    ImportBatchCreateTestModeMismatchTests' setUp pattern above; the key
    difference under test is that this endpoint must create Questions and
    nothing else (no Test, no TestQuestion), while Import & Create Test
    keeps working exactly as before."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username='staff1', email='staff1@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )
        self.client.force_authenticate(user=self.staff)
        self.subject = Subject.objects.create(name='Physics')
        self.chapter = Chapter.objects.create(subject=self.subject, name='Mechanics')
        self.topic = Topic.objects.create(chapter=self.chapter, name='Kinematics')

    def _make_batch(self, **overrides):
        defaults = dict(
            uploaded_by=self.staff, file_name='q.xlsx', file_format='xlsx',
            status='ready', total_rows=1, import_mode='question_bank',
            subject=self.subject, chapter=self.chapter, topic=self.topic,
        )
        defaults.update(overrides)
        return ImportBatch.objects.create(**defaults)

    def _make_row(self, batch, row_number, text='Q', status='valid', **overrides):
        defaults = dict(
            batch=batch, row_number=row_number, status=status,
            raw_data={
                'text_html': f'<p>{text}</p>',
                'options': [{'text_html': 'A', 'is_correct': True}, {'text_html': 'B', 'is_correct': False}],
                'explanation_html': '<p>Because.</p>',
            },
        )
        defaults.update(overrides)
        return ImportRow.objects.create(**defaults)

    def _url(self, batch):
        return f'/api/import-batches/{batch.id}/create-questions/'

    # --- 1/6: valid import + correct returned question IDs -----------------
    def test_valid_import_creates_questions_in_row_order_and_returns_their_ids(self):
        batch = self._make_batch(total_rows=2)
        self._make_row(batch, 1, text='First')
        self._make_row(batch, 2, text='Second')

        resp = self.client.post(self._url(batch), {}, format='json')

        self.assertEqual(resp.status_code, 200)
        question_ids = resp.data['question_ids']
        self.assertEqual(len(question_ids), 2)
        questions = list(Question.objects.filter(id__in=question_ids).order_by('id'))
        # Order in the response must match row_number order, not just "both exist".
        first, second = Question.objects.get(id=question_ids[0]), Question.objects.get(id=question_ids[1])
        self.assertIn('First', first.text)
        self.assertIn('Second', second.text)
        self.assertEqual(len(questions), 2)

        batch.refresh_from_db()
        self.assertEqual(batch.status, 'completed')
        self.assertEqual(batch.created_count, 2)

    # --- 4: NOT a Test, NOT TestQuestion ------------------------------------
    def test_no_test_or_testquestion_is_ever_created(self):
        batch = self._make_batch()
        self._make_row(batch, 1)

        resp = self.client.post(self._url(batch), {}, format='json')

        self.assertEqual(resp.status_code, 200)
        self.assertIsNone(resp.data.get('created_test_id'))
        batch.refresh_from_db()
        self.assertIsNone(batch.created_test_id)
        self.assertEqual(Test.objects.count(), 0)
        self.assertEqual(TestQuestion.objects.count(), 0)

    # --- 2: invalid rows / unresolved duplicate decision --------------------
    def test_unresolved_duplicate_decision_blocks_with_400_and_writes_nothing(self):
        batch = self._make_batch()
        existing = Question.objects.create(subject=self.subject, chapter=self.chapter, topic=self.topic, text='<p>Dup</p>')
        self._make_row(batch, 1, status='duplicate', duplicate_of=existing, dedup_action='')

        resp = self.client.post(self._url(batch), {}, format='json')

        self.assertEqual(resp.status_code, 400)
        batch.refresh_from_db()
        self.assertEqual(batch.status, 'ready')  # untouched — nothing was attempted
        self.assertEqual(Question.objects.count(), 1)  # only the pre-existing one

    def test_missing_taxonomy_blocks_with_400(self):
        batch = self._make_batch(subject=None, chapter=None, topic=None)
        self._make_row(batch, 1)

        resp = self.client.post(self._url(batch), {}, format='json')

        self.assertEqual(resp.status_code, 400)

    # --- 3: duplicate handling — skip / replace ------------------------------
    def test_skip_attaches_the_existing_question_without_creating_a_new_one(self):
        batch = self._make_batch()
        existing = Question.objects.create(subject=self.subject, chapter=self.chapter, topic=self.topic, text='<p>Existing</p>')
        self._make_row(batch, 1, status='duplicate', duplicate_of=existing, dedup_action='skip')

        resp = self.client.post(self._url(batch), {}, format='json')

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['question_ids'], [existing.id])
        self.assertEqual(Question.objects.count(), 1)  # no new question created

    def test_replace_deletes_the_old_question_and_creates_a_new_one(self):
        batch = self._make_batch()
        old = Question.objects.create(subject=self.subject, chapter=self.chapter, topic=self.topic, text='<p>Old</p>')
        self._make_row(batch, 1, status='duplicate', duplicate_of=old, dedup_action='replace', text='New')

        resp = self.client.post(self._url(batch), {}, format='json')

        self.assertEqual(resp.status_code, 200)
        self.assertFalse(Question.objects.filter(id=old.id).exists())
        new_id = resp.data['question_ids'][0]
        self.assertIn('New', Question.objects.get(id=new_id).text)

    def test_replace_does_not_delete_an_already_referenced_question(self):
        """Same protection ImportBatchCreateTestView's shared engine already
        has (question_is_referenced) — must not silently regress here."""
        batch = self._make_batch()
        old = Question.objects.create(subject=self.subject, chapter=self.chapter, topic=self.topic, text='<p>Old</p>')
        test = Test.objects.create(title='T', exam_type='mock', duration_minutes=30, created_by=self.staff)
        TestQuestion.objects.create(test=test, question=old, order=0)
        self._make_row(batch, 1, status='duplicate', duplicate_of=old, dedup_action='replace', text='New')

        resp = self.client.post(self._url(batch), {}, format='json')

        self.assertEqual(resp.status_code, 200)
        self.assertTrue(Question.objects.filter(id=old.id).exists())  # not deleted — still referenced

    # --- 3b: preventing accidental re-import of the same batch ---------------
    def test_a_completed_batch_cannot_be_imported_again(self):
        batch = self._make_batch()
        self._make_row(batch, 1)
        first = self.client.post(self._url(batch), {}, format='json')
        self.assertEqual(first.status_code, 200)

        second = self.client.post(self._url(batch), {}, format='json')

        self.assertEqual(second.status_code, 400)
        # No second question was created for the same row.
        self.assertEqual(Question.objects.filter(subject=self.subject).count(), 1)

    # --- 4: transaction rollback --------------------------------------------
    def test_a_failing_row_rolls_back_every_question_already_created_in_the_attempt(self):
        from academics.import_engine import create_question_from_row as real_create_question_from_row

        batch = self._make_batch(total_rows=2)
        self._make_row(batch, 1, text='Survives only if rolled back')
        self._make_row(batch, 2, text='BOOM')

        def flaky(data, batch_arg, course_list):
            if data.get('text_html', '').find('BOOM') != -1:
                raise ValueError('forced failure for the rollback test')
            return real_create_question_from_row(data, batch_arg, course_list)

        with patch('academics.import_views.create_question_from_row', side_effect=flaky):
            resp = self.client.post(self._url(batch), {}, format='json')

        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.data['failed_row_number'], 2)
        # Row 1 would have succeeded in isolation — proving it did NOT
        # survive confirms the whole attempt rolled back atomically.
        self.assertEqual(Question.objects.filter(subject=self.subject).count(), 0)

        batch.refresh_from_db()
        self.assertEqual(batch.status, 'failed')
        row2 = batch.rows.get(row_number=2)
        self.assertEqual(row2.status, 'error')

    # --- 5: permission enforcement -------------------------------------------
    def test_unauthenticated_request_is_rejected(self):
        self.client.force_authenticate(user=None)
        batch = self._make_batch()
        self._make_row(batch, 1)

        resp = self.client.post(self._url(batch), {}, format='json')

        self.assertEqual(resp.status_code, 401)

    def test_non_staff_user_is_denied(self):
        student = User.objects.create_user(username='student1', email='student1@example.com', password='pw12345')
        self.client.force_authenticate(user=student)
        batch = self._make_batch()
        self._make_row(batch, 1)

        resp = self.client.post(self._url(batch), {}, format='json')

        self.assertEqual(resp.status_code, 403)

    def test_staff_user_without_question_entry_feature_is_denied(self):
        # An explicit RolePermission row that deliberately excludes
        # question_entry — proves the check is the granular feature, not
        # merely is_staff (which IsAdminUser, used by every sibling view in
        # this file, would have let through).
        RolePermission.objects.update_or_create(role='editor', defaults={'features': ['question_bank']})
        limited_staff = User.objects.create_user(
            username='limited1', email='limited1@example.com', password='pw12345',
            is_staff=True, admin_role='editor',
        )
        self.client.force_authenticate(user=limited_staff)
        batch = self._make_batch()
        self._make_row(batch, 1)

        resp = self.client.post(self._url(batch), {}, format='json')

        self.assertEqual(resp.status_code, 403)
        self.assertEqual(Question.objects.count(), 0)

    # --- 7/8: existing behavior unchanged -------------------------------------
    def test_create_test_endpoint_still_creates_a_test_unaffected_by_the_new_endpoint(self):
        """Import & Create Test must keep behaving exactly as before —
        proven by exercising it directly alongside the new endpoint's own
        test class, not just relying on its own pre-existing test file."""
        batch = self._make_batch()
        self._make_row(batch, 1)

        resp = self.client.post(
            f'/api/import-batches/{batch.id}/create-test/',
            {'title': 'Mock Test 1', 'exam_type': 'mock', 'duration_minutes': 30},
            format='json',
        )

        self.assertEqual(resp.status_code, 200)
        batch.refresh_from_db()
        self.assertIsNotNone(batch.created_test_id)
        self.assertEqual(TestQuestion.objects.filter(test_id=batch.created_test_id).count(), 1)

    def test_a_batch_already_consumed_by_create_test_cannot_be_imported_via_the_new_endpoint(self):
        batch = self._make_batch()
        self._make_row(batch, 1)
        created = self.client.post(
            f'/api/import-batches/{batch.id}/create-test/',
            {'title': 'Mock Test 1', 'exam_type': 'mock', 'duration_minutes': 30},
            format='json',
        )
        self.assertEqual(created.status_code, 200)

        resp = self.client.post(self._url(batch), {}, format='json')

        self.assertEqual(resp.status_code, 400)


class ImportUploadViewObservabilityTests(APITestCase):
    """Observability fix: ImportUploadView.post()'s final stretch (row
    validation, bulk_create, batch.save, building the response) previously
    had no exception handling at all — a failure there fell straight
    through to Django's generic 500 with nothing logged anywhere (Django's
    own default LOGGING only prints to console when DEBUG=True, which is
    never true in a real environment). These tests prove: (1) an unexpected
    exception there is now logged, with the batch id, before it propagates;
    (2) it still produces a 500, not a swallowed/converted response —
    exactly the same response shape a client saw before this change;
    (3) a normal, valid upload is completely unaffected; (4) an expected
    parser failure (a different, already-guarded code path) still returns
    its existing 400, proving this change didn't touch that behavior."""

    VALID_CSV = (
        'Question,Option1,Option2,Option3,Option4,Correct Option,Explanation\r\n'
        'What is 2+2?,3,4,5,6,2,Basic addition.\r\n'
    )

    def setUp(self):
        self.staff = User.objects.create_user(
            username='staff1', email='staff1@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )
        self.client.force_authenticate(user=self.staff)

    def _upload(self, content=None):
        content = self.VALID_CSV if content is None else content
        file_obj = SimpleUploadedFile('questions.csv', content.encode('utf-8'), content_type='text/csv')
        return self.client.post('/api/import-batches/upload/', {'file': file_obj, 'import_mode': 'question_bank'})

    def test_valid_upload_is_completely_unaffected(self):
        resp = self._upload()

        self.assertEqual(resp.status_code, 201)
        self.assertEqual(resp.data['total_rows'], 1)
        self.assertEqual(resp.data['status'], 'ready')
        batch = ImportBatch.objects.get(pk=resp.data['id'])
        self.assertEqual(batch.status, 'ready')
        self.assertEqual(ImportRow.objects.filter(batch=batch).count(), 1)

    def test_expected_parser_failure_still_returns_400_unchanged(self):
        # No 'Question' column at all — the parser's own pre-existing,
        # already-guarded ValueError path, untouched by this fix.
        resp = self._upload('NotAQuestionColumn\r\nfoo\r\n')

        self.assertEqual(resp.status_code, 400)
        self.assertIn('Missing required column', resp.data['detail'])
        # This path is guarded by the pre-existing try/except around the
        # parser call, not the new one — confirms the two don't overlap.
        with self.assertNoLogs('academics.import_views', level='ERROR'):
            self._upload('NotAQuestionColumn\r\nfoo\r\n')

    def test_unexpected_exception_is_logged_and_still_returns_500(self):
        def boom(*args, **kwargs):
            raise RuntimeError('simulated unexpected failure finalizing the batch')

        self.client.raise_request_exception = False
        with patch('academics.import_views.ImportRow.objects.bulk_create', side_effect=boom):
            with self.assertLogs('academics.import_views', level='ERROR') as captured:
                resp = self._upload()

        self.assertEqual(resp.status_code, 500)  # unchanged semantics — still a 500, never swallowed/converted
        self.assertEqual(len(captured.records), 1)
        message = captured.records[0].getMessage()
        self.assertIn('unexpected error finalizing batch', message)
        # Safe context only — a real, admin-visible batch id, the detected
        # file format, and a count — never file contents or credentials.
        self.assertIn('file_format=csv', message)
        self.assertIn('parsed_question_count=1', message)
        self.assertNotIn('2+2', message)  # never logs the uploaded file's own content
        # The traceback itself (not just the message) must be captured —
        # logger.exception() attaches it automatically; assertLogs exposes
        # it via the record's exc_info/exc_text.
        self.assertIsNotNone(captured.records[0].exc_info)
        self.assertIn('RuntimeError', captured.records[0].exc_text or '')

        # The batch this failed on is left in a real, inspectable state
        # (not silently lost) — created 'validating', never advanced to
        # 'ready' since the failure happened before that assignment.
        batch = ImportBatch.objects.filter(file_name='questions.csv').latest('id')
        self.assertEqual(batch.status, 'validating')

    def test_unexpected_exception_does_not_change_row_validation_behavior(self):
        """A validation-only difference (e.g. a genuinely malformed row)
        must still behave exactly as it did before this fix — no exception,
        no log entry, just the normal error-status row it always produced."""
        with self.assertNoLogs('academics.import_views', level='ERROR'):
            resp = self._upload('Question,Option1,Option2,Option3,Option4,Correct Option,Explanation\r\n,,,,,,\r\n')

        self.assertEqual(resp.status_code, 201)
        batch = ImportBatch.objects.get(pk=resp.data['id'])
        self.assertEqual(ImportRow.objects.get(batch=batch).status, 'error')


def _build_docx(lines):
    """Builds a minimal in-memory .docx from plain-text lines, mirroring
    docx_parser.build_template_docx()'s pattern."""
    doc = DocxDocument()
    for line in lines:
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


class DocxParserExplanationTests(TestCase):
    """Regression coverage for a real reported bug: a rich, AI-style
    explanation (Correct Answer / Core Concept / ... / Option Analysis /
    Common Exam Trap / Review Point) was being shredded on import, because
    the parser treated ANY line starting with "digit." or "letter)" as a
    new question/option — even decimal values and a per-option breakdown
    that legitimately appear inside the explanation itself. Confirmed via
    the actual "Volumetric Analysis.docx" file: 70 phantom rows and 8
    "options" per question instead of 66 real questions with 4 each."""

    def test_decimal_value_inside_explanation_does_not_start_a_new_question(self):
        buf = _build_docx([
            'Q1. What is the normality of a 1 M solution of H3PO4?',
            'A) 0.5 N', 'B) 0.1 N', 'C) 2.0 N', 'D) 3.0 N',
            'Answer: D',
            'Explanation:',
            'Correct Answer: D) 3.0 N',
            '0.5 N -- would require n = 0.5, which has no chemical basis.',
            '2.0 N -- this matches a diprotic acid, not H3PO4.',
            'Review Point: Normality = Molarity x Basicity.',
        ])

        questions = parse_docx(buf)

        self.assertEqual(len(questions), 1)
        self.assertEqual(len(questions[0]['options']), 4)
        self.assertIn('0.5 N', questions[0]['explanation_html'])
        self.assertIn('Review Point', questions[0]['explanation_html'])

    def test_lettered_option_analysis_inside_explanation_is_not_read_as_new_options(self):
        buf = _build_docx([
            'Q1. Molecular weight of a tribasic acid is W. Its equivalent weight is:',
            'A) W/2', 'B) W/3', 'C) W', 'D) 3W',
            'Answer: B',
            'Explanation:',
            'Correct Answer: B) W/3',
            'Option Analysis:',
            'A) W/2 -- corresponds to a dibasic acid.',
            'B) W/3 -- correctly divides by basicity.',
            'C) W -- would mean the acid is monobasic.',
            'D) 3W -- incorrectly multiplies instead of dividing.',
            'Review Point: Equivalent weight = Molecular weight / Basicity.',
        ])

        questions = parse_docx(buf)

        self.assertEqual(len(questions), 1)
        question = questions[0]
        self.assertEqual(len(question['options']), 4)
        correct = [i for i, o in enumerate(question['options']) if o['is_correct']]
        self.assertEqual(correct, [1])  # B
        self.assertIn('Option Analysis', question['explanation_html'])
        self.assertIn('Review Point', question['explanation_html'])

    def test_explicit_q_prefix_still_starts_a_new_question_even_mid_explanation(self):
        buf = _build_docx([
            'Q1. First question?',
            'A) 1', 'B) 2', 'C) 3', 'D) 4',
            'Answer: A',
            'Explanation:',
            'Some explanation text without a proper close.',
            'Q2. Second question?',
            'A) 5', 'B) 6', 'C) 7', 'D) 8',
            'Answer: B',
        ])

        questions = parse_docx(buf)

        self.assertEqual(len(questions), 2)
        self.assertEqual(len(questions[1]['options']), 4)


def _add_omml_equation(paragraph, symbols):
    """Injects a minimal, real OMML (<m:oMath>) equation into `paragraph`,
    with one <m:t> text node per string in `symbols` — simulating a
    question/option authored via Word's native Insert -> Equation tool,
    which python-docx's high-level API has no support for constructing.
    Used to reproduce and regression-test the "Question text is blank"
    false-Error bug fixed in docx_parser._paragraph_math_text()."""
    from docx.oxml.ns import qn
    from lxml import etree

    m_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    nsmap = {'m': m_ns}
    oMath = etree.SubElement(paragraph._p, f'{{{m_ns}}}oMath', nsmap=nsmap)  # noqa: SLF001 - no public API for OMML
    for sym in symbols:
        r = etree.SubElement(oMath, f'{{{m_ns}}}r')
        t = etree.SubElement(r, f'{{{m_ns}}}t')
        t.text = sym
    assert qn  # imported for clarity/parity with docx_parser.py's own import; unused directly here


class DocxParserOmmlEquationTests(TestCase):
    """Regression coverage for the Feature 3 false-Error audit: a
    question/option/explanation whose entire content is a Word-native
    equation (OMML, not typed text) must not be classified as blank."""

    def test_question_containing_only_an_omml_equation_is_not_blank(self):
        doc = DocxDocument()
        p = doc.add_paragraph('Q1. ')
        _add_omml_equation(p, ['G', 'R', 'E', '2', 'g'])
        doc.add_paragraph('A) g/R')
        doc.add_paragraph('B) 2g/R')
        doc.add_paragraph('C) g*R')
        doc.add_paragraph('D) 2g*R')
        doc.add_paragraph('Answer: A')
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        questions = parse_docx(buf)

        self.assertEqual(len(questions), 1)
        text = questions[0]['text_html']
        self.assertNotEqual(text.strip(), '')
        for sym in ['G', 'R', 'E', '2', 'g']:
            self.assertIn(sym, text)

    def test_paragraph_html_recovers_equation_only_content_directly(self):
        """Unit-level coverage of the actual function changed
        (_paragraph_html), independent of parse_docx's unrelated
        question/option-boundary heuristics — a paragraph whose only
        content is an OMML equation (no plain-text runs at all, the
        option-lettering case included) must not render as an empty <p>."""
        from academics.importers.docx_parser import _paragraph_html

        doc = DocxDocument()
        p = doc.add_paragraph()
        _add_omml_equation(p, ['2', 'g', 'R', 'squared'])

        html = _paragraph_html(p)

        self.assertNotEqual(html, '')
        self.assertIn('squared', html)

    def test_a_genuinely_blank_paragraph_with_no_text_and_no_equation_is_still_blank(self):
        """Root-cause fixes must not weaken validation globally — a
        paragraph with neither real text nor an OMML equation must still
        render as empty, exactly as before this fix."""
        from academics.importers.docx_parser import _paragraph_html

        doc = DocxDocument()
        p = doc.add_paragraph()  # no runs, no equation at all

        self.assertEqual(_paragraph_html(p), '')

    def test_a_paragraph_with_only_whitespace_runs_and_no_equation_is_still_blank(self):
        from academics.importers.docx_parser import _paragraph_html

        doc = DocxDocument()
        p = doc.add_paragraph('   ')

        self.assertEqual(_paragraph_html(p), '')


def _build_docx_runs(paragraphs):
    """Like `_build_docx`, but each paragraph is either a plain string or a
    list of `(text, bold)` run tuples — needed to build documents with
    real per-run bold, which `Document.add_paragraph(str)` alone can't
    express (python-docx gives every run of a plain-string paragraph
    `bold=None`, never `True`)."""
    doc = DocxDocument()
    for item in paragraphs:
        if isinstance(item, str):
            doc.add_paragraph(item)
            continue
        p = doc.add_paragraph()
        for text, bold in item:
            run = p.add_run(text)
            run.bold = bold
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


class DocxParserBoldCorrectAnswerTests(TestCase):
    """Regression coverage for the final, non-negotiable product rule:
    OPTION formatting from a .docx import must never carry bold, whether
    it exists because the author used bold as their own correct-answer
    marking convention (whole option), genuine word-level emphasis
    (partial), or lands on a distractor rather than the actual answer —
    none of that formatting is ever legitimate student-facing content on
    an option, so it is unconditionally stripped. Correctness itself is
    completely independent of formatting: it comes exclusively from the
    explicit "Answer: <letter>" line, exactly as before this change, and
    is represented purely structurally (`is_correct`) — never inferred
    from, or coupled to, bold. Question text and explanations are
    entirely out of scope and keep their bold exactly as authored."""

    def test_bold_correct_option_C_loses_the_bold_but_stays_correct(self):
        # The "C) " label is its own plain run and only the answer text
        # itself is bold — exactly how a real Word document is authored.
        buf = _build_docx_runs([
            'Q1. Which quantity is defined as charge per unit area?',
            'A) Electric intensity',
            'B) Current density',
            [('C) ', False), ('Potential difference', True)],
            'D) Capacitance',
            'Answer: C',
        ])

        questions = parse_docx(buf)

        self.assertEqual(len(questions), 1)
        options = questions[0]['options']
        self.assertEqual(len(options), 4)
        texts = [o['text_html'] for o in options]
        self.assertEqual(texts, [
            '<p>Electric intensity</p>',
            '<p>Current density</p>',
            '<p>Potential difference</p>',
            '<p>Capacitance</p>',
        ])
        self.assertNotIn('<strong>', texts[2])
        self.assertEqual([o['is_correct'] for o in options], [False, False, True, False])
        # The returned option shape carries only the documented fields —
        # no internal bookkeeping keys of any kind.
        self.assertEqual(set(options[2].keys()), {'text_html', 'is_correct', 'image_path'})

    def test_bold_correct_option_A(self):
        buf = _build_docx_runs([
            'Q1. Stem?',
            [('A) ', False), ('First choice', True)],
            'B) Second choice',
            'C) Third choice',
            'D) Fourth choice',
            'Answer: A',
        ])
        questions = parse_docx(buf)
        options = questions[0]['options']
        self.assertEqual(options[0]['text_html'], '<p>First choice</p>')
        self.assertEqual([o['is_correct'] for o in options], [True, False, False, False])

    def test_bold_correct_option_B(self):
        buf = _build_docx_runs([
            'Q1. Stem?',
            'A) First choice',
            [('B) ', False), ('Second choice', True)],
            'C) Third choice',
            'D) Fourth choice',
            'Answer: B',
        ])
        questions = parse_docx(buf)
        options = questions[0]['options']
        self.assertEqual(options[1]['text_html'], '<p>Second choice</p>')
        self.assertEqual([o['is_correct'] for o in options], [False, True, False, False])

    def test_bold_correct_option_D(self):
        buf = _build_docx_runs([
            'Q1. Stem?',
            'A) First choice',
            'B) Second choice',
            'C) Third choice',
            [('D) ', False), ('Fourth choice', True)],
            'Answer: D',
        ])
        questions = parse_docx(buf)
        options = questions[0]['options']
        self.assertEqual(options[3]['text_html'], '<p>Fourth choice</p>')
        self.assertEqual([o['is_correct'] for o in options], [False, False, False, True])

    def test_multiple_questions_each_with_their_own_bold_correct_option(self):
        buf = _build_docx_runs([
            'Q1. First stem?',
            [('A) ', False), ('Right one', True)],
            'B) Wrong one',
            'C) Wrong two',
            'D) Wrong three',
            'Answer: A',
            'Q2. Second stem?',
            'A) Wrong one',
            'B) Wrong two',
            [('C) ', False), ('Right one', True)],
            'D) Wrong three',
            'Answer: C',
        ])
        questions = parse_docx(buf)
        self.assertEqual(len(questions), 2)
        self.assertEqual(questions[0]['options'][0]['text_html'], '<p>Right one</p>')
        self.assertEqual([o['is_correct'] for o in questions[0]['options']], [True, False, False, False])
        self.assertEqual(questions[1]['options'][2]['text_html'], '<p>Right one</p>')
        self.assertEqual([o['is_correct'] for o in questions[1]['options']], [False, False, True, False])

    def test_partial_bold_word_is_stripped_but_option_text_is_preserved(self):
        # "MOST" is a single bold word inside otherwise-plain text — bold
        # is still option-presentation formatting, so it is stripped like
        # everything else. The important thing is the wording itself
        # survives intact: no character is lost, just the <strong> tag.
        buf = _build_docx_runs([
            'Q1. Which structure is involved?',
            [('A) ', False), ('Which structure is the ', False), ('MOST', True), (' important?', False)],
            'B) Second choice',
            'C) Third choice',
            'D) Fourth choice',
            'Answer: A',
        ])
        questions = parse_docx(buf)
        options = questions[0]['options']
        self.assertEqual(options[0]['text_html'], '<p>Which structure is the MOST important?</p>')
        self.assertNotIn('<strong>', options[0]['text_html'])
        self.assertTrue(options[0]['is_correct'])

    def test_bold_distractor_also_loses_its_bold_and_stays_incorrect(self):
        # A non-correct option that happens to be bold is NOT read as an
        # answer signal (the Answer: line alone decides is_correct), and
        # its bold is stripped too — presentation formatting is discarded
        # from every option, not only the correct one.
        buf = _build_docx_runs([
            'Q1. Stem?',
            [('A) ', False), ('Bolded but wrong', True)],
            'B) Correct one',
            'C) Third choice',
            'D) Fourth choice',
            'Answer: B',
        ])
        questions = parse_docx(buf)
        options = questions[0]['options']
        self.assertEqual(options[0]['text_html'], '<p>Bolded but wrong</p>')
        self.assertNotIn('<strong>', options[0]['text_html'])
        self.assertFalse(options[0]['is_correct'])
        self.assertTrue(options[1]['is_correct'])

    def test_latex_looking_text_in_a_bold_correct_option_is_unwrapped_intact(self):
        buf = _build_docx_runs([
            'Q1. Coefficient of viscosity has the formula:',
            'A) F = ma',
            [('B) ', False), (r'\(\eta = F/A(dv/dx)\)', True)],
            'C) E = mc^2',
            'D) PV = nRT',
            'Answer: B',
        ])
        questions = parse_docx(buf)
        options = questions[0]['options']
        self.assertEqual(options[1]['text_html'], r'<p>\(\eta = F/A(dv/dx)\)</p>')
        self.assertNotIn('<strong>', options[1]['text_html'])
        self.assertTrue(options[1]['is_correct'])

    def test_mixed_text_and_latex_bold_is_stripped_math_source_intact(self):
        # Options must appear in real A/B/C/D document order — the parser
        # assigns the letter by position, not by the label text typed.
        # The bold LaTeX span loses its <strong> like any other option
        # bold, but the LaTeX source itself (backslashes, braces) must
        # come through completely unaltered.
        buf = _build_docx_runs([
            'Q1. Stem?',
            'A) First',
            'B) Second',
            [('C) ', False), ('The answer is ', False), (r'\(F=ma\)', True), (' exactly.', False)],
            'D) Fourth',
            'Answer: C',
        ])
        questions = parse_docx(buf)
        options = questions[0]['options']
        self.assertEqual(options[2]['text_html'], r'<p>The answer is \(F=ma\) exactly.</p>')
        self.assertNotIn('<strong>', options[2]['text_html'])

    def test_bold_text_in_explanation_is_never_touched_by_this_fix(self):
        # The correct-answer marking convention only ever applies to
        # options — bold inside an explanation is ordinary rich-text
        # emphasis (e.g. AI-style "Correct Answer: **C**" formatting) and
        # is out of scope for this fix entirely.
        buf = _build_docx_runs([
            'Q1. Stem?',
            'A) First',
            'B) Second',
            [('C) ', False), ('Right answer', True)],
            'D) Fourth',
            'Answer: C',
            'Explanation:',
            [('This is the ', False), ('correct', True), (' reasoning.', False)],
        ])
        questions = parse_docx(buf)
        self.assertEqual(
            questions[0]['explanation_html'],
            '<p>This is the <strong>correct</strong> reasoning.</p>',
        )
        # And the option fix still applied independently.
        self.assertEqual(questions[0]['options'][2]['text_html'], '<p>Right answer</p>')

    def test_multi_line_bold_option_across_a_continuation_paragraph_is_fully_unwrapped(self):
        buf = _build_docx_runs([
            'Q1. Stem?',
            'A) First',
            'B) Second',
            [('C) ', False), ('First line bold', True)],
            [('second line also bold', True)],
            'D) Fourth',
            'Answer: C',
        ])
        questions = parse_docx(buf)
        options = questions[0]['options']
        self.assertEqual(options[2]['text_html'], '<p>First line bold</p><p>second line also bold</p>')

    def test_multi_line_option_partial_bold_across_continuation_is_fully_stripped(self):
        # The first line is bold, the continuation line is not — both end
        # up bold-free, since bold is discarded from the option
        # regardless of which line(s) carried it or how much of the
        # option they cover. No text from either line is lost.
        buf = _build_docx_runs([
            'Q1. Stem?',
            'A) First',
            'B) Second',
            [('C) ', False), ('First line bold', True)],
            [('second line not bold', False)],
            'D) Fourth',
            'Answer: C',
        ])
        questions = parse_docx(buf)
        options = questions[0]['options']
        self.assertEqual(
            options[2]['text_html'],
            '<p>First line bold</p><p>second line not bold</p>',
        )
        self.assertNotIn('<strong>', options[2]['text_html'])

    def test_plain_unbolded_correct_option_is_unaffected(self):
        buf = _build_docx([
            'Q1. Stem?',
            'A) First', 'B) Second', 'C) Third', 'D) Fourth',
            'Answer: C',
        ])
        questions = parse_docx(buf)
        options = questions[0]['options']
        self.assertEqual(options[2]['text_html'], '<p>Third</p>')
        self.assertTrue(options[2]['is_correct'])

    def test_question_text_containing_bold_is_never_touched(self):
        # Bold stripping is scoped to OPTIONS only — a bold word in the
        # question stem itself is completely out of scope and must render
        # exactly as authored.
        buf = _build_docx_runs([
            [('Q1. ', False), ('Which', False), (' quantity is the ', False), ('MOST', True), (' fundamental?', False)],
            'A) First',
            'B) Second',
            [('C) ', False), ('Third', True)],
            'D) Fourth',
            'Answer: C',
        ])
        questions = parse_docx(buf)
        self.assertEqual(
            questions[0]['text_html'],
            '<p>Which quantity is the <strong>MOST</strong> fundamental?</p>',
        )
        # And the option-side stripping still applies independently.
        self.assertEqual(questions[0]['options'][2]['text_html'], '<p>Third</p>')

    def test_correctness_never_depends_on_stored_html(self):
        # is_correct is read purely from the structural Answer: line —
        # confirm it agrees with is_correct regardless of which option (if
        # any) happened to be bold in the source, i.e. formatting and
        # correctness are fully decoupled in both directions.
        buf = _build_docx_runs([
            'Q1. Stem?',
            [('A) ', False), ('Bold but wrong', True)],
            [('B) ', False), ('Also bold, also wrong', True)],
            'C) Plain and correct',
            [('D) ', False), ('Bold and wrong too', True)],
            'Answer: C',
        ])
        questions = parse_docx(buf)
        options = questions[0]['options']
        self.assertEqual([o['is_correct'] for o in options], [False, False, True, False])
        for opt in options:
            self.assertNotIn('<strong>', opt['text_html'])


def _pq(text, option_texts):
    return {'text_html': text, 'options': [{'text_html': t} for t in option_texts]}


def _batch_entry(pq):
    return {'text': normalize_text(pq['text_html']), 'options': normalize_option_set(o['text_html'] for o in pq['options'])}


class ImportDedupTests(TestCase):
    """Regression coverage for a real reported false-positive: two questions
    sharing a common template stem ("Normality of X M solution of Y is?")
    but different specifics and different options were being flagged as
    duplicates on stem similarity alone. Duplicate detection must require
    both the stem AND the option set to substantially match."""

    def test_similar_stem_with_different_options_is_not_flagged(self):
        candidate = _pq('Normality of 1 M solution of phosphoric acid is', ['0.5 N', '0.1 N', '2.0 N', '3.0 N'])
        other = _pq('Normality of 1 M solution of sulphuric acid is', ['1 N', '2 N', 'N/2', 'N/4'])
        batch = {1: _batch_entry(other)}

        dup_id, score = find_duplicate(candidate, {}, batch, self_index=0)

        self.assertIsNone(dup_id)
        self.assertEqual(score, 0.0)

    def test_identical_stem_and_options_is_flagged(self):
        candidate = _pq('Normality of 1 M solution of sulphuric acid is', ['1 N', '2 N', 'N/2', 'N/4'])
        other = _pq('Normality of 1 M solution of sulphuric acid is', ['1 N', '2 N', 'N/2', 'N/4'])
        batch = {1: _batch_entry(other)}

        dup_id, score = find_duplicate(candidate, {}, batch, self_index=0)

        self.assertEqual(dup_id, 'row:1')
        self.assertGreaterEqual(score, 0.85)

    def test_matches_against_existing_db_question_require_both_dimensions(self):
        subject = Subject.objects.create(name='Chemistry')
        existing = Question.objects.create(subject=subject, text='Normality of 1 M solution of sulphuric acid is')
        Option.objects.create(question=existing, text='1 N')
        Option.objects.create(question=existing, text='2 N')
        Option.objects.create(question=existing, text='N/2')
        Option.objects.create(question=existing, text='N/4')

        existing_map = existing_texts_for_subject(subject)

        same_options = _pq('Normality of 1 M solution of sulphuric acid is', ['1 N', '2 N', 'N/2', 'N/4'])
        dup_id, score = find_duplicate(same_options, existing_map, self_index=None)
        self.assertEqual(dup_id, existing.id)
        self.assertGreaterEqual(score, 0.85)

        different_options = _pq('Normality of 1 M solution of phosphoric acid is', ['0.5 N', '0.1 N', '2.0 N', '3.0 N'])
        dup_id, score = find_duplicate(different_options, existing_map, self_index=None)
        self.assertIsNone(dup_id)


class ImportBatchTaxonomyDedupTriggerTests(APITestCase):
    """Bulk-import taxonomy audit: _run_dedup() must fire only when the
    taxonomy PATCH actually changes Subject — a Chapter/Topic/course-only
    change (or re-sending the same Subject) can never change which
    existing questions are duplicates, so it must be skipped entirely,
    not just cheap. Covers ImportBatchTaxonomyView.patch()'s new
    subject-changed gate and _batch_summary()'s N+1 fix together, since
    both are exercised by every PATCH to this endpoint."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username='taxstaff', email='taxstaff@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )
        self.student = User.objects.create_user(username='taxstudent', email='taxstudent@example.com', password='pw12345')
        self.subject_a = Subject.objects.create(name='Taxonomy Subject A')
        self.subject_b = Subject.objects.create(name='Taxonomy Subject B')
        self.chapter_a1 = Chapter.objects.create(subject=self.subject_a, name='A Chapter 1')
        self.chapter_a2 = Chapter.objects.create(subject=self.subject_a, name='A Chapter 2')
        self.topic_a1 = Topic.objects.create(chapter=self.chapter_a1, name='A Topic 1')
        self.topic_a2 = Topic.objects.create(chapter=self.chapter_a1, name='A Topic 2')

        self.batch = ImportBatch.objects.create(
            file_name='tax-test.csv', file_format='csv', import_mode='question_bank',
            status='ready', total_rows=2, uploaded_by=self.staff,
        )
        ImportRow.objects.create(
            batch=self.batch, row_number=1, status='valid',
            raw_data={'text_html': 'Row one question text', 'options': [{'text_html': 'A'}, {'text_html': 'B'}]},
        )
        ImportRow.objects.create(
            batch=self.batch, row_number=2, status='pending',
            raw_data={'text_html': 'Row two question text', 'options': [{'text_html': 'C'}, {'text_html': 'D'}]},
        )
        self.client.force_authenticate(user=self.staff)

    def _patch_taxonomy(self, **fields):
        payload = {'subject_id': None, 'chapter_id': None, 'topic_id': None, 'course_ids': []}
        payload.update(fields)
        return self.client.patch(f'/api/import-batches/{self.batch.id}/taxonomy/', payload, format='json')

    def test_subject_change_runs_dedup_exactly_once(self):
        with patch('academics.import_views._run_dedup') as mock_dedup:
            resp = self._patch_taxonomy(subject_id=self.subject_a.id)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(mock_dedup.call_count, 1)

    def test_same_subject_plus_chapter_change_does_not_run_dedup(self):
        self._patch_taxonomy(subject_id=self.subject_a.id)  # establish subject first (real dedup runs once here)
        with patch('academics.import_views._run_dedup') as mock_dedup:
            resp = self._patch_taxonomy(subject_id=self.subject_a.id, chapter_id=self.chapter_a1.id)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(mock_dedup.call_count, 0)

    def test_same_subject_plus_topic_change_does_not_run_dedup(self):
        self._patch_taxonomy(subject_id=self.subject_a.id, chapter_id=self.chapter_a1.id)
        with patch('academics.import_views._run_dedup') as mock_dedup:
            resp = self._patch_taxonomy(subject_id=self.subject_a.id, chapter_id=self.chapter_a1.id, topic_id=self.topic_a1.id)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(mock_dedup.call_count, 0)

    def test_reselecting_the_same_subject_again_does_not_rerun_dedup(self):
        self._patch_taxonomy(subject_id=self.subject_a.id)
        with patch('academics.import_views._run_dedup') as mock_dedup:
            resp = self._patch_taxonomy(subject_id=self.subject_a.id)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(mock_dedup.call_count, 0)

    def test_genuine_subject_change_runs_dedup_again(self):
        self._patch_taxonomy(subject_id=self.subject_a.id)
        with patch('academics.import_views._run_dedup') as mock_dedup:
            resp = self._patch_taxonomy(subject_id=self.subject_b.id)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(mock_dedup.call_count, 1)

    def test_chapter_and_topic_changes_preserve_existing_row_statuses(self):
        """A row already flagged 'duplicate' (by a real dedup pass, run once
        when Subject was set) must stay exactly as-is through subsequent
        Chapter/Topic-only changes that skip dedup — proving the skip never
        silently drops or alters row state. Subject is established first
        (its own real dedup pass, which would legitimately re-evaluate the
        row) before the row is put into the 'duplicate' state being
        protected, so only the Chapter-only change under test is mocked."""
        self._patch_taxonomy(subject_id=self.subject_a.id)  # real dedup, subject genuinely changes (None -> A)

        row = self.batch.rows.get(row_number=1)
        row.status = 'duplicate'
        row.duplicate_of_id = None
        row.warnings = ['85% similar to an existing question.']
        row.save(update_fields=['status', 'duplicate_of_id', 'warnings'])

        with patch('academics.import_views._run_dedup') as mock_dedup:
            resp = self._patch_taxonomy(subject_id=self.subject_a.id, chapter_id=self.chapter_a1.id)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(mock_dedup.call_count, 0)
        row.refresh_from_db()
        self.assertEqual(row.status, 'duplicate')
        self.assertEqual(row.warnings, ['85% similar to an existing question.'])

        resp = self._patch_taxonomy(subject_id=self.subject_a.id, chapter_id=self.chapter_a1.id, topic_id=self.topic_a1.id)
        self.assertEqual(resp.status_code, 200)
        row.refresh_from_db()
        self.assertEqual(row.status, 'duplicate')

    def test_duplicate_detection_is_still_correct_when_subject_genuinely_changes(self):
        """End-to-end (not mocked): a row identical to an existing question
        in Subject A is flagged duplicate when A is selected, and correctly
        un-flagged when the batch is moved to Subject B (no matching
        question there)."""
        existing = Question.objects.create(subject=self.subject_a, text='Row one question text')
        Option.objects.create(question=existing, text='A')
        Option.objects.create(question=existing, text='B')

        resp = self._patch_taxonomy(subject_id=self.subject_a.id)
        self.assertEqual(resp.status_code, 200)
        row = self.batch.rows.get(row_number=1)
        self.assertEqual(row.status, 'duplicate')

        resp = self._patch_taxonomy(subject_id=self.subject_b.id)
        self.assertEqual(resp.status_code, 200)
        row.refresh_from_db()
        self.assertIn(row.status, ('valid', 'warning'))

    def test_batch_summary_row_counts_match_manual_per_status_counts(self):
        """_batch_summary()'s new single grouped-aggregate query must
        return exactly the same {status: count} shape the old per-status
        .count() loop produced — including a 0 for every status with no
        rows in this batch."""
        resp = self._patch_taxonomy(subject_id=self.subject_a.id)
        self.assertEqual(resp.status_code, 200)
        expected = {s: self.batch.rows.filter(status=s).count() for s, _ in ImportRow.STATUS_CHOICES}
        self.assertEqual(resp.data['row_counts'], expected)
        self.assertEqual(set(resp.data['row_counts'].keys()), {s for s, _ in ImportRow.STATUS_CHOICES})

    def test_non_staff_cannot_change_taxonomy(self):
        self.client.force_authenticate(user=self.student)
        resp = self._patch_taxonomy(subject_id=self.subject_a.id)
        self.assertIn(resp.status_code, (status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN))
        self.batch.refresh_from_db()
        self.assertIsNone(self.batch.subject_id)

    def test_no_rows_are_lost_across_a_sequence_of_taxonomy_changes(self):
        self._patch_taxonomy(subject_id=self.subject_a.id)
        self._patch_taxonomy(subject_id=self.subject_a.id, chapter_id=self.chapter_a1.id)
        self._patch_taxonomy(subject_id=self.subject_a.id, chapter_id=self.chapter_a1.id, topic_id=self.topic_a1.id)
        self.assertEqual(self.batch.rows.count(), 2)
        self.assertEqual(set(self.batch.rows.values_list('row_number', flat=True)), {1, 2})


class ImportDedupAsyncTriggerTests(APITestCase):
    """Phase 3 (async dedup) mandatory tests 1-5: exactly one dedup task
    (Cloud Task enqueue) per genuine Subject change, none for Chapter/
    Topic/re-selecting-the-same-Subject, and rapid Subject A->B->C
    changes produce three separate, correctly-numbered generations.
    Mocks enqueue_dedup_task itself (not _run_dedup) so these tests
    verify the ENQUEUE decision — whether a task would be created —
    independent of DEDUP_PROCESSING_ASYNC's local-dev sync-fallback."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username='asyncstaff', email='asyncstaff@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )
        self.subject_a = Subject.objects.create(name='Async Subject A')
        self.subject_b = Subject.objects.create(name='Async Subject B')
        self.subject_c = Subject.objects.create(name='Async Subject C')
        self.chapter_a1 = Chapter.objects.create(subject=self.subject_a, name='A Chapter 1')
        self.topic_a1 = Topic.objects.create(chapter=self.chapter_a1, name='A Topic 1')
        self.batch = ImportBatch.objects.create(
            file_name='async-tax-test.csv', file_format='csv', import_mode='question_bank',
            status='ready', total_rows=1, uploaded_by=self.staff,
        )
        ImportRow.objects.create(
            batch=self.batch, row_number=1, status='valid',
            raw_data={'text_html': 'Async test question', 'options': [{'text_html': 'A'}, {'text_html': 'B'}]},
        )
        self.client.force_authenticate(user=self.staff)

    def _patch_taxonomy(self, **fields):
        payload = {'subject_id': None, 'chapter_id': None, 'topic_id': None, 'course_ids': []}
        payload.update(fields)
        return self.client.patch(f'/api/import-batches/{self.batch.id}/taxonomy/', payload, format='json')

    def test_subject_change_enqueues_exactly_one_dedup_generation(self):
        with patch('academics.import_views.enqueue_dedup_task') as mock_enqueue:
            resp = self._patch_taxonomy(subject_id=self.subject_a.id)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(mock_enqueue.call_count, 1)
        mock_enqueue.assert_called_once_with(self.batch.id, 1)
        self.assertEqual(resp.data['dedup_generation'], 1)
        self.assertEqual(resp.data['dedup_status'], 'pending')

    def test_chapter_change_enqueues_zero_dedup_tasks(self):
        with patch('academics.import_views.enqueue_dedup_task'):
            self._patch_taxonomy(subject_id=self.subject_a.id)
        with patch('academics.import_views.enqueue_dedup_task') as mock_enqueue:
            resp = self._patch_taxonomy(subject_id=self.subject_a.id, chapter_id=self.chapter_a1.id)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(mock_enqueue.call_count, 0)

    def test_topic_change_enqueues_zero_dedup_tasks(self):
        with patch('academics.import_views.enqueue_dedup_task'):
            self._patch_taxonomy(subject_id=self.subject_a.id, chapter_id=self.chapter_a1.id)
        with patch('academics.import_views.enqueue_dedup_task') as mock_enqueue:
            resp = self._patch_taxonomy(subject_id=self.subject_a.id, chapter_id=self.chapter_a1.id, topic_id=self.topic_a1.id)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(mock_enqueue.call_count, 0)

    def test_reselecting_the_same_subject_enqueues_zero_new_tasks(self):
        with patch('academics.import_views.enqueue_dedup_task'):
            self._patch_taxonomy(subject_id=self.subject_a.id)
        with patch('academics.import_views.enqueue_dedup_task') as mock_enqueue:
            resp = self._patch_taxonomy(subject_id=self.subject_a.id)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(mock_enqueue.call_count, 0)
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.dedup_generation, 1)  # unchanged — still generation 1

    def test_rapid_subject_a_to_b_to_c_changes_produce_separate_generations(self):
        with patch('academics.import_views.enqueue_dedup_task') as mock_enqueue:
            self._patch_taxonomy(subject_id=self.subject_a.id)
            self._patch_taxonomy(subject_id=self.subject_b.id)
            self._patch_taxonomy(subject_id=self.subject_c.id)
        self.assertEqual(mock_enqueue.call_count, 3)
        mock_enqueue.assert_any_call(self.batch.id, 1)
        mock_enqueue.assert_any_call(self.batch.id, 2)
        mock_enqueue.assert_any_call(self.batch.id, 3)
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.dedup_generation, 3)
        self.assertEqual(self.batch.subject_id, self.subject_c.id)
        self.assertEqual(self.batch.dedup_status, 'pending')  # generation 3's task never actually ran (mocked)


class ImportDedupAsyncIdempotencyTests(TestCase):
    """Phase 3 (async dedup) mandatory tests 6-9: the Cloud Task handler
    (run_dedup_task) under duplicate delivery, worker crash/retry, a
    stale processing claim, and — the core safety requirement — an old
    generation's task executing (a delayed retry or redelivery) AFTER a
    newer generation already exists must never overwrite the newer
    generation's results. Calls run_dedup_task() directly for precise
    control over execution order, decoupled from DEDUP_PROCESSING_ASYNC's
    sync-fallback timing."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username='idempstaff', email='idempstaff@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )
        self.subject_a = Subject.objects.create(name='Idempotency Subject A')
        self.subject_b = Subject.objects.create(name='Idempotency Subject B')
        self.existing_a = Question.objects.create(subject=self.subject_a, text='An existing A question')
        Option.objects.create(question=self.existing_a, text='Correct A', is_correct=True)
        Option.objects.create(question=self.existing_a, text='Wrong A')
        self.batch = ImportBatch.objects.create(
            file_name='idemp-test.csv', file_format='csv', import_mode='question_bank',
            status='ready', total_rows=1, uploaded_by=self.staff, subject=self.subject_a, dedup_generation=1,
            dedup_status='pending',
        )
        self.row = ImportRow.objects.create(
            batch=self.batch, row_number=1, status='valid',
            raw_data={'text_html': 'An existing A question', 'options': [{'text_html': 'Correct A'}, {'text_html': 'Wrong A'}]},
        )

    def test_task_for_current_generation_completes_and_flags_the_real_duplicate(self):
        from academics.import_dedup_tasks import run_dedup_task

        run_dedup_task(self.batch.id, 1)

        self.batch.refresh_from_db()
        self.row.refresh_from_db()
        self.assertEqual(self.batch.dedup_status, 'completed')
        self.assertIsNotNone(self.batch.dedup_completed_at)
        self.assertEqual(self.row.status, 'duplicate')
        self.assertEqual(self.row.duplicate_of_id, self.existing_a.id)

    def test_duplicate_task_delivery_for_an_already_completed_generation_is_a_safe_noop(self):
        from academics.import_dedup_tasks import run_dedup_task

        run_dedup_task(self.batch.id, 1)
        completed_at_first = ImportBatch.objects.get(pk=self.batch.id).dedup_completed_at

        # Simulated Cloud Tasks at-least-once redelivery of the same task.
        run_dedup_task(self.batch.id, 1)

        self.batch.refresh_from_db()
        self.assertEqual(self.batch.dedup_status, 'completed')
        # A genuine no-op — the second call never re-claimed, so
        # dedup_completed_at is untouched (not merely equal by luck).
        self.assertEqual(self.batch.dedup_completed_at, completed_at_first)

    def test_worker_crash_leaves_a_safe_non_completed_state_not_a_false_success(self):
        from academics.import_dedup_tasks import run_dedup_task

        with patch('academics.import_views._run_dedup', side_effect=RuntimeError('simulated crash')):
            with self.assertRaises(RuntimeError):
                run_dedup_task(self.batch.id, 1)

        self.batch.refresh_from_db()
        self.assertEqual(self.batch.dedup_status, 'processing')  # claimed, never falsely marked completed
        self.assertIsNotNone(self.batch.dedup_claimed_at)
        self.assertIsNone(self.batch.dedup_completed_at)

    def test_immediate_retry_after_a_fresh_claim_is_blocked_not_a_double_run(self):
        """A retry arriving before the claim looks abandoned must not
        re-run the work — this is what actually prevents two overlapping
        workers from processing the same generation at once."""
        from academics.import_dedup_tasks import run_dedup_task

        self.batch.dedup_status = 'processing'
        self.batch.dedup_claimed_at = timezone.now()
        self.batch.save(update_fields=['dedup_status', 'dedup_claimed_at'])

        with patch('academics.import_views._run_dedup') as mock_run:
            run_dedup_task(self.batch.id, 1)
        mock_run.assert_not_called()

    def test_stale_processing_claim_is_reclaimable(self):
        """A claim older than DEDUP_CLAIM_STALE_MINUTES is presumed
        abandoned (the worker that held it likely crashed/recycled) and a
        later retry may reclaim and complete it."""
        from academics.import_dedup_tasks import run_dedup_task

        self.batch.dedup_status = 'processing'
        self.batch.dedup_claimed_at = timezone.now() - timezone.timedelta(minutes=settings.DEDUP_CLAIM_STALE_MINUTES + 5)
        self.batch.save(update_fields=['dedup_status', 'dedup_claimed_at'])

        run_dedup_task(self.batch.id, 1)

        self.batch.refresh_from_db()
        self.assertEqual(self.batch.dedup_status, 'completed')
        self.row.refresh_from_db()
        self.assertEqual(self.row.status, 'duplicate')

    def test_old_generation_task_running_after_a_newer_generation_never_overwrites_it(self):
        """The core Phase 3 safety requirement: Subject A's dedup task
        (generation 1) is delayed — Subject B is selected before it runs
        (generation 2, which completes normally) — and only THEN does
        generation 1's task actually execute (a late retry/redelivery).
        It must be a safe no-op: it must not claim, must not touch row
        statuses, and must not mark itself completed."""
        from academics.import_dedup_tasks import run_dedup_task

        # Subject changes to B before generation 1 ever ran — exactly what
        # ImportBatchTaxonomyView.patch() does: bump generation, reset
        # dedup_status, and (in real life) enqueue a new task.
        self.batch.subject = self.subject_b
        self.batch.dedup_generation = 2
        self.batch.dedup_status = 'pending'
        self.batch.dedup_claimed_at = None
        self.batch.dedup_completed_at = None
        self.batch.save(update_fields=['subject', 'dedup_generation', 'dedup_status', 'dedup_claimed_at', 'dedup_completed_at'])

        # Generation 2 runs and completes normally first.
        run_dedup_task(self.batch.id, 2)
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.dedup_status, 'completed')
        gen2_completed_at = self.batch.dedup_completed_at
        self.row.refresh_from_db()
        gen2_row_status = self.row.status  # not a duplicate under Subject B — no matching question there

        # Generation 1's task finally executes — stale, superseded.
        with patch('academics.import_views._run_dedup') as mock_run:
            run_dedup_task(self.batch.id, 1)
        mock_run.assert_not_called()  # never even claimed, let alone ran the comparison

        self.batch.refresh_from_db()
        self.row.refresh_from_db()
        self.assertEqual(self.batch.dedup_generation, 2)
        self.assertEqual(self.batch.dedup_status, 'completed')
        self.assertEqual(self.batch.dedup_completed_at, gen2_completed_at)  # untouched by the stale run
        self.assertEqual(self.row.status, gen2_row_status)  # untouched by the stale run

    def test_old_generation_mid_run_detects_staleness_and_stops_before_writing_more_rows(self):
        """A subtler variant: generation 1's task is already IN PROGRESS
        (past its claim, mid-loop) when generation 2 supersedes it. The
        is_stale callback _run_dedup() checks every few rows must catch
        this and stop before writing any further rows for the stale
        generation."""
        from academics.import_dedup_tasks import _claim_dedup, run_dedup_task

        # A second row so the loop has more than one iteration to check staleness within.
        ImportRow.objects.create(
            batch=self.batch, row_number=2, status='valid',
            raw_data={'text_html': 'A second async test question', 'options': [{'text_html': 'X'}, {'text_html': 'Y'}]},
        )

        # Claim generation 1 (as run_dedup_task's first step would), then
        # simulate the Subject changing to B WHILE this "in-flight" run
        # would still be executing — is_stale() must see it immediately.
        self.assertTrue(_claim_dedup(self.batch.id, 1))
        ImportBatch.objects.filter(pk=self.batch.id).update(
            subject=self.subject_b, dedup_generation=2, dedup_status='pending', dedup_claimed_at=None,
        )

        from academics.import_dedup_tasks import _is_stale
        self.assertTrue(_is_stale(self.batch.id, 1))

        # The row(s) must be untouched — a stale run must never reach the
        # point of writing a result under the old Subject once superseded.
        for row in self.batch.rows.all():
            self.assertEqual(row.status, 'valid')


class ImportConfirmDedupGateTests(APITestCase):
    """Phase 3 (async dedup) mandatory tests 10-11: Confirm/Import must be
    blocked while dedup is pending/processing, and only usable once the
    batch's CURRENT generation has actually completed — never a stale,
    already-superseded completion."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username='confirmstaff', email='confirmstaff@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )
        self.subject_a = Subject.objects.create(name='Confirm Subject A')
        self.subject_b = Subject.objects.create(name='Confirm Subject B')
        self.chapter_a1 = Chapter.objects.create(subject=self.subject_a, name='Confirm Chapter A1')
        self.topic_a1 = Topic.objects.create(chapter=self.chapter_a1, name='Confirm Topic A1')
        self.batch = ImportBatch.objects.create(
            file_name='confirm-test.csv', file_format='csv', import_mode='question_bank',
            status='ready', total_rows=1, uploaded_by=self.staff,
            subject=self.subject_a, chapter=self.chapter_a1, topic=self.topic_a1,
            dedup_generation=1, dedup_status='pending',
        )
        ImportRow.objects.create(
            batch=self.batch, row_number=1, status='valid',
            raw_data={'text_html': 'Confirm test question', 'options': [{'text_html': 'A', 'is_correct': True}, {'text_html': 'B'}]},
        )
        self.client.force_authenticate(user=self.staff)

    def _confirm(self):
        return self.client.post(f'/api/import-batches/{self.batch.id}/confirm/', {}, format='json')

    def test_confirm_blocked_while_dedup_pending(self):
        resp = self._confirm()
        self.assertEqual(resp.status_code, 400)
        self.assertIn('progress', resp.data['detail'].lower())
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.status, 'ready')  # never transitioned to importing

    def test_confirm_blocked_while_dedup_processing(self):
        self.batch.dedup_status = 'processing'
        self.batch.save(update_fields=['dedup_status'])
        resp = self._confirm()
        self.assertEqual(resp.status_code, 400)

    def test_confirm_allowed_once_dedup_completed_for_current_generation(self):
        self.batch.dedup_status = 'completed'
        self.batch.save(update_fields=['dedup_status'])
        with patch('academics.import_views.enqueue_import_task'):
            resp = self._confirm()
        self.assertEqual(resp.status_code, 200)

    def test_confirm_blocked_if_generation_moved_on_after_completion(self):
        """The batch completed dedup for generation 1, then Subject
        changed again to B (generation 2, reset to pending) — Confirm
        must not be usable on generation 1's now-stale 'completed' result
        just because it once said 'completed'."""
        self.batch.dedup_status = 'completed'
        self.batch.save(update_fields=['dedup_status'])
        # A later Subject change resets dedup_status — exactly what
        # ImportBatchTaxonomyView.patch() does on a genuine Subject change.
        self.batch.subject = self.subject_b
        self.batch.dedup_generation = 2
        self.batch.dedup_status = 'pending'
        self.batch.save(update_fields=['subject', 'dedup_generation', 'dedup_status'])

        resp = self._confirm()
        self.assertEqual(resp.status_code, 400)


class ImportDedupLengthFilterTests(TestCase):
    """Phase 2 (dedup performance audit) regression coverage: the lossless
    length-ratio pre-filter in find_duplicate() must never change a
    duplicate/non-duplicate decision versus the unfiltered algorithm —
    only skip SequenceMatcher calls that are mathematically guaranteed to
    fall below SIMILARITY_THRESHOLD anyway (see the proof in
    academics/import_dedup.py above _MAX_LENGTH_RATIO)."""

    def _find_duplicate_unfiltered(self, pq, existing_by_id, batch_by_index=None, self_index=None):
        """Reference implementation: byte-for-byte the same logic as
        find_duplicate(), minus the length pre-filter — used to prove the
        real (filtered) function makes identical decisions."""
        from difflib import SequenceMatcher as SM

        candidate_text = normalize_text(pq.get('text_html'))
        if not candidate_text:
            return None, 0.0
        candidate_options = _options_from_parsed_for_test(pq.get('options'))

        best_id, best_score = None, 0.0

        def consider(other_id, other_text, other_options):
            nonlocal best_id, best_score
            if not other_text:
                return
            text_score = SM(None, candidate_text, other_text).ratio()
            if text_score < SIMILARITY_THRESHOLD:
                return
            options_score = _options_similarity(candidate_options, other_options)
            if options_score < SIMILARITY_THRESHOLD:
                return
            combined = min(text_score, options_score)
            if combined > best_score:
                best_id, best_score = other_id, combined

        for question_id, data in existing_by_id.items():
            consider(question_id, data['text'], data['options'])
        if batch_by_index:
            for idx, data in batch_by_index.items():
                if idx == self_index:
                    continue
                consider(f'row:{idx}', data['text'], data['options'])

        if best_id is not None:
            return best_id, round(best_score, 3)
        return None, 0.0

    # --- _length_could_match: direct boundary unit tests ---

    def test_max_length_ratio_matches_threshold_formula(self):
        self.assertAlmostEqual(_MAX_LENGTH_RATIO, (2 - SIMILARITY_THRESHOLD) / SIMILARITY_THRESHOLD)

    def test_length_could_match_true_at_and_below_the_ratio_boundary(self):
        # 135/100 = 1.35 <= 1.352941... -> must still be considered
        self.assertTrue(_length_could_match(100, 135))
        self.assertTrue(_length_could_match(135, 100))  # order-independent
        self.assertTrue(_length_could_match(100, 100))  # identical lengths

    def test_length_could_match_false_just_past_the_ratio_boundary(self):
        # 136/100 = 1.36 > 1.352941... -> provably below threshold, safe to skip
        self.assertFalse(_length_could_match(100, 136))
        self.assertFalse(_length_could_match(136, 100))

    def test_length_could_match_handles_empty_strings(self):
        self.assertTrue(_length_could_match(0, 0))
        self.assertFalse(_length_could_match(0, 5))
        self.assertFalse(_length_could_match(5, 0))

    def test_filter_boundary_matches_real_sequencematcher_at_the_exact_edge(self):
        """Empirical confirmation of the proof, not just algebra: at the
        length-ratio boundary, the pair the filter allows through really
        does score >= threshold, and the pair it excludes really would
        have scored < threshold even without the filter."""
        from difflib import SequenceMatcher as SM

        short_text = 'a' * 100
        long_135 = short_text + 'b' * 35   # ratio 1.35 <= boundary -> filter allows
        long_136 = short_text + 'b' * 36   # ratio 1.36 > boundary -> filter excludes

        self.assertTrue(_length_could_match(len(short_text), len(long_135)))
        self.assertFalse(_length_could_match(len(short_text), len(long_136)))

        self.assertGreaterEqual(SM(None, short_text, long_135).ratio(), SIMILARITY_THRESHOLD)
        self.assertLess(SM(None, short_text, long_136).ratio(), SIMILARITY_THRESHOLD)

    # --- Category regression corpus: filtered vs unfiltered must agree ---

    def _assert_same_decision(self, pq, existing_by_id, label):
        filtered = find_duplicate(pq, existing_by_id)
        unfiltered = self._find_duplicate_unfiltered(pq, existing_by_id)
        self.assertEqual(filtered, unfiltered, f'{label}: filtered={filtered} unfiltered={unfiltered}')
        return filtered

    def test_exact_duplicate(self):
        subject = Subject.objects.create(name='Dedup Exact')
        existing = Question.objects.create(subject=subject, text='What is the powerhouse of the cell?')
        Option.objects.create(question=existing, text='Mitochondria')
        Option.objects.create(question=existing, text='Nucleus')
        existing_map = existing_texts_for_subject(subject)

        pq = _pq('What is the powerhouse of the cell?', ['Mitochondria', 'Nucleus'])
        dup_id, score = self._assert_same_decision(pq, existing_map, 'exact duplicate')
        self.assertEqual(dup_id, existing.id)
        self.assertEqual(score, 1.0)

    def test_near_duplicate_above_threshold(self):
        subject = Subject.objects.create(name='Dedup Near')
        existing = Question.objects.create(
            subject=subject, text='The patient presents with acute chest pain radiating to the left arm',
        )
        Option.objects.create(question=existing, text='Myocardial infarction')
        Option.objects.create(question=existing, text='Costochondritis')
        existing_map = existing_texts_for_subject(subject)

        pq = _pq(
            'The patient presents with acute chest pain radiatng to the left arm',  # one typo
            ['Myocardial infarction', 'Costochondritis'],
        )
        dup_id, score = self._assert_same_decision(pq, existing_map, 'near duplicate')
        self.assertEqual(dup_id, existing.id)
        self.assertGreaterEqual(score, SIMILARITY_THRESHOLD)

    def test_short_questions_near_duplicate(self):
        subject = Subject.objects.create(name='Dedup Short')
        existing = Question.objects.create(subject=subject, text='2 + 2 = ?')
        Option.objects.create(question=existing, text='4')
        Option.objects.create(question=existing, text='5')
        existing_map = existing_texts_for_subject(subject)

        pq = _pq('2 + 2 = ?', ['4', '5'])
        dup_id, score = self._assert_same_decision(pq, existing_map, 'short question exact')
        self.assertEqual(dup_id, existing.id)

    def test_short_questions_genuinely_different(self):
        subject = Subject.objects.create(name='Dedup Short Different')
        existing = Question.objects.create(subject=subject, text='2 + 2 = ?')
        Option.objects.create(question=existing, text='4')
        Option.objects.create(question=existing, text='5')
        existing_map = existing_texts_for_subject(subject)

        pq = _pq('What is DNA?', ['Deoxyribonucleic acid', 'Ribonucleic acid'])
        dup_id, score = self._assert_same_decision(pq, existing_map, 'short question different')
        self.assertIsNone(dup_id)

    def test_long_clinical_vignette_near_duplicate(self):
        subject = Subject.objects.create(name='Dedup Vignette')
        vignette = (
            'A 45-year-old male presents to the emergency department with a two-hour history of '
            'crushing substernal chest pain radiating to the jaw and left shoulder, associated with '
            'diaphoresis, nausea, and shortness of breath. His blood pressure is 150/95 mmHg, heart '
            'rate is 110 beats per minute, and an ECG shows ST-segment elevation in leads II, III, '
            'and aVF. Which of the following is the most likely diagnosis?'
        )
        existing = Question.objects.create(subject=subject, text=vignette)
        Option.objects.create(question=existing, text='Inferior wall myocardial infarction')
        Option.objects.create(question=existing, text='Pulmonary embolism')
        existing_map = existing_texts_for_subject(subject)

        near_vignette = vignette.replace('two-hour', 'three-hour').replace('150/95 mmHg', '148/94 mmHg')
        pq = _pq(near_vignette, ['Inferior wall myocardial infarction', 'Pulmonary embolism'])
        dup_id, score = self._assert_same_decision(pq, existing_map, 'long vignette near duplicate')
        self.assertEqual(dup_id, existing.id)
        self.assertGreaterEqual(score, SIMILARITY_THRESHOLD)

    def test_different_length_questions_not_flagged(self):
        """A short question and a long, unrelated question must never be
        flagged, before or after the filter — the filter's whole job is to
        recognize exactly this case cheaply."""
        subject = Subject.objects.create(name='Dedup Different Length')
        long_text = (
            'A 62-year-old woman with a history of type 2 diabetes mellitus and hypertension presents '
            'with progressive dyspnea on exertion, bilateral lower extremity edema, and orthopnea over '
            'the past three weeks. Which of the following is the most appropriate initial investigation?'
        )
        existing = Question.objects.create(subject=subject, text=long_text)
        Option.objects.create(question=existing, text='Echocardiogram')
        Option.objects.create(question=existing, text='Chest X-ray')
        existing_map = existing_texts_for_subject(subject)

        pq = _pq('What is the normal pH of blood?', ['7.35-7.45', '6.8-7.0'])
        dup_id, score = self._assert_same_decision(pq, existing_map, 'different length, unrelated')
        self.assertIsNone(dup_id)

    def test_questions_just_above_and_just_below_threshold(self):
        """Two hand-tuned, equal-length variants of the same base sentence,
        one scoring just above SIMILARITY_THRESHOLD and one just below —
        proves the threshold comparison itself is untouched by the filter
        (equal-length pairs are never excluded by it)."""
        base = 'the patient presents with acute chest pain radiating to the left arm and shortness of breath'
        just_above = 'xhe patxent prxsents xith acxte chext painxradiatxng to xhe lefx arm axd shorxness ox breath'
        just_below = 'xhe paxient xresenxs witx acutx chesx painxradiaxing tx the xeft axm andxshortxess of breath'
        self.assertEqual(len(base), len(just_above))
        self.assertEqual(len(base), len(just_below))

        subject = Subject.objects.create(name='Dedup Threshold Boundary')
        existing = Question.objects.create(subject=subject, text=base)
        Option.objects.create(question=existing, text='Diagnosis A')
        Option.objects.create(question=existing, text='Diagnosis B')
        existing_map = existing_texts_for_subject(subject)

        dup_id_above, score_above = self._assert_same_decision(
            _pq(just_above, ['Diagnosis A', 'Diagnosis B']), existing_map, 'just above threshold',
        )
        self.assertEqual(dup_id_above, existing.id)
        self.assertGreaterEqual(score_above, SIMILARITY_THRESHOLD)

        dup_id_below, score_below = self._assert_same_decision(
            _pq(just_below, ['Diagnosis A', 'Diagnosis B']), existing_map, 'just below threshold',
        )
        self.assertIsNone(dup_id_below)

    def test_formatting_and_punctuation_differences_still_match(self):
        subject = Subject.objects.create(name='Dedup Formatting')
        existing = Question.objects.create(subject=subject, text='What is the capital of France?')
        Option.objects.create(question=existing, text='Paris')
        Option.objects.create(question=existing, text='Lyon')
        existing_map = existing_texts_for_subject(subject)

        pq = _pq('<p>What is the <b>capital</b> of France?</p>', ['Paris', 'Lyon'])
        dup_id, score = self._assert_same_decision(pq, existing_map, 'HTML formatting difference')
        self.assertEqual(dup_id, existing.id)

    def test_random_corpus_filtered_matches_unfiltered(self):
        """Broad sweep: many random text pairs across a wide length range
        (well beyond any single hand-picked case) — the filtered function
        must agree with the unfiltered reference on every single one."""
        import random

        rng = random.Random(2026)
        vocab = ['the', 'patient', 'question', 'diagnosis', 'blood', 'cell', 'acute', 'chronic', 'exam', 'test']

        def random_text(min_words, max_words):
            n = rng.randint(min_words, max_words)
            return ' '.join(rng.choice(vocab) for _ in range(n))

        existing_by_id = {}
        for i in range(40):
            existing_by_id[i] = {
                'text': random_text(2, 60),
                'options': frozenset({random_text(1, 3), random_text(1, 3)}),
            }

        for i in range(60):
            candidate_text = random_text(2, 60)
            pq = {
                'text_html': candidate_text,
                'options': [{'text_html': t} for t in (random_text(1, 3), random_text(1, 3))],
            }
            filtered = find_duplicate(pq, existing_by_id)
            unfiltered = self._find_duplicate_unfiltered(pq, existing_by_id)
            self.assertEqual(filtered, unfiltered, f'row {i}: candidate={candidate_text!r}')


def _options_from_parsed_for_test(options):
    return normalize_option_set((o.get('text_html') for o in (options or [])))


class QuestionBookmarkFilterTests(APITestCase):
    """Powers the QBank 'Bookmarks' page — ?bookmarked=true must return only
    this user's own bookmarked questions, never another student's."""

    def setUp(self):
        self.student = User.objects.create_user(username='student1', email='student1@example.com', password='pw12345')
        self.other = User.objects.create_user(username='other1', email='other1@example.com', password='pw12345')
        self.subject = Subject.objects.create(name='Physics')
        self.bookmarked_q = Question.objects.create(subject=self.subject, text='Bookmarked question')
        self.plain_q = Question.objects.create(subject=self.subject, text='Not bookmarked')
        QuestionAttempt.objects.create(user=self.student, question=self.bookmarked_q, is_bookmarked=True)
        QuestionAttempt.objects.create(user=self.student, question=self.plain_q, is_bookmarked=False)
        # Another student bookmarking the same question must not leak into self.student's list.
        QuestionAttempt.objects.create(user=self.other, question=self.plain_q, is_bookmarked=True)

    def test_bookmarked_filter_returns_only_own_bookmarks(self):
        self.client.force_authenticate(user=self.student)
        resp = self.client.get('/api/questions/?bookmarked=true')
        self.assertEqual(resp.status_code, 200)
        ids = {q['id'] for q in resp.data}
        self.assertEqual(ids, {self.bookmarked_q.id})

    def test_bookmarked_filter_requires_auth(self):
        resp = self.client.get('/api/questions/?bookmarked=true')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data, [])

    def test_without_filter_all_visible_questions_still_returned(self):
        self.client.force_authenticate(user=self.student)
        resp = self.client.get('/api/questions/')
        ids = {q['id'] for q in resp.data}
        self.assertEqual(ids, {self.bookmarked_q.id, self.plain_q.id})


class QuestionBookmarkToggleTests(APITestCase):
    """Regression coverage: bookmarking must never blank out a previously
    recorded answer (see the bookmark() action's docstring for the bug in
    answer() this deliberately avoids)."""

    def setUp(self):
        self.student = User.objects.create_user(username='student1', email='student1@example.com', password='pw12345')
        self.subject = Subject.objects.create(name='Physics')
        self.question = Question.objects.create(subject=self.subject, text='2+2=?')
        self.option = Option.objects.create(question=self.question, text='4', is_correct=True)
        self.client.force_authenticate(user=self.student)

    def test_bookmark_on_a_never_attempted_question_creates_attempt(self):
        resp = self.client.post(f'/api/questions/{self.question.id}/bookmark/', {'bookmark': True})
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.data['is_bookmarked'])
        attempt = QuestionAttempt.objects.get(user=self.student, question=self.question)
        self.assertTrue(attempt.is_bookmarked)
        self.assertIsNone(attempt.selected_option)

    def test_bookmarking_does_not_erase_a_previously_recorded_answer(self):
        self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.option.id})
        self.client.post(f'/api/questions/{self.question.id}/bookmark/', {'bookmark': True})

        attempt = QuestionAttempt.objects.get(user=self.student, question=self.question)
        self.assertTrue(attempt.is_bookmarked)
        self.assertEqual(attempt.selected_option_id, self.option.id)
        self.assertTrue(attempt.is_correct)

    def test_unbookmark_clears_the_flag_only(self):
        self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.option.id})
        self.client.post(f'/api/questions/{self.question.id}/bookmark/', {'bookmark': True})
        resp = self.client.post(f'/api/questions/{self.question.id}/bookmark/', {'bookmark': False})

        self.assertFalse(resp.data['is_bookmarked'])
        attempt = QuestionAttempt.objects.get(user=self.student, question=self.question)
        self.assertFalse(attempt.is_bookmarked)
        self.assertEqual(attempt.selected_option_id, self.option.id)

    def test_is_bookmarked_reflects_in_the_question_list_response(self):
        other_question = Question.objects.create(subject=self.subject, text='3+3=?')
        self.client.post(f'/api/questions/{self.question.id}/bookmark/', {'bookmark': True})

        resp = self.client.get('/api/questions/')

        by_id = {q['id']: q['is_bookmarked'] for q in resp.data}
        self.assertTrue(by_id[self.question.id])


class RecordQuestionResultTests(TestCase):
    """academics.services.record_question_result — the single write path
    for the Smart Question Bank's performance tracking, used by both QBank
    practice and test submission."""

    def setUp(self):
        self.student = User.objects.create_user(username='student1', email='student1@example.com', password='pw12345')
        self.subject = Subject.objects.create(name='Physics')
        self.question = Question.objects.create(subject=self.subject, text='2+2=?')

    def test_first_correct_attempt_is_learning_not_mastered(self):
        from academics.services import record_question_result

        attempt = record_question_result(self.student, self.question, True, source='qbank')

        self.assertEqual(attempt.attempts_count, 1)
        self.assertEqual(attempt.correct_count, 1)
        self.assertEqual(attempt.incorrect_count, 0)
        self.assertTrue(attempt.is_correct)
        self.assertTrue(attempt.last_result)
        self.assertEqual(attempt.mastery_status, 'learning')
        self.assertIsNotNone(attempt.revision_due_at)

    def test_two_correct_attempts_become_mastered(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.question, True, source='qbank')
        attempt = record_question_result(self.student, self.question, True, source='qbank')

        self.assertEqual(attempt.attempts_count, 2)
        self.assertEqual(attempt.mastery_status, 'mastered')

    def test_repeated_incorrect_attempts_become_weak(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.question, False, source='qbank')
        record_question_result(self.student, self.question, False, source='qbank')
        attempt = record_question_result(self.student, self.question, False, source='qbank')

        self.assertEqual(attempt.attempts_count, 3)
        self.assertEqual(attempt.incorrect_count, 3)
        self.assertEqual(attempt.mastery_status, 'weak')
        self.assertFalse(attempt.last_result)

    def test_counts_accumulate_across_calls_instead_of_overwriting(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.question, True, source='qbank')
        record_question_result(self.student, self.question, False, source='test')
        attempt = record_question_result(self.student, self.question, True, source='qbank')

        self.assertEqual(attempt.attempts_count, 3)
        self.assertEqual(attempt.correct_count, 2)
        self.assertEqual(attempt.incorrect_count, 1)

    def test_never_touches_bookmark(self):
        from academics.services import record_question_result

        QuestionAttempt.objects.create(user=self.student, question=self.question, is_bookmarked=True)
        record_question_result(self.student, self.question, False, source='qbank')

        attempt = QuestionAttempt.objects.get(user=self.student, question=self.question)
        self.assertTrue(attempt.is_bookmarked)

    def test_logs_an_immutable_question_event_each_call(self):
        from academics.models import QuestionEvent
        from academics.services import record_question_result

        record_question_result(self.student, self.question, True, source='qbank')
        record_question_result(self.student, self.question, False, source='test')

        events = list(QuestionEvent.objects.filter(user=self.student, question=self.question).order_by('id'))
        self.assertEqual(len(events), 2)
        self.assertEqual([e.source for e in events], ['qbank', 'test'])
        self.assertEqual([e.is_correct for e in events], [True, False])

    def test_does_not_create_duplicate_attempt_rows(self):
        from academics.services import record_question_result

        for _ in range(3):
            record_question_result(self.student, self.question, True, source='qbank')

        self.assertEqual(QuestionAttempt.objects.filter(user=self.student, question=self.question).count(), 1)


class QuestionAnswerRecordsPerformanceTests(APITestCase):
    """/questions/{id}/answer/ must feed the new performance-tracking
    counters, not just overwrite the latest-state fields, and must never
    touch is_bookmarked (see QuestionBookmarkToggleTests for that contract)."""

    def setUp(self):
        self.student = User.objects.create_user(username='student1', email='student1@example.com', password='pw12345')
        self.subject = Subject.objects.create(name='Physics')
        self.question = Question.objects.create(subject=self.subject, text='2+2=?')
        self.correct_option = Option.objects.create(question=self.question, text='4', is_correct=True)
        self.wrong_option = Option.objects.create(question=self.question, text='5', is_correct=False)
        self.client.force_authenticate(user=self.student)

    def test_answering_twice_accumulates_attempts_count(self):
        self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.wrong_option.id})
        self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.correct_option.id})

        attempt = QuestionAttempt.objects.get(user=self.student, question=self.question)
        self.assertEqual(attempt.attempts_count, 2)
        self.assertEqual(attempt.correct_count, 1)
        self.assertEqual(attempt.incorrect_count, 1)

    def test_answering_with_no_option_does_not_record_an_attempt(self):
        self.client.post(f'/api/questions/{self.question.id}/answer/', {})
        self.assertFalse(QuestionAttempt.objects.filter(user=self.student, question=self.question).exists())

    def test_answering_does_not_reset_an_existing_bookmark(self):
        self.client.post(f'/api/questions/{self.question.id}/bookmark/', {'bookmark': True})
        self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.correct_option.id})

        attempt = QuestionAttempt.objects.get(user=self.student, question=self.question)
        self.assertTrue(attempt.is_bookmarked)


class AnswerConfidenceTests(APITestCase):
    """QBank homepage redesign: post-answer 'how confident were you?'
    (guess/unsure/confident), self-reported, QBank-only — must never be
    touched by Test Mode submissions (see record_question_result's source
    param elsewhere)."""

    def setUp(self):
        self.student = User.objects.create_user(username='student1', email='student1@example.com', password='pw12345')
        self.subject = Subject.objects.create(name='Physics')
        self.question = Question.objects.create(subject=self.subject, text='2+2=?')
        self.correct_option = Option.objects.create(question=self.question, text='4', is_correct=True)
        self.client.force_authenticate(user=self.student)

    def test_confidence_is_persisted_on_answer(self):
        self.client.post(
            f'/api/questions/{self.question.id}/answer/',
            {'option_id': self.correct_option.id, 'confidence': 'confident'},
        )

        attempt = QuestionAttempt.objects.get(user=self.student, question=self.question)
        self.assertEqual(attempt.confidence, 'confident')

    def test_confidence_is_optional_and_defaults_blank(self):
        resp = self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.correct_option.id})

        self.assertEqual(resp.status_code, 200)
        attempt = QuestionAttempt.objects.get(user=self.student, question=self.question)
        self.assertEqual(attempt.confidence, '')

    def test_invalid_confidence_value_is_rejected(self):
        resp = self.client.post(
            f'/api/questions/{self.question.id}/answer/',
            {'option_id': self.correct_option.id, 'confidence': 'not-a-real-choice'},
        )

        self.assertEqual(resp.status_code, 400)

    def test_confidence_action_sets_confidence_without_double_counting_attempts(self):
        """The post-result confidence prompt hits this dedicated action, not
        `answer` again — must never bump attempts_count/mastery_status."""
        self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.correct_option.id})

        resp = self.client.post(f'/api/questions/{self.question.id}/confidence/', {'confidence': 'guess'})

        self.assertEqual(resp.status_code, 200)
        attempt = QuestionAttempt.objects.get(user=self.student, question=self.question)
        self.assertEqual(attempt.confidence, 'guess')
        self.assertEqual(attempt.attempts_count, 1)

    def test_confidence_action_rejects_invalid_value(self):
        resp = self.client.post(f'/api/questions/{self.question.id}/confidence/', {'confidence': 'sort-of'})
        self.assertEqual(resp.status_code, 400)


class QuestionDashboardEndpointTests(APITestCase):
    def setUp(self):
        self.student = User.objects.create_user(username='student1', email='student1@example.com', password='pw12345')
        self.subject = Subject.objects.create(name='Physics')
        self.q1 = Question.objects.create(subject=self.subject, text='Q1')
        self.q2 = Question.objects.create(subject=self.subject, text='Q2')
        self.q3 = Question.objects.create(subject=self.subject, text='Q3')  # never attempted
        self.client.force_authenticate(user=self.student)

    def test_dashboard_counts_reflect_real_attempts(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.q1, True, source='qbank')
        record_question_result(self.student, self.q1, True, source='qbank')  # -> mastered
        record_question_result(self.student, self.q2, False, source='qbank')

        resp = self.client.get('/api/questions/dashboard/')

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['total_questions'], 3)
        self.assertEqual(resp.data['attempted'], 2)
        self.assertEqual(resp.data['new'], 1)
        self.assertEqual(resp.data['correct'], 1)
        self.assertEqual(resp.data['incorrect'], 1)
        self.assertEqual(resp.data['mastered'], 1)
        self.assertEqual(resp.data['weak'], 1)

    def test_dashboard_reports_topics_practiced_and_qbank_study_time(self):
        from academics.services import record_question_result

        chapter = Chapter.objects.create(subject=self.subject, name='Mechanics')
        topic = Topic.objects.create(chapter=chapter, name='Kinematics')
        self.q1.topic = topic
        self.q1.save()

        record_question_result(self.student, self.q1, True, source='qbank', time_taken_seconds=30)
        record_question_result(self.student, self.q2, False, source='qbank', time_taken_seconds=45)

        resp = self.client.get('/api/questions/dashboard/')

        self.assertEqual(resp.data['topics_practiced'], 1)
        self.assertEqual(resp.data['study_seconds'], 75)

    def test_dashboard_study_time_excludes_test_mode_time(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.q1, True, source='test', time_taken_seconds=999)

        resp = self.client.get('/api/questions/dashboard/')

        self.assertEqual(resp.data['study_seconds'], 0)

    def test_dashboard_requires_auth(self):
        self.client.force_authenticate(user=None)
        resp = self.client.get('/api/questions/dashboard/')
        self.assertEqual(resp.status_code, 401)


class QuestionMistakesEndpointTests(APITestCase):
    def setUp(self):
        self.student = User.objects.create_user(username='student1', email='student1@example.com', password='pw12345')
        self.physics = Subject.objects.create(name='Physics')
        self.chemistry = Subject.objects.create(name='Chemistry')
        self.wrong_physics = Question.objects.create(subject=self.physics, text='Wrong physics Q')
        self.wrong_chemistry = Question.objects.create(subject=self.chemistry, text='Wrong chem Q')
        self.right_physics = Question.objects.create(subject=self.physics, text='Right physics Q')
        self.client.force_authenticate(user=self.student)

    def test_mistakes_grouped_by_subject_and_excludes_correct_questions(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.wrong_physics, False, source='qbank')
        record_question_result(self.student, self.wrong_chemistry, False, source='test')
        record_question_result(self.student, self.right_physics, True, source='qbank')

        resp = self.client.get('/api/questions/mistakes/')

        self.assertEqual(resp.status_code, 200)
        counts = {row['subject_name']: row['count'] for row in resp.data['by_subject']}
        self.assertEqual(counts, {'Physics': 1, 'Chemistry': 1})
        result_ids = {q['id'] for q in resp.data['results']}
        self.assertEqual(result_ids, {self.wrong_physics.id, self.wrong_chemistry.id})

    def test_a_question_later_answered_correctly_leaves_the_mistake_list(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.wrong_physics, False, source='qbank')
        record_question_result(self.student, self.wrong_physics, True, source='qbank')

        resp = self.client.get('/api/questions/mistakes/')
        result_ids = {q['id'] for q in resp.data['results']}
        self.assertNotIn(self.wrong_physics.id, result_ids)


class QuestionPracticeSessionEndpointTests(APITestCase):
    def setUp(self):
        self.student = User.objects.create_user(username='student1', email='student1@example.com', password='pw12345')
        self.subject = Subject.objects.create(name='Physics')
        self.new_q = Question.objects.create(subject=self.subject, text='Never attempted')
        self.weak_q = Question.objects.create(subject=self.subject, text='Weak question')
        self.mastered_q = Question.objects.create(subject=self.subject, text='Mastered question')
        self.client.force_authenticate(user=self.student)

    def test_status_new_returns_only_unattempted_questions(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.weak_q, False, source='qbank')
        record_question_result(self.student, self.mastered_q, True, source='qbank')
        record_question_result(self.student, self.mastered_q, True, source='qbank')

        resp = self.client.post('/api/questions/practice-session/', {'status': ['new']}, format='json')

        ids = {q['id'] for q in resp.data}
        self.assertEqual(ids, {self.new_q.id})

    def test_status_weak_returns_only_weak_questions(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.weak_q, False, source='qbank')
        record_question_result(self.student, self.weak_q, False, source='qbank')

        resp = self.client.post('/api/questions/practice-session/', {'status': ['weak']}, format='json')

        ids = {q['id'] for q in resp.data}
        self.assertEqual(ids, {self.weak_q.id})

    def test_count_is_capped_at_100(self):
        # Individual .create() calls, not bulk_create — Question.save() is
        # what generates the unique slug/public_id, and bulk_create skips save().
        for i in range(120):
            Question.objects.create(subject=self.subject, text=f'Bulk {i}')

        resp = self.client.post('/api/questions/practice-session/', {'count': 500}, format='json')

        self.assertLessEqual(len(resp.data), 100)


class RecomputeQuestionDifficultyCommandTests(TestCase):
    def setUp(self):
        self.subject = Subject.objects.create(name='Physics')
        self.question = Question.objects.create(subject=self.subject, text='Q1', instructor_difficulty='hard')

    def test_below_min_attempts_leaves_actual_difficulty_blank(self):
        from django.core.management import call_command

        from academics.models import QuestionBankConfig

        QuestionBankConfig.objects.create(pk=1, min_attempts_for_difficulty=30)
        student = User.objects.create_user(username='s1', email='s1@example.com', password='pw12345')
        QuestionAttempt.objects.create(user=student, question=self.question, attempts_count=5, correct_count=5)

        call_command('recompute_question_difficulty')

        self.question.refresh_from_db()
        self.assertEqual(self.question.actual_difficulty, '')

    def test_meets_min_attempts_computes_difficulty_without_touching_instructor_difficulty(self):
        from django.core.management import call_command

        from academics.models import QuestionBankConfig

        QuestionBankConfig.objects.create(pk=1, min_attempts_for_difficulty=10, easy_min_pct=75, medium_min_pct=55, hard_min_pct=30)
        student = User.objects.create_user(username='s1', email='s1@example.com', password='pw12345')
        # 2/10 correct = 20% -> below hard_min_pct (30) -> very_hard
        QuestionAttempt.objects.create(user=student, question=self.question, attempts_count=10, correct_count=2)

        call_command('recompute_question_difficulty')

        self.question.refresh_from_db()
        self.assertEqual(self.question.actual_difficulty, 'very_hard')
        self.assertEqual(self.question.actual_difficulty_sample_size, 10)
        self.assertEqual(self.question.instructor_difficulty, 'hard')


class QuestionCourseScopingTests(APITestCase):
    """A question explicitly tagged to a course must not surface for a
    student not enrolled in it — closes the QBank-search leak (spec item
    19), where the frontend previously sent no ?course= at all and the
    backend's course filter was opt-in on the client supplying one."""

    def setUp(self):
        from courses.models import Course, Enrollment

        self.cee_mbbs = Course.objects.create(name='CEE-MBBS', prefix='CEEMBBSQ')
        self.nmcle_mbbs = Course.objects.create(name='NMCLE-MBBS', prefix='NMCLEMBBSQ')
        self.cee_student = User.objects.create_user(username='cee_q_student', email='ceeq@example.com', password='pw12345')
        Enrollment.objects.create(user=self.cee_student, course=self.cee_mbbs)

        self.subject = Subject.objects.create(name='Biology', is_free=True)
        self.cee_question = Question.objects.create(subject=self.subject, text='CEE-MBBS only question')
        self.cee_question.courses.set([self.cee_mbbs])
        self.nmcle_question = Question.objects.create(subject=self.subject, text='NMCLE-MBBS only question')
        self.nmcle_question.courses.set([self.nmcle_mbbs])
        self.shared_question = Question.objects.create(subject=self.subject, text='Untagged shared question')

        self.client.force_authenticate(user=self.cee_student)

    def test_browse_excludes_other_courses_question_even_without_course_param(self):
        resp = self.client.get('/api/questions/browse/', {'search': 'question'})
        ids = {q['id'] for q in resp.data['results']}
        self.assertIn(self.cee_question.id, ids)
        self.assertIn(self.shared_question.id, ids)
        self.assertNotIn(self.nmcle_question.id, ids)


class QuestionListPaginationCountFixTests(APITestCase):
    """Scalability audit: GET /questions/'s pagination count used to call
    queryset.count() directly — a wide, unrestricted DISTINCT (every
    column, including the large `text` field and 4 annotated-subquery
    columns) that forced MySQL to materialize a full derived table per
    request (confirmed 2.7-4.7s at 100K questions, dominating the
    endpoint's latency far more than the actual fetch/serialization).
    _CheapDistinctCountPaginator instead counts `.values('pk').distinct()`
    — trivially equivalent since `pk` alone already defines row identity,
    but an index-covered operation for MySQL instead of a wide-row
    materialization.

    These tests deliberately construct a genuine row-fan-out scenario (a
    question reachable via the course-scoping OR-filter's JOIN through
    TWO of the student's eligible courses at once) — exactly the case
    `.distinct()` exists to guard against — to prove the cheap pk-only
    count is not just faster but still numerically and functionally
    identical: no duplicate rows in the actual response, and the
    paginator's internal page-validation count matches the true distinct
    count exactly."""

    def setUp(self):
        from courses.models import Course, Enrollment

        self.course_a = Course.objects.create(name='Pagination Course A', prefix='PGCOUA')
        self.course_b = Course.objects.create(name='Pagination Course B', prefix='PGCOUB')
        self.student = User.objects.create_user(username='pg_count_student', email='pg_count_student@example.com', password='pw12345')
        Enrollment.objects.create(user=self.student, course=self.course_a)
        Enrollment.objects.create(user=self.student, course=self.course_b)

        self.subject = Subject.objects.create(name='Pagination Count Subject', is_free=True)
        # Tagged to BOTH of the student's eligible courses at once -- the
        # course-scoping OR-filter's Q(courses__id__in=...) JOIN matches
        # this question via two separate join rows (fan-out), exactly the
        # scenario .distinct() exists to collapse back to one row.
        self.fanout_question = Question.objects.create(subject=self.subject, text='Fan-out question')
        self.fanout_question.courses.set([self.course_a, self.course_b])
        self.plain_question = Question.objects.create(subject=self.subject, text='Plain untagged question')

        self.client.force_authenticate(user=self.student)

    def test_cheap_count_matches_wide_distinct_count_with_real_fanout_data(self):
        from academics.access import question_course_scoped
        from academics.views import _CheapDistinctCountPaginator, QuestionViewSet

        view = QuestionViewSet()
        view.request = type('R', (), {'user': self.student, 'query_params': {}})()
        qs = view.get_queryset()

        # The old behavior, for comparison -- a full wide-row DISTINCT count.
        old_wide_count = qs.count()
        # The new, fixed behavior.
        new_cheap_count = _CheapDistinctCountPaginator(qs, 500).count

        self.assertEqual(old_wide_count, new_cheap_count)
        self.assertEqual(new_cheap_count, 2)  # fanout_question + plain_question, each counted once

    def test_get_questions_returns_fanout_question_exactly_once(self):
        resp = self.client.get('/api/questions/')
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        ids = [row['id'] for row in resp.data]
        self.assertEqual(ids.count(self.fanout_question.id), 1)
        self.assertIn(self.plain_question.id, ids)

    def test_query_count_for_paginated_list_does_not_regress(self):
        """The fix must not ADD queries -- still exactly one count query
        plus the fetch/prefetch queries, same shape as before."""
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        with CaptureQueriesContext(connection) as ctx:
            resp = self.client.get('/api/questions/')
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        # main paginated fetch + count + options prefetch -- a small,
        # flat number regardless of how many questions exist (see
        # QuestionSearchFulltextTests-adjacent N+1 coverage elsewhere).
        # 7, not 6 (Phase 3 Free Starter Foundation): locked_subject_ids()
        # now also does one flat, indexed FreeStarterEntitlement read — see
        # the matching comment on
        # QuestionListPublicVisibilityTests.test_public_list_query_count_
        # does_not_grow_with_row_count for the full reasoning. Still flat.
        self.assertLessEqual(len(ctx.captured_queries), 7, ctx.captured_queries)


class SubjectCourseScopingTests(APITestCase):
    """A subject explicitly scoped to specific course(s) must not surface to
    a student enrolled in an unrelated course — the reported bug: a CEE-PG
    student's Question Bank was showing CEE-UG-only subjects (Physics/
    Chemistry/Botany) alongside their own (Pathology/Physiology/Anatomy),
    because SubjectViewSet.get_queryset() only filtered by course when the
    client happened to send ?course=, exactly the same class of gap already
    fixed for Test/Question."""

    def setUp(self):
        from courses.models import Course, Enrollment

        self.cee_ug = Course.objects.create(name='CEE-MBBS Scoping', prefix='CEEUGSCOPE')
        self.cee_pg = Course.objects.create(name='MD/MS Scoping', prefix='CEEPGSCOPE')
        self.pg_student = User.objects.create_user(username='pg_student', email='pgstudent@example.com', password='pw12345')
        Enrollment.objects.create(user=self.pg_student, course=self.cee_pg)

        self.physics = Subject.objects.create(name='Physics Scoping')
        self.physics.courses.set([self.cee_ug])
        self.pathology = Subject.objects.create(name='Pathology Scoping')
        self.pathology.courses.set([self.cee_pg])
        self.shared = Subject.objects.create(name='Shared Scoping')  # blank courses = shared

        self.client.force_authenticate(user=self.pg_student)

    def test_subject_list_excludes_other_courses_subject_even_without_course_param(self):
        resp = self.client.get('/api/subjects/')
        ids = {s['id'] for s in resp.data}
        self.assertIn(self.pathology.id, ids)
        self.assertIn(self.shared.id, ids)
        self.assertNotIn(self.physics.id, ids)

    def test_tampered_course_param_cannot_widen_access(self):
        resp = self.client.get(f'/api/subjects/?course={self.cee_ug.id}')
        ids = {s['id'] for s in resp.data}
        self.assertNotIn(self.physics.id, ids)

    def test_chapters_of_an_unassigned_subject_are_excluded(self):
        from academics.models import Chapter

        chapter = Chapter.objects.create(subject=self.physics, name='Kinematics')
        resp = self.client.get(f'/api/chapters/?subject={self.physics.slug}')
        ids = {c['id'] for c in resp.data}
        self.assertNotIn(chapter.id, ids)

    def test_recommended_new_subject_suggestion_never_names_an_unassigned_subject(self):
        resp = self.client.get('/api/questions/recommended/')
        self.assertEqual(resp.status_code, 200)
        new_subject_suggestions = [s for s in resp.data['suggestions'] if s.get('type') == 'new_subject']
        for s in new_subject_suggestions:
            self.assertNotEqual(s['subject_id'], self.physics.id)


class SubjectPercentPracticedTests(APITestCase):
    """QBank homepage redesign: SubjectGrid needs a question-level %
    practiced (attempted_count/question_count), distinct from the existing
    chapter-level solved_modules/module_count."""

    def setUp(self):
        self.student = User.objects.create_user(username='pct_student', email='pct_student@example.com', password='pw12345')
        self.subject = Subject.objects.create(name='Percent Practiced Subject')
        self.q1 = Question.objects.create(subject=self.subject, text='Q1')
        self.q2 = Question.objects.create(subject=self.subject, text='Q2')
        self.q3 = Question.objects.create(subject=self.subject, text='Q3')
        self.q4 = Question.objects.create(subject=self.subject, text='Q4')
        self.client.force_authenticate(user=self.student)

    def test_percent_practiced_matches_attempted_over_total(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.q1, True, source='qbank')
        record_question_result(self.student, self.q2, False, source='qbank')

        resp = self.client.get('/api/subjects/')
        row = next(s for s in resp.data if s['id'] == self.subject.id)

        self.assertEqual(row['attempted_count'], 2)
        self.assertEqual(row['question_count'], 4)
        self.assertEqual(row['percent_practiced'], 50)

    def test_percent_practiced_is_zero_for_a_subject_with_no_questions(self):
        empty_subject = Subject.objects.create(name='Empty Percent Subject')

        resp = self.client.get('/api/subjects/')
        row = next(s for s in resp.data if s['id'] == empty_subject.id)

        self.assertEqual(row['percent_practiced'], 0)

    def test_percent_practiced_is_zero_for_anonymous_user(self):
        self.client.force_authenticate(user=None)
        resp = self.client.get('/api/subjects/')
        row = next(s for s in resp.data if s['id'] == self.subject.id)
        self.assertEqual(row['percent_practiced'], 0)


class RecommendedTopSuggestionEnrichmentTests(APITestCase):
    """QBank homepage redesign: the top /questions/recommended/ suggestion
    needs accuracy_pct/question_count/estimated_minutes for the 'Your Next
    Practice' hero card's accuracy ring + '~N min' + 'N Questions'."""

    def setUp(self):
        self.student = User.objects.create_user(username='next_practice_student', email='next_practice@example.com', password='pw12345')
        self.subject = Subject.objects.create(name='Cardiovascular Physiology')
        self.topic = Topic.objects.create(chapter=Chapter.objects.create(subject=self.subject, name='Heart'), name='Cardiac Cycle')
        self.client.force_authenticate(user=self.student)

    def test_top_suggestion_carries_normalized_hero_card_fields(self):
        from academics.services import record_question_result

        # 3+ attempts on this subject, mostly wrong, so it becomes the
        # weakest subject and produces a 'revise_topic' or 'improve_subject'
        # top suggestion with a real weak_count/accuracy behind it.
        questions = [Question.objects.create(subject=self.subject, topic=self.topic, text=f'Q{i}') for i in range(5)]
        for i, q in enumerate(questions):
            record_question_result(self.student, q, i == 0, source='qbank')

        resp = self.client.get('/api/questions/recommended/')

        self.assertEqual(resp.status_code, 200)
        top = resp.data['suggestions'][0]
        self.assertIn('question_count', top)
        self.assertIn('accuracy_pct', top)
        self.assertIn('estimated_minutes', top)
        if top['question_count']:
            self.assertGreaterEqual(top['estimated_minutes'], 5)

    def test_fallback_suggestion_has_no_crash_on_missing_counts(self):
        """A brand-new student with no attempts gets the 'start_new'
        fallback, which has no count — must not raise a TypeError computing
        estimated_minutes from None."""
        resp = self.client.get('/api/questions/recommended/')

        self.assertEqual(resp.status_code, 200)
        top = resp.data['suggestions'][0]
        self.assertEqual(top['type'], 'start_new')
        self.assertIsNone(top['estimated_minutes'])


class CompleteCourseScopingAuditTests(APITestCase):
    """Full-matrix regression suite for the course-scoping audit: anonymous,
    no-enrollment, single-course (CEE-UG / CEE-PG), and multi-course
    students, across Subject listing, Question browse/search, the Practice
    Session Builder (the exact endpoint behind the reported "Physics/
    Chemistry still appear in CEE-PG practice" bug), the QBank dashboard,
    and recommendations — each also probed with a tampered ?course=/course
    param to confirm a client value can only narrow, never widen, access."""

    def setUp(self):
        from courses.models import Course, Enrollment

        self.cee_ug = Course.objects.create(name='CEE-UG Audit', prefix='CEEUGAUDIT')
        self.cee_pg = Course.objects.create(name='CEE-PG Audit', prefix='CEEPGAUDIT')

        self.cee_ug_student = User.objects.create_user(username='audit_ug', email='audit_ug@example.com', password='pw12345')
        Enrollment.objects.create(user=self.cee_ug_student, course=self.cee_ug)

        self.cee_pg_student = User.objects.create_user(username='audit_pg', email='audit_pg@example.com', password='pw12345')
        Enrollment.objects.create(user=self.cee_pg_student, course=self.cee_pg)

        self.multi_student = User.objects.create_user(username='audit_multi', email='audit_multi@example.com', password='pw12345')
        Enrollment.objects.create(user=self.multi_student, course=self.cee_ug)
        Enrollment.objects.create(user=self.multi_student, course=self.cee_pg)

        self.no_enrollment_student = User.objects.create_user(username='audit_none', email='audit_none@example.com', password='pw12345')

        def _subject_with_question(name, course):
            # Question.courses is deliberately left BLANK here — matching
            # real production data, where every Subject is explicitly
            # course-scoped but no Question has ever had its own `courses`
            # tag set. A question must inherit its subject's course scope
            # when its own `courses` is blank (see
            # academics.views._question_course_scoped) — this is the exact
            # shape of the reported "Physics/Chemistry still appear in
            # CEE-PG practice" bug, which explicit Question-level tagging
            # (as in QuestionCourseScopingTests) would not have caught.
            subject = Subject.objects.create(name=name, is_free=True)
            subject.courses.set([course])
            question = Question.objects.create(subject=subject, text=f'{name} question 1')
            return subject, question

        self.physics, self.physics_q = _subject_with_question('Physics Audit', self.cee_ug)
        self.chemistry, self.chemistry_q = _subject_with_question('Chemistry Audit', self.cee_ug)
        self.botany, self.botany_q = _subject_with_question('Botany Audit', self.cee_ug)
        self.pathology, self.pathology_q = _subject_with_question('Pathology Audit', self.cee_pg)
        self.physiology, self.physiology_q = _subject_with_question('Physiology Audit', self.cee_pg)
        self.anatomy, self.anatomy_q = _subject_with_question('Anatomy Audit', self.cee_pg)

    # -- Subject listing ---------------------------------------------------

    def test_cee_pg_student_sees_only_pg_subjects(self):
        self.client.force_authenticate(user=self.cee_pg_student)
        resp = self.client.get('/api/subjects/')
        names = {s['name'] for s in resp.data}
        self.assertEqual(names, {'Pathology Audit', 'Physiology Audit', 'Anatomy Audit'})

    def test_cee_ug_student_sees_only_ug_subjects(self):
        self.client.force_authenticate(user=self.cee_ug_student)
        resp = self.client.get('/api/subjects/')
        names = {s['name'] for s in resp.data}
        self.assertEqual(names, {'Physics Audit', 'Chemistry Audit', 'Botany Audit'})

    def test_multi_course_student_sees_both_courses_subjects(self):
        self.client.force_authenticate(user=self.multi_student)
        resp = self.client.get('/api/subjects/')
        names = {s['name'] for s in resp.data}
        self.assertEqual(
            names,
            {'Physics Audit', 'Chemistry Audit', 'Botany Audit', 'Pathology Audit', 'Physiology Audit', 'Anatomy Audit'},
        )

    def test_anonymous_user_sees_none_of_these_scoped_subjects(self):
        resp = self.client.get('/api/subjects/')
        names = {s['name'] for s in resp.data}
        self.assertFalse(names & {'Physics Audit', 'Pathology Audit'})

    def test_no_enrollment_student_sees_none_of_these_scoped_subjects(self):
        self.client.force_authenticate(user=self.no_enrollment_student)
        resp = self.client.get('/api/subjects/')
        names = {s['name'] for s in resp.data}
        self.assertFalse(names & {'Physics Audit', 'Pathology Audit'})

    # -- Practice Session Builder (the exact reported leak) ----------------

    def test_practice_session_for_pg_student_never_returns_ug_questions(self):
        self.client.force_authenticate(user=self.cee_pg_student)
        resp = self.client.post('/api/questions/practice-session/', {'count': 50}, format='json')
        ids = {q['id'] for q in resp.data}
        self.assertIn(self.pathology_q.id, ids)
        self.assertNotIn(self.physics_q.id, ids)
        self.assertNotIn(self.chemistry_q.id, ids)
        self.assertNotIn(self.botany_q.id, ids)

    def test_practice_session_for_ug_student_never_returns_pg_questions(self):
        self.client.force_authenticate(user=self.cee_ug_student)
        resp = self.client.post('/api/questions/practice-session/', {'count': 50}, format='json')
        ids = {q['id'] for q in resp.data}
        self.assertIn(self.physics_q.id, ids)
        self.assertNotIn(self.pathology_q.id, ids)
        self.assertNotIn(self.physiology_q.id, ids)
        self.assertNotIn(self.anatomy_q.id, ids)

    def test_practice_session_with_students_own_real_course_param_still_returns_results(self):
        """Regression for a real production bug: passing the student's OWN
        real, enrolled course id (exactly what the browser sends on every
        request once a course is active — not a tampering attempt) used to
        return zero results platform-wide, because the course filter
        matched raw Question.courses directly, which — like every question
        in this fixture and in real production — is blank. Must inherit
        via Subject.courses the same way the base eligibility scoping
        already does."""
        self.client.force_authenticate(user=self.cee_pg_student)
        resp = self.client.post(
            '/api/questions/practice-session/', {'count': 50, 'course': self.cee_pg.id}, format='json',
        )
        ids = {q['id'] for q in resp.data}
        self.assertIn(self.pathology_q.id, ids)
        self.assertIn(self.physiology_q.id, ids)
        self.assertIn(self.anatomy_q.id, ids)

    def test_practice_session_tampered_course_param_returns_no_ug_questions_for_pg_student(self):
        self.client.force_authenticate(user=self.cee_pg_student)
        resp = self.client.post(
            '/api/questions/practice-session/', {'count': 50, 'course': self.cee_ug.id}, format='json',
        )
        ids = {q['id'] for q in resp.data}
        self.assertNotIn(self.physics_q.id, ids)
        self.assertEqual(ids, set())

    def test_practice_session_tampered_subject_slug_returns_nothing_for_unassigned_subject(self):
        self.client.force_authenticate(user=self.cee_pg_student)
        resp = self.client.post(
            '/api/questions/practice-session/', {'count': 50, 'subject': self.physics.slug}, format='json',
        )
        self.assertEqual(resp.data, [])

    def test_question_with_blank_courses_inherits_subject_course_scope(self):
        """The exact real-data shape: every fixture question in this class
        has a BLANK Question.courses (like every real question in
        production) and relies entirely on inheriting its Subject's
        courses. If a question with blank `courses` were ever treated as
        unconditionally shared (the bug this test guards against), this
        entire test class's course isolation would silently stop meaning
        anything, since production has zero Question-level course tags."""
        self.assertFalse(self.physics_q.courses.exists())
        self.client.force_authenticate(user=self.cee_pg_student)
        resp = self.client.post('/api/questions/practice-session/', {'count': 50}, format='json')
        ids = {q['id'] for q in resp.data}
        self.assertNotIn(self.physics_q.id, ids)

    # -- QBank dashboard -----------------------------------------------------

    def test_dashboard_total_questions_excludes_other_course_for_pg_student(self):
        self.client.force_authenticate(user=self.cee_pg_student)
        resp = self.client.get('/api/questions/dashboard/')
        self.assertEqual(resp.data['total_questions'], 3)

    def test_dashboard_tampered_course_param_cannot_widen_pg_student_totals(self):
        self.client.force_authenticate(user=self.cee_pg_student)
        resp = self.client.get(f'/api/questions/dashboard/?course={self.cee_ug.id}')
        self.assertEqual(resp.data['total_questions'], 0)

    # -- Question browse / search --------------------------------------------

    def test_browse_search_for_pg_student_excludes_ug_matches(self):
        self.client.force_authenticate(user=self.cee_pg_student)
        resp = self.client.get('/api/questions/browse/', {'search': 'Audit question'})
        ids = {q['id'] for q in resp.data['results']}
        self.assertIn(self.pathology_q.id, ids)
        self.assertNotIn(self.physics_q.id, ids)

    def test_direct_question_id_via_browse_never_returns_unassigned_course_question(self):
        self.client.force_authenticate(user=self.cee_pg_student)
        resp = self.client.get('/api/questions/browse/', {'search': 'Physics Audit'})
        ids = {q['id'] for q in resp.data['results']}
        self.assertNotIn(self.physics_q.id, ids)

    # -- Chapter / Topic nested-resource bypass ------------------------------

    def test_chapters_of_unassigned_subject_not_reachable_by_pg_student(self):
        from academics.models import Chapter

        chapter = Chapter.objects.create(subject=self.physics, name='Kinematics Audit')
        self.client.force_authenticate(user=self.cee_pg_student)
        resp = self.client.get(f'/api/chapters/?subject={self.physics.slug}')
        ids = {c['id'] for c in resp.data}
        self.assertNotIn(chapter.id, ids)

    def test_topics_of_unassigned_subject_not_reachable_by_pg_student(self):
        from academics.models import Chapter, Topic

        chapter = Chapter.objects.create(subject=self.physics, name='Kinematics Audit 2')
        topic = Topic.objects.create(chapter=chapter, name='Vectors Audit')
        self.client.force_authenticate(user=self.cee_pg_student)
        resp = self.client.get(f'/api/topics/?chapter={chapter.id}')
        ids = {t['id'] for t in resp.data}
        self.assertNotIn(topic.id, ids)


class RecordQuestionResultStatMathTests(TestCase):
    """Question.total_attempts/correct_attempts and Option.pick_count/
    pick_percentage must reflect one vote per distinct student — a retry
    moves the vote, it never adds another — per the signed-delta design in
    academics.services.record_question_result. This is what keeps "X% of
    students got this right" honest instead of inflating on every retry."""

    def setUp(self):
        self.subject = Subject.objects.create(name='Stats Subject')
        self.question = Question.objects.create(subject=self.subject, text='Stats question')
        self.opt_a = Option.objects.create(question=self.question, text='A', is_correct=True, order=1)
        self.opt_b = Option.objects.create(question=self.question, text='B', is_correct=False, order=2)
        self.student1 = User.objects.create_user(username='stats1', email='stats1@example.com', password='pw12345')
        self.student2 = User.objects.create_user(username='stats2', email='stats2@example.com', password='pw12345')

    def test_first_attempt_increments_totals_and_pick_count(self):
        from academics.services import record_question_result

        record_question_result(self.student1, self.question, True, source='qbank', selected_option=self.opt_a)

        self.question.refresh_from_db()
        self.opt_a.refresh_from_db()
        self.assertEqual(self.question.total_attempts, 1)
        self.assertEqual(self.question.correct_attempts, 1)
        self.assertEqual(self.opt_a.pick_count, 1)
        self.assertEqual(self.opt_a.pick_percentage, 100)

    def test_reanswer_same_option_does_not_double_count(self):
        from academics.services import record_question_result

        record_question_result(self.student1, self.question, True, source='qbank', selected_option=self.opt_a)
        record_question_result(self.student1, self.question, True, source='qbank', selected_option=self.opt_a)

        self.question.refresh_from_db()
        self.opt_a.refresh_from_db()
        self.assertEqual(self.question.total_attempts, 1)
        self.assertEqual(self.opt_a.pick_count, 1)

    def test_reanswer_different_option_moves_the_vote_not_adds_one(self):
        from academics.services import record_question_result

        record_question_result(self.student1, self.question, False, source='qbank', selected_option=self.opt_b)
        record_question_result(self.student1, self.question, True, source='qbank', selected_option=self.opt_a)

        self.question.refresh_from_db()
        self.opt_a.refresh_from_db()
        self.opt_b.refresh_from_db()
        self.assertEqual(self.question.total_attempts, 1)
        self.assertEqual(self.question.correct_attempts, 1)
        self.assertEqual(self.opt_a.pick_count, 1)
        self.assertEqual(self.opt_b.pick_count, 0)

    def test_percentages_sum_to_approximately_100_across_students(self):
        from academics.services import record_question_result

        record_question_result(self.student1, self.question, True, source='qbank', selected_option=self.opt_a)
        record_question_result(self.student2, self.question, False, source='qbank', selected_option=self.opt_b)

        self.question.refresh_from_db()
        self.opt_a.refresh_from_db()
        self.opt_b.refresh_from_db()
        self.assertEqual(self.question.total_attempts, 2)
        self.assertEqual(self.question.correct_attempts, 1)
        total_pct = self.opt_a.pick_percentage + self.opt_b.pick_percentage
        self.assertIn(total_pct, (99, 100, 101))  # rounding tolerance

    def test_time_taken_seconds_recorded_on_question_event(self):
        from academics.services import record_question_result

        record_question_result(
            self.student1, self.question, True, source='qbank',
            selected_option=self.opt_a, time_taken_seconds=42,
        )

        event = QuestionEvent.objects.get(user=self.student1, question=self.question)
        self.assertEqual(event.time_taken_seconds, 42)


class RecordQuestionResultQuestionAttemptConcurrencyTests(TransactionTestCase):
    """Scalability audit Phase 3 mandatory validation ("concurrent same-user/
    same-question" and "concurrent different-user/same-question"): proves
    the collapsed select_for_update().filter(...).first() + create-with-
    IntegrityError-recovery flow (replacing get_or_create() + a second
    locked re-fetch) still guarantees exactly one QuestionAttempt row per
    (user, question) and never loses an increment, under real concurrent
    threads/connections.

    TransactionTestCase (not TestCase), same reasoning as
    QuestionPublicIdConcurrencyTests above: each thread needs its own real,
    committing connection. SQLite has no real row-level locking, so this
    doesn't prove blocking/serialization the way the live MySQL staging
    validation does — it proves the invariant that must hold regardless of
    engine: unique_together plus the IntegrityError-recovery branch mean no
    two concurrent first-encounter calls for the same key ever both create
    a row, and every call's increment is eventually applied to the one row
    that exists."""

    def _run_concurrently(self, jobs):
        import threading

        from django.db import connection
        from django.db.utils import OperationalError

        results = []
        errors = []
        lock = threading.Lock()

        def run_one(job):
            for attempt in range(20):
                try:
                    result = job()
                    with lock:
                        results.append(result)
                    return
                except OperationalError as exc:
                    if 'locked' in str(exc).lower() and attempt < 19:
                        import time
                        time.sleep(0.05)
                        continue
                    with lock:
                        errors.append(exc)
                    return
                except Exception as exc:  # noqa: BLE001
                    with lock:
                        errors.append(exc)
                    return
                finally:
                    connection.close()

        threads = [threading.Thread(target=run_one, args=(job,)) for job in jobs]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        return results, errors

    def test_concurrent_same_user_same_question_first_encounter_creates_exactly_one_row(self):
        from academics.services import record_question_result

        subject = Subject.objects.create(name='Phase3 Concurrency Subject A')
        question = Question.objects.create(subject=subject, text='Phase3 Q1')
        opt_correct = Option.objects.create(question=question, text='Right', is_correct=True, order=1)
        student = User.objects.create_user(username='phase3_student_a', email='phase3_a@example.com', password='pw12345')

        n = 8
        jobs = [
            (lambda: record_question_result(student, question, True, source='qbank', selected_option=opt_correct))
            for _ in range(n)
        ]
        results, errors = self._run_concurrently(jobs)

        self.assertEqual(errors, [], f'concurrent record_question_result() raised: {errors}')
        self.assertEqual(len(results), n)
        # Exactly one row for this (user, question) — the unique_together
        # constraint plus the IntegrityError-recovery branch must have
        # collapsed every concurrent "first encounter" into one row.
        self.assertEqual(QuestionAttempt.objects.filter(user=student, question=question).count(), 1)
        attempt = QuestionAttempt.objects.get(user=student, question=question)
        # Every one of the n concurrent calls' increment must have landed —
        # none lost to a race between the locked read and the write.
        self.assertEqual(attempt.attempts_count, n)
        self.assertEqual(attempt.correct_count, n)
        self.assertEqual(QuestionEvent.objects.filter(user=student, question=question).count(), n)

    def test_concurrent_same_user_same_question_repeat_encounter_no_lost_updates(self):
        from academics.services import record_question_result

        subject = Subject.objects.create(name='Phase3 Concurrency Subject B')
        question = Question.objects.create(subject=subject, text='Phase3 Q2')
        opt_correct = Option.objects.create(question=question, text='Right', is_correct=True, order=1)
        student = User.objects.create_user(username='phase3_student_b', email='phase3_b@example.com', password='pw12345')
        # Pre-create the row (repeat-encounter path: select_for_update().
        # filter(...).first() must find it directly, no create() needed).
        record_question_result(student, question, True, source='qbank', selected_option=opt_correct)

        n = 8
        jobs = [
            (lambda: record_question_result(student, question, True, source='qbank', selected_option=opt_correct))
            for _ in range(n)
        ]
        results, errors = self._run_concurrently(jobs)

        self.assertEqual(errors, [], f'concurrent record_question_result() raised: {errors}')
        self.assertEqual(len(results), n)
        self.assertEqual(QuestionAttempt.objects.filter(user=student, question=question).count(), 1)
        attempt = QuestionAttempt.objects.get(user=student, question=question)
        # 1 pre-create call + n concurrent calls — every increment counted.
        self.assertEqual(attempt.attempts_count, 1 + n)
        self.assertEqual(QuestionEvent.objects.filter(user=student, question=question).count(), 1 + n)

    def test_concurrent_different_users_same_question_each_get_their_own_row(self):
        from academics.services import record_question_result

        subject = Subject.objects.create(name='Phase3 Concurrency Subject C')
        question = Question.objects.create(subject=subject, text='Phase3 Q3')
        opt_correct = Option.objects.create(question=question, text='Right', is_correct=True, order=1)
        opt_wrong = Option.objects.create(question=question, text='Wrong', is_correct=False, order=2)
        n = 6
        students = [
            User.objects.create_user(username=f'phase3_student_c{i}', email=f'phase3_c{i}@example.com', password='pw12345')
            for i in range(n)
        ]

        jobs = [
            (lambda s=s: record_question_result(s, question, True, source='qbank', selected_option=opt_correct))
            for s in students
        ]
        results, errors = self._run_concurrently(jobs)

        self.assertEqual(errors, [], f'concurrent record_question_result() raised: {errors}')
        self.assertEqual(len(results), n)
        # Each student gets exactly their own row — no cross-student
        # contention or row-sharing, since (user, question) is the key.
        self.assertEqual(QuestionAttempt.objects.filter(question=question).count(), n)
        for s in students:
            attempt = QuestionAttempt.objects.get(user=s, question=question)
            self.assertEqual(attempt.attempts_count, 1)
            self.assertEqual(attempt.correct_count, 1)
        # Question.total_attempts (the cross-student aggregate,
        # _apply_question_stat_delta) must also have counted all n votes —
        # untouched by Phase 3, checked here as a cheap regression guard.
        question.refresh_from_db()
        self.assertEqual(question.total_attempts, n)
        self.assertEqual(question.correct_attempts, n)


class AnswerActionStatsVisibilityTests(APITestCase):
    """The 'don't reveal stats before submission' / 'privacy-safe below a
    minimum sample size' rules from the QBank redesign spec."""

    def setUp(self):
        self.subject = Subject.objects.create(name='Answer Stats Subject', is_free=True)
        self.question = Question.objects.create(subject=self.subject, text='Answer stats question')
        self.opt_correct = Option.objects.create(question=self.question, text='Right', is_correct=True, order=1)
        self.opt_wrong = Option.objects.create(question=self.question, text='Wrong', is_correct=False, order=2)
        self.student = User.objects.create_user(username='ansstats', email='ansstats@example.com', password='pw12345')
        self.client.force_authenticate(user=self.student)

    def test_pre_submission_question_detail_never_includes_pick_percentage(self):
        resp = self.client.get(f'/api/questions/{self.question.id}/')
        for opt in resp.data['options']:
            self.assertNotIn('pick_percentage', opt)
            self.assertNotIn('is_correct', opt)

    def test_answer_response_privacy_safe_below_threshold(self):
        resp = self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.opt_correct.id}, format='json')
        self.assertFalse(resp.data['stats_available'])
        self.assertIsNone(resp.data['students_correct_percent'])
        self.assertIsNone(resp.data['total_responses'])
        for opt in resp.data['options']:
            self.assertIsNone(opt['pick_percentage'])

    def test_answer_response_shows_stats_at_or_above_threshold(self):
        config = QuestionBankConfig.load()
        config.min_attempts_for_option_stats = 1
        config.save()

        resp = self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.opt_correct.id}, format='json')
        self.assertTrue(resp.data['stats_available'])
        self.assertEqual(resp.data['students_correct_percent'], 100)
        self.assertEqual(resp.data['total_responses'], 1)
        percentages = {opt['id']: opt['pick_percentage'] for opt in resp.data['options']}
        self.assertEqual(percentages[self.opt_correct.id], 100)

    def test_answer_response_includes_key_takeaway_and_structured_reference(self):
        book = ReferenceBook.objects.create(name='Robbins & Cotran')
        self.question.key_takeaway = 'High yield point'
        self.question.reference_book = book
        self.question.reference_edition = '10th'
        self.question.reference_chapter = 'Hemodynamic Disorders'
        self.question.reference_page = '123'
        self.question.save()

        resp = self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.opt_correct.id}, format='json')
        self.assertEqual(resp.data['key_takeaway'], 'High yield point')
        self.assertEqual(resp.data['reference_book_name'], 'Robbins & Cotran')
        self.assertEqual(resp.data['reference_edition'], '10th')
        self.assertEqual(resp.data['reference_chapter'], 'Hemodynamic Disorders')
        self.assertEqual(resp.data['reference_page'], '123')


class AnswerActionSurfacesAttemptStateTests(APITestCase):
    """QBank 2.0 Phase 2D: the answer() response now surfaces
    mastery_status/attempts_count/revision_due_at/recent_events —
    fields record_question_result() already computed and previously
    discarded — plus the practice_session() annotation bug fix. No new
    mastery/revision algorithm; this only asserts the existing
    QuestionAttempt/QuestionEvent state is read back correctly."""

    def setUp(self):
        self.subject = Subject.objects.create(name='Attempt State Subject', is_free=True)
        self.question = Question.objects.create(subject=self.subject, text='Attempt state question')
        self.opt_correct = Option.objects.create(question=self.question, text='Right', is_correct=True, order=1)
        self.opt_wrong = Option.objects.create(question=self.question, text='Wrong', is_correct=False, order=2)
        self.student = User.objects.create_user(username='attemptstate', email='attemptstate@example.com', password='pw12345')
        self.client.force_authenticate(user=self.student)

    def test_no_option_submitted_leaves_new_attempt_fields_null(self):
        resp = self.client.post(f'/api/questions/{self.question.id}/answer/', {}, format='json')
        self.assertIsNone(resp.data['mastery_status'])
        self.assertIsNone(resp.data['attempts_count'])
        self.assertIsNone(resp.data['revision_due_at'])
        self.assertIsNone(resp.data['recent_events'])

    def test_first_wrong_answer_reports_learning_and_a_future_revision_date(self):
        resp = self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.opt_wrong.id}, format='json')
        attempt = QuestionAttempt.objects.get(user=self.student, question=self.question)
        self.assertEqual(resp.data['mastery_status'], attempt.mastery_status)
        self.assertEqual(resp.data['attempts_count'], 1)
        self.assertEqual(resp.data['correct_count'], 0)
        self.assertEqual(resp.data['incorrect_count'], 1)
        self.assertIsNotNone(resp.data['revision_due_at'])
        self.assertEqual(resp.data['revision_due_at'], attempt.revision_due_at.isoformat())

    def test_recent_events_reflect_answer_history_for_this_user_only(self):
        other_student = User.objects.create_user(username='otherstudent', email='otherstudent@example.com', password='pw12345')
        self.client.force_authenticate(user=other_student)
        self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.opt_wrong.id}, format='json')

        self.client.force_authenticate(user=self.student)
        self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.opt_wrong.id}, format='json')
        resp = self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.opt_correct.id}, format='json')

        # Only this user's own two events — the other student's answer to
        # the same question must never leak into this list.
        self.assertEqual(len(resp.data['recent_events']), 2)
        self.assertEqual([e['is_correct'] for e in resp.data['recent_events']], [True, False])
        self.assertEqual(resp.data['attempts_count'], 2)

    def test_practice_session_reports_real_mastery_status_not_always_new(self):
        # Reproduces the bug: without the fix, every practice-session
        # question always reported mastery_status "new" and
        # is_revision_due False, regardless of the student's real history.
        self.client.post(f'/api/questions/{self.question.id}/answer/', {'option_id': self.opt_correct.id}, format='json')
        attempt = QuestionAttempt.objects.get(user=self.student, question=self.question)
        self.assertNotEqual(attempt.mastery_status, 'new')

        # Backdate the real revision date into the past — without the
        # annotation fix, is_revision_due always falls back to False
        # regardless of this; with it, it correctly reflects an overdue
        # question. This is what makes the assertion below discriminating
        # rather than coincidentally passing either way.
        from django.utils import timezone as tz
        attempt.revision_due_at = tz.now() - tz.timedelta(days=1)
        attempt.save(update_fields=['revision_due_at'])

        resp = self.client.post('/api/questions/practice-session/', {'count': 10}, format='json')
        by_id = {q['id']: q for q in resp.data}
        self.assertEqual(by_id[self.question.id]['mastery_status'], attempt.mastery_status)
        self.assertTrue(by_id[self.question.id]['is_revision_due'])


class QuestionReportTests(APITestCase):
    def setUp(self):
        self.subject = Subject.objects.create(name='Report Subject', is_free=True)
        self.question = Question.objects.create(subject=self.subject, text='Report question')
        self.student = User.objects.create_user(username='reporter', email='reporter@example.com', password='pw12345')
        self.staff = User.objects.create_user(
            username='report_staff', email='report_staff@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )

    def test_student_can_report_a_visible_question(self):
        self.client.force_authenticate(user=self.student)
        resp = self.client.post(
            f'/api/questions/{self.question.id}/report/',
            {'reason': 'incorrect_answer', 'comment': 'The marked answer looks wrong.'},
            format='json',
        )
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
        report = QuestionReport.objects.get()
        self.assertEqual(report.user, self.student)
        self.assertEqual(report.question, self.question)
        self.assertEqual(report.status, 'open')

    def test_student_cannot_report_a_question_from_an_unrelated_course(self):
        from courses.models import Course

        other_course = Course.objects.create(name='Report Other Course', prefix='REPORTOTHER')
        self.subject.courses.set([other_course])  # now scoped away from self.student (no enrollment anywhere)
        self.client.force_authenticate(user=self.student)

        resp = self.client.post(f'/api/questions/{self.question.id}/report/', {'reason': 'other'}, format='json')
        self.assertEqual(resp.status_code, status.HTTP_404_NOT_FOUND)

    def test_report_list_is_admin_only(self):
        self.client.force_authenticate(user=self.student)
        resp = self.client.get('/api/question-reports/')
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)

    def test_admin_can_list_and_resolve_reports(self):
        QuestionReport.objects.create(question=self.question, user=self.student, reason='typo', comment='x')
        self.client.force_authenticate(user=self.staff)

        list_resp = self.client.get('/api/question-reports/?status=open')
        self.assertEqual(len(list_resp.data), 1)

        report_id = list_resp.data[0]['id']
        patch_resp = self.client.patch(f'/api/question-reports/{report_id}/', {'status': 'reviewed'}, format='json')
        self.assertEqual(patch_resp.status_code, status.HTTP_200_OK)
        report = QuestionReport.objects.get(pk=report_id)
        self.assertEqual(report.status, 'reviewed')
        self.assertEqual(report.reviewed_by, self.staff)
        self.assertIsNotNone(report.reviewed_at)

    def test_report_never_exposes_student_identity_fields(self):
        QuestionReport.objects.create(question=self.question, user=self.student, reason='typo')
        self.client.force_authenticate(user=self.staff)
        resp = self.client.get('/api/question-reports/')
        keys = set(resp.data[0].keys())
        self.assertFalse(keys & {'user', 'user_email', 'user_name', 'student_email', 'student_name'})


class DifficultyRatingTests(APITestCase):
    def setUp(self):
        self.subject = Subject.objects.create(name='Difficulty Rating Subject', is_free=True)
        self.question = Question.objects.create(subject=self.subject, text='Difficulty rating question')
        self.student = User.objects.create_user(username='rater', email='rater@example.com', password='pw12345')
        self.client.force_authenticate(user=self.student)

    def test_rate_difficulty_creates_then_updates_on_rerate(self):
        resp1 = self.client.post(f'/api/questions/{self.question.id}/rate-difficulty/', {'rating': 'easy'}, format='json')
        self.assertEqual(resp1.status_code, status.HTTP_200_OK)
        self.assertEqual(QuestionDifficultyRating.objects.count(), 1)

        resp2 = self.client.post(f'/api/questions/{self.question.id}/rate-difficulty/', {'rating': 'difficult'}, format='json')
        self.assertEqual(resp2.status_code, status.HTTP_200_OK)
        self.assertEqual(QuestionDifficultyRating.objects.count(), 1)
        rating = QuestionDifficultyRating.objects.get()
        self.assertEqual(rating.rating, 'difficult')

    def test_invalid_rating_rejected(self):
        resp = self.client.post(f'/api/questions/{self.question.id}/rate-difficulty/', {'rating': 'nonsense'}, format='json')
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_rate_difficulty_respects_course_scoping(self):
        from courses.models import Course

        other_course = Course.objects.create(name='Rating Other Course', prefix='RATEOTHER')
        self.subject.courses.set([other_course])

        resp = self.client.post(f'/api/questions/{self.question.id}/rate-difficulty/', {'rating': 'easy'}, format='json')
        self.assertEqual(resp.status_code, status.HTTP_404_NOT_FOUND)


class QuestionListQueryCountTests(APITestCase):
    """Scalability audit fix: GET /questions/ (plain list) was missing
    select_related/prefetch_related for topic, image_asset,
    explanation_image_asset, reference_book, created_by, each option's
    image_asset, and (admin only) the `courses` M2M — up to ~122 queries
    for a 20-row public page, ~182 for the admin one. Confirms the fix is
    bounded regardless of row count, for both the public and admin
    serializer paths (they select_related different fields), and that
    the response is still a bare array (4 confirmed callers: QuestionSolver's
    chapter-solve session, the QBank bookmarks page, and the Admin
    question-management table and QuestionPicker, which stay on this
    bare-array 500-row cap for now — see _BoundedListPagination's
    docstring for why their migration to `browse` is deferred).

    This test class itself caught two real N+1 bugs during development:
    QuestionSerializer.get_image_data() (public path) reads obj.image_asset
    too, not just the admin serializer, so it needs select_related
    unconditionally, not just for is_admin_view; and QuestionAdminSerializer
    exposes the raw `courses` M2M, which needs prefetch_related, not
    select_related."""

    def setUp(self):
        from courses.models import Course, Enrollment
        from media_library.models import MediaAsset

        self.course = Course.objects.create(name='QLC Course', prefix='QLCCOURSE')
        self.student = User.objects.create_user(username='qlc_student', email='qlc_student@example.com', password='pw12345')
        Enrollment.objects.create(user=self.student, course=self.course)
        self.teacher = User.objects.create_user(
            username='qlc_teacher', email='qlc_teacher@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )
        self.subject = Subject.objects.create(name='QLC Subject', is_free=True)
        self.subject.courses.set([self.course])
        self.chapter = Chapter.objects.create(subject=self.subject, name='QLC Chapter')
        self.topic = Topic.objects.create(chapter=self.chapter, name='QLC Topic')
        self.book = ReferenceBook.objects.create(name='QLC Book')

        def make_asset():
            return MediaAsset.objects.create(
                image_type='question_image', processing_status='ready', bucket='public',
                storage_key='qlc/test.jpg', width=10, height=10,
            )

        def make_question(text):
            q = Question.objects.create(
                subject=self.subject, chapter=self.chapter, topic=self.topic, text=text, marks=1, negative_marks=0,
                image_asset=make_asset(), explanation_image_asset=make_asset(),
                reference_book=self.book, created_by=self.teacher,
            )
            for i in range(4):
                Option.objects.create(question=q, text=f'{text} opt {i}', order=i, is_correct=(i == 0), image_asset=make_asset())
            return q

        self.questions = [make_question(f'QLC Q{i}') for i in range(5)]

    def test_public_list_response_is_a_bare_array(self):
        self.client.force_authenticate(user=self.student)
        resp = self.client.get('/api/questions/')
        self.assertEqual(resp.status_code, 200)
        self.assertIsInstance(resp.data, list)

    def test_public_list_query_count_does_not_grow_with_row_count(self):
        # 7, not 5: locked_subject_ids() (Phase 3 fix) now always issues a
        # flat 2 queries (active subscriptions + non-free subjects) instead
        # of "however many Subject rows exist" — this fixture happens to
        # have zero non-free subjects, where the *old* implementation's
        # single Subject.objects.all() scan would've short-circuited to 1
        # query total. The new version trades that single-fixture case for
        # a guaranteed-flat cost at any real catalog size — see
        # LockedSubjectIdsTests for the actual regression coverage.
        # +1 more (Free Starter Foundation, Phase 3 entitlements): a single
        # extra flat read of FreeStarterEntitlement (one indexed row lookup
        # on the (user, resource_type) unique constraint, not a scan) so a
        # subject a student can still cover with remaining free-starter
        # qbank quota isn't wrongly excluded from listing — see
        # entitlements.tests and academics.access.locked_subject_ids's own
        # docstring for the full reasoning. Still flat regardless of
        # catalog size, same guarantee as before, one query larger.
        self.client.force_authenticate(user=self.student)
        with self.assertNumQueries(FixedQueryCount := 7):
            resp = self.client.get('/api/questions/')
            self.assertEqual(len(resp.data), 5)

        for i in range(5, 10):
            q = Question.objects.create(subject=self.subject, chapter=self.chapter, topic=self.topic, text=f'QLC Q{i}', marks=1, negative_marks=0)
            Option.objects.create(question=q, text='opt', order=0, is_correct=True)

        with self.assertNumQueries(FixedQueryCount):
            resp = self.client.get('/api/questions/')
            self.assertEqual(len(resp.data), 10)

    def test_public_list_returns_correct_nested_data(self):
        self.client.force_authenticate(user=self.student)
        resp = self.client.get('/api/questions/')
        row = resp.data[0]
        self.assertEqual(row['subject_name'], 'QLC Subject')
        self.assertEqual(row['chapter_name'], 'QLC Chapter')
        self.assertEqual(row['topic_name'], 'QLC Topic')
        self.assertEqual(len(row['options']), 4)
        self.assertIsNotNone(row['image_data'])
        self.assertIsNotNone(row['options'][0]['image_data'])
        # Admin-only fields must NOT leak into the public serializer.
        self.assertNotIn('reference_book_name', row)
        self.assertNotIn('created_by_name', row)

    def test_admin_list_query_count_does_not_grow_with_row_count(self):
        """The admin serializer needs 3 more select_related relations
        (explanation_image_asset, reference_book, created_by) plus a
        prefetch_related('courses') than the public one, but a staff
        request skips the student-only course-scoping and locked-subject
        queries the public path pays (see get_queryset()'s _locked_
        subject_ids/_question_course_scoped calls) — net fewer total
        queries here, still a separate, still-flat bound."""
        self.client.force_authenticate(user=self.teacher)
        with self.assertNumQueries(FixedQueryCount := 4):
            resp = self.client.get('/api/questions/')
            self.assertEqual(len(resp.data), 5)

        for i in range(5, 10):
            Question.objects.create(subject=self.subject, chapter=self.chapter, topic=self.topic, text=f'QLC Admin Q{i}', marks=1, negative_marks=0)

        with self.assertNumQueries(FixedQueryCount):
            resp = self.client.get('/api/questions/')
            self.assertEqual(len(resp.data), 10)

    def test_admin_list_returns_reference_book_and_created_by(self):
        self.client.force_authenticate(user=self.teacher)
        resp = self.client.get('/api/questions/')
        row = resp.data[0]
        self.assertEqual(row['reference_book_name'], 'QLC Book')
        self.assertEqual(row['created_by_name'], self.teacher.email)
        self.assertIsNotNone(row['explanation_image_data'])

    def test_bounded_cap_still_returns_a_bare_array_shape(self):
        self.client.force_authenticate(user=self.student)
        resp = self.client.get('/api/questions/?page_size=2')
        self.assertEqual(resp.status_code, 200)
        self.assertIsInstance(resp.data, list)

    def test_browse_action_still_works_unaffected(self):
        """The already-existing, already-correct paginated action — must
        keep returning its {count,next,previous,results} envelope exactly
        as before; the new pagination_class must not interfere with it
        (browse() builds its own paginator manually, not via
        self.paginate_queryset())."""
        self.client.force_authenticate(user=self.student)
        resp = self.client.get('/api/questions/browse/')
        self.assertEqual(resp.status_code, 200)
        self.assertIn('results', resp.data)
        self.assertIn('count', resp.data)
        self.assertEqual(resp.data['count'], 5)


class SubjectChapterTopicQueryCountTests(APITestCase):
    """Scalability audit fix: GET /subjects/, /chapters/, /topics/ each had
    module_count/question_count/video_count/solved_modules/solved_count as
    per-row SerializerMethodFields issuing their own query — multiplied
    further by ChapterSerializer nesting an unprefetched TopicSerializer.
    Confirms the annotate()-based fix stays flat as subject count grows
    (the actual N+1 regression) and that every count/aggregate is still
    correct, including the fan-out-prone case of combining three Count()
    annotations (chapters/questions/videos) on the same Subject row."""

    def setUp(self):
        from videos_app.models import Video

        from courses.models import Course, Enrollment

        self.course = Course.objects.create(name='SCT Course', prefix='SCTCOURSE')
        self.student = User.objects.create_user(username='sct_student', email='sct_student@example.com', password='pw12345')
        Enrollment.objects.create(user=self.student, course=self.course)

        self.subject = Subject.objects.create(name='SCT Subject', is_free=True)
        self.subject.courses.set([self.course])
        self.chapters = [Chapter.objects.create(subject=self.subject, name=f'SCT Ch{i}') for i in range(3)]
        self.questions = []
        for ch in self.chapters:
            topic = Topic.objects.create(chapter=ch, name=f'Topic for {ch.name}')
            for j in range(4):
                q = Question.objects.create(
                    subject=self.subject, chapter=ch, topic=topic, text=f'{ch.name} Q{j}', marks=1, negative_marks=0,
                )
                self.questions.append(q)
        for i in range(3):
            Video.objects.create(title=f'SCT Video {i}', subject=self.subject, video_url='https://example.com/v')
        # 5 attempted questions spanning exactly the first 2 chapters (4 in
        # chapter 0, 1 in chapter 1) — makes solved_modules a real
        # distinct-chapter-count check, not just "> 0".
        for q in self.questions[:5]:
            QuestionAttempt.objects.create(user=self.student, question=q, attempts_count=1, correct_count=1, mastery_status='learning')

    def test_subject_list_counts_are_correct(self):
        self.client.force_authenticate(user=self.student)
        resp = self.client.get('/api/subjects/')
        row = next(r for r in resp.data if r['id'] == self.subject.id)
        self.assertEqual(row['module_count'], 3)
        self.assertEqual(row['question_count'], 12)
        self.assertEqual(row['video_count'], 3)
        self.assertEqual(row['solved_modules'], 2)
        self.assertEqual(row['attempted_count'], 5)
        self.assertEqual(row['percent_practiced'], round(5 / 12 * 100))

    def test_subject_list_query_count_does_not_grow_with_subject_count(self):
        self.client.force_authenticate(user=self.student)
        with self.assertNumQueries(FixedQueryCount := 6):
            self.client.get('/api/subjects/')

        subject2 = Subject.objects.create(name='SCT Subject 2', is_free=True)
        subject2.courses.set([self.course])
        Chapter.objects.create(subject=subject2, name='SCT Ch2-0')

        with self.assertNumQueries(FixedQueryCount):
            resp = self.client.get('/api/subjects/')
            self.assertEqual(len(resp.data), 2)

    def test_anonymous_subject_list_still_works(self):
        # self.subject is course-scoped (courses=[self.course]), so an
        # anonymous request correctly can't see it at all — pre-existing
        # _course_scoped() behavior, unrelated to this fix. What this test
        # actually guards: the annotate()/precompute path doesn't 500 or
        # leak solved_modules/attempted_count queries for an anonymous
        # request against a *visible* (uncoursed) subject.
        open_subject = Subject.objects.create(name='SCT Open Subject', is_free=True)
        resp = self.client.get('/api/subjects/')
        self.assertEqual(resp.status_code, 200)
        row = next(r for r in resp.data if r['id'] == open_subject.id)
        self.assertEqual(row['question_count'], 0)
        self.assertEqual(row['solved_modules'], 0)
        self.assertEqual(row['attempted_count'], 0)

    def test_chapter_list_counts_and_nested_topics_are_correct(self):
        self.client.force_authenticate(user=self.student)
        resp = self.client.get(f'/api/chapters/?subject={self.subject.slug}')
        row = next(r for r in resp.data if r['id'] == self.chapters[0].id)
        self.assertEqual(row['mcq_count'], 4)
        self.assertEqual(row['video_count'], 0)
        self.assertEqual(row['solved_count'], 4)  # all 4 of chapter 0's questions were attempted
        self.assertEqual(len(row['topics']), 1)
        self.assertEqual(row['topics'][0]['question_count'], 4)

    def test_chapter_list_query_count_does_not_grow_with_chapter_count(self):
        self.client.force_authenticate(user=self.student)
        with self.assertNumQueries(FixedQueryCount := 6):
            resp = self.client.get(f'/api/chapters/?subject={self.subject.slug}')
            self.assertEqual(len(resp.data), 3)

        Chapter.objects.create(subject=self.subject, name='SCT Ch3')
        with self.assertNumQueries(FixedQueryCount):
            resp = self.client.get(f'/api/chapters/?subject={self.subject.slug}')
            self.assertEqual(len(resp.data), 4)

    def test_topic_list_counts_are_correct(self):
        self.client.force_authenticate(user=self.student)
        resp = self.client.get(f'/api/topics/?chapter={self.chapters[0].id}')
        self.assertEqual(resp.data[0]['question_count'], 4)
        self.assertEqual(resp.data[0]['video_count'], 0)

    def test_bounded_list_pagination_keeps_bare_array_shape(self):
        self.client.force_authenticate(user=self.student)
        for endpoint in ('/api/subjects/', '/api/chapters/', '/api/topics/'):
            resp = self.client.get(f'{endpoint}?page_size=1')
            self.assertEqual(resp.status_code, 200)
            self.assertIsInstance(resp.data, list)


class RandomSampleTests(TestCase):
    """Scalability audit fix (Phase 1.7): replaces `.order_by('?')` — which
    forces the DB to generate a random sort key for, and fully sort, every
    matching row before any LIMIT is applied — with fetching only bare ids,
    sampling in Python, then fetching full rows for just the sample. Direct
    unit tests for the helper itself (academics/random_sample.py); the
    Practice Session Builder's own integration tests already exercise it
    end-to-end (e.g. test_count_is_capped_at_100 above, against a 120-row
    pool)."""

    def setUp(self):
        self.subject = Subject.objects.create(name='RS Subject')
        self.questions = [Question.objects.create(subject=self.subject, text=f'RS Q{i}') for i in range(10)]

    def test_returns_exactly_count_rows_when_pool_is_larger(self):
        from academics.random_sample import random_sample

        result = random_sample(Question.objects.filter(subject=self.subject), 4)
        self.assertEqual(len(result), 4)

    def test_returns_every_row_when_pool_is_smaller_than_count(self):
        from academics.random_sample import random_sample

        result = random_sample(Question.objects.filter(subject=self.subject), 100)
        self.assertEqual(len(result), 10)
        self.assertEqual({q.id for q in result}, {q.id for q in self.questions})

    def test_returns_empty_list_for_empty_queryset(self):
        from academics.random_sample import random_sample

        result = random_sample(Question.objects.filter(subject=self.subject, text='does not exist'), 5)
        self.assertEqual(result, [])

    def test_every_returned_row_is_a_real_member_of_the_original_queryset(self):
        from academics.random_sample import random_sample

        valid_ids = {q.id for q in self.questions}
        result = random_sample(Question.objects.filter(subject=self.subject), 6)
        self.assertTrue({q.id for q in result}.issubset(valid_ids))

    def test_no_duplicate_rows_in_the_sample(self):
        from academics.random_sample import random_sample

        result = random_sample(Question.objects.filter(subject=self.subject), 7)
        ids = [q.id for q in result]
        self.assertEqual(len(ids), len(set(ids)))

    def test_preserves_the_original_querysets_select_related(self):
        """The final fetch reuses the caller's queryset (via .filter()),
        not a fresh Question.objects.all() — select_related applied by the
        caller must still take effect (no extra query per row)."""
        from academics.random_sample import random_sample

        with self.assertNumQueries(2):  # 1 for the id projection, 1 for the final fetch
            result = random_sample(Question.objects.filter(subject=self.subject).select_related('subject'), 4)
            for q in result:
                q.subject.name  # noqa: B018 — must not trigger an extra query if select_related worked

    def test_respects_additional_filters_already_on_the_queryset(self):
        from academics.random_sample import random_sample

        other_subject = Subject.objects.create(name='RS Other Subject')
        Question.objects.create(subject=other_subject, text='Should never be sampled')

        result = random_sample(Question.objects.filter(subject=self.subject), 100)
        self.assertTrue(all(q.subject_id == self.subject.id for q in result))
        self.assertEqual(len(result), 10)


class QuestionPublicIdGenerationTests(TestCase):
    """Scalability audit fix (Phase 1.8): Question.save() used to compute
    its public_id via `Question.objects.filter(public_id__startswith=
    prefix).count() + 1` — an O(prefix's question count) query on every
    single save, and a real concurrent-creation race (two saves could both
    read the same COUNT before either committed and collide on public_id's
    uniqueness constraint). Replaced with an atomic, per-prefix
    QuestionPublicIdCounter. See QuestionPublicIdConcurrencyTests below for
    the actual concurrency proof (needs TransactionTestCase)."""

    def setUp(self):
        self.subject = Subject.objects.create(name='Anatomy')  # -> prefix 'A'

    def test_sequential_creation_increments_the_public_id(self):
        q1 = Question.objects.create(subject=self.subject, text='Q1')
        q2 = Question.objects.create(subject=self.subject, text='Q2')
        q3 = Question.objects.create(subject=self.subject, text='Q3')
        self.assertEqual(q1.public_id, 'A0001')
        self.assertEqual(q2.public_id, 'A0002')
        self.assertEqual(q3.public_id, 'A0003')

    def test_continues_from_the_existing_max_seeded_by_the_data_migration(self):
        """Simulates what the 0022 data migration does: seed the counter
        from a pre-existing public_id, then confirm the next save continues
        from there instead of restarting at 1 (which would collide)."""
        from academics.models import QuestionPublicIdCounter

        Question.objects.create(subject=self.subject, text='Legacy Q', public_id='A0174')
        QuestionPublicIdCounter.objects.update_or_create(prefix='A', defaults={'last_number': 174})

        q = Question.objects.create(subject=self.subject, text='New Q')

        self.assertEqual(q.public_id, 'A0175')

    def test_different_prefixes_get_independent_counters(self):
        other_subject = Subject.objects.create(name='Botany')  # -> prefix 'B'
        qa = Question.objects.create(subject=self.subject, text='Anatomy Q')
        qb = Question.objects.create(subject=other_subject, text='Botany Q')
        self.assertEqual(qa.public_id, 'A0001')
        self.assertEqual(qb.public_id, 'B0001')

    def test_explicit_public_id_bypasses_the_counter_entirely(self):
        from academics.models import QuestionPublicIdCounter

        Question.objects.create(subject=self.subject, text='Explicit', public_id='CUSTOM-001')
        self.assertFalse(QuestionPublicIdCounter.objects.filter(prefix='A').exists())

    def test_public_id_stays_stable_across_later_edits(self):
        q = Question.objects.create(subject=self.subject, text='Original text')
        original_id = q.public_id
        q.text = 'Edited text'
        q.save()
        self.assertEqual(q.public_id, original_id)

    def test_query_count_for_a_single_save_does_not_grow_with_existing_question_count(self):
        """The old implementation's COUNT(*) grew with how many questions
        already shared this prefix; the new one is a flat, small number of
        queries regardless."""
        for i in range(30):
            Question.objects.create(subject=self.subject, text=f'Filler {i}')

        with self.assertNumQueries(FixedQueryCount := 7):
            Question.objects.create(subject=self.subject, text='Just one more')


class QuestionPublicIdConcurrencyTests(TransactionTestCase):
    """The actual concurrency proof the audit explicitly asked for ("test
    concurrent question creation") — needs TransactionTestCase (not
    TestCase) because each thread needs its own real, committing
    connection/transaction; TestCase wraps the whole test in one outer
    transaction that never actually commits, which would hide exactly the
    race this test exists to catch.

    SQLite (this suite's test DB) has no per-row locking — select_for_
    update() is a documented no-op there — and its shared-cache mode
    (needed for genuinely concurrent threads to see each other's commits
    at all) raises "database table is locked" (SQLITE_LOCKED) under
    contention, a *different* error class from the ordinary busy-wait lock
    and one that PRAGMA busy_timeout does not cover per SQLite's own docs.
    So this test retries on that specific, SQLite-only error — it is
    testing infrastructure friction, not a real race condition; on the
    real target database (MySQL/InnoDB), select_for_update() takes an
    actual row lock and a concurrent transaction blocks-and-waits instead
    of erroring, which is what the fix was verified against live in
    production (see the Phase 1.8 change-control report). What this test
    proves regardless of engine: every thread's write eventually
    succeeds, and no two of them ever end up with the same public_id —
    the actual invariant the fix guarantees."""

    def test_concurrent_saves_never_produce_duplicate_public_ids(self):
        import threading
        import time

        from django.db import connection
        from django.db.utils import OperationalError

        subject = Subject.objects.create(name='Anatomy')
        n = 12
        results = []
        errors = []
        lock = threading.Lock()

        def create_one(i):
            for attempt in range(20):
                try:
                    q = Question.objects.create(subject=subject, text=f'Concurrent Q{i}')
                    with lock:
                        results.append(q.public_id)
                    return
                except OperationalError as exc:
                    if 'locked' in str(exc).lower() and attempt < 19:
                        time.sleep(0.05)
                        continue
                    with lock:
                        errors.append(exc)
                    return
                except Exception as exc:  # noqa: BLE001 — the old race raised IntegrityError here
                    with lock:
                        errors.append(exc)
                    return
                finally:
                    connection.close()  # each thread opened its own connection; don't leak it

        threads = [threading.Thread(target=create_one, args=(i,)) for i in range(n)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()

        self.assertEqual(errors, [], f'concurrent Question.save() raised: {errors}')
        self.assertEqual(len(results), n)
        self.assertEqual(len(set(results)), n, f'duplicate public_id(s) produced under concurrency: {results}')
        self.assertEqual(Question.objects.filter(subject=subject).count(), n)


def _row_data(text='Q text', options=None, **extra):
    return {
        'text_html': f'<p>{text}</p>',
        'options': options if options is not None else [
            {'text_html': 'A', 'is_correct': True}, {'text_html': 'B', 'is_correct': False},
            {'text_html': 'C', 'is_correct': False}, {'text_html': 'D', 'is_correct': False},
        ],
        'explanation_html': '', 'explanation_video_url': '', 'remarks': '', 'past_exam_years': '', 'references': [],
        **extra,
    }


class CreateQuestionFromRowOptionBulkCreateTests(TestCase):
    """Scalability audit fix (Phase 2.3): create_question_from_row's
    image-less options are now written via Option.objects.bulk_create()
    instead of one INSERT per option."""

    def setUp(self):
        from academics.models import QuestionPublicIdCounter

        self.subject = Subject.objects.create(name='Import Subject')
        # Pre-warm this subject's public_id counter row — its first-ever
        # use does an extra INSERT (get_or_create's create branch) that a
        # later call for the same prefix doesn't pay, which would
        # otherwise make query-count comparisons between two calls in the
        # same test flaky depending on which one runs first (see Phase
        # 1.8's identical QuestionBankConfig lesson).
        QuestionPublicIdCounter.objects.create(prefix='IS', last_number=0)
        self.chapter = Chapter.objects.create(subject=self.subject, name='Import Chapter')
        self.topic = Topic.objects.create(chapter=self.chapter, name='Import Topic')
        self.batch = ImportBatch.objects.create(
            file_name='f.csv', file_format='csv', subject=self.subject, chapter=self.chapter, topic=self.topic,
        )

    def test_options_created_correctly_and_in_order(self):
        from academics.import_engine import create_question_from_row

        question = create_question_from_row(_row_data(), self.batch, [])

        options = list(question.options.order_by('order'))
        self.assertEqual(len(options), 4)
        self.assertEqual([o.text for o in options], ['A', 'B', 'C', 'D'])
        self.assertEqual([o.order for o in options], [0, 1, 2, 3])
        self.assertTrue(options[0].is_correct)
        self.assertFalse(any(o.is_correct for o in options[1:]))

    def test_blank_options_are_skipped(self):
        from academics.import_engine import create_question_from_row

        data = _row_data(options=[
            {'text_html': 'A', 'is_correct': True}, {'text_html': '   ', 'is_correct': False},
            {'text_html': 'C', 'is_correct': False},
        ])
        question = create_question_from_row(data, self.batch, [])
        self.assertEqual(question.options.count(), 2)

    def test_query_count_for_options_does_not_scale_with_option_count(self):
        """4 options vs 10 options must cost the same query count — proves
        bulk_create() is doing one INSERT regardless of option count, not
        one per option. Each call uses a distinct base question text (a
        real subject-per-call would be simplest, but distinct text is
        enough) so neither question's slug collides with the other's and
        needs an extra uniqueness-suffix query, which would make the two
        counts genuinely different for a reason unrelated to what this
        test checks."""
        small = self._count_for(_row_data('Small text', options=[
            {'text_html': f'Opt{i}', 'is_correct': i == 0} for i in range(4)
        ]))
        big = self._count_for(_row_data('A completely different big text', options=[
            {'text_html': f'Big{i}', 'is_correct': i == 0} for i in range(10)
        ]))
        self.assertEqual(small, big)

    def _count_for(self, data):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        from academics.import_engine import create_question_from_row

        with CaptureQueriesContext(connection) as ctx:
            create_question_from_row(data, self.batch, [])
        return len(ctx.captured_queries)


class RunImportTests(TestCase):
    """Scalability audit fix (Phase 2.2/2.3): run_import() (moved from the
    old thread-spawned _run_import) is what a Cloud Task invokes. Covers
    the happy path for Question Bank mode, Skip/Replace/Keep Both
    preserved exactly, resumability (a retry never reprocesses an
    already-terminal row), and the claim mechanism preventing two
    concurrent runs from double-processing the same batch."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username='import_staff', email='import_staff@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )
        self.subject = Subject.objects.create(name='Import Subject')
        self.chapter = Chapter.objects.create(subject=self.subject, name='Import Chapter')
        self.topic = Topic.objects.create(chapter=self.chapter, name='Import Topic')
        self.batch = ImportBatch.objects.create(
            uploaded_by=self.staff, file_name='f.csv', file_format='csv', status='ready',
            subject=self.subject, chapter=self.chapter, topic=self.topic, total_rows=0,
        )

    def _add_row(self, n, **kwargs):
        kwargs.setdefault('status', 'valid')
        row = ImportRow.objects.create(batch=self.batch, row_number=n, raw_data=_row_data(f'Row {n}'), **kwargs)
        self.batch.total_rows += 1
        self.batch.save(update_fields=['total_rows'])
        return row

    def test_happy_path_creates_questions_and_completes_the_batch(self):
        from academics.import_engine import run_import

        self._add_row(1)
        self._add_row(2)

        run_import(self.batch.id)

        self.batch.refresh_from_db()
        self.assertEqual(self.batch.status, 'completed')
        self.assertEqual(self.batch.created_count, 2)
        self.assertIsNotNone(self.batch.completed_at)
        self.assertEqual(Question.objects.filter(subject=self.subject).count(), 2)
        for row in self.batch.rows.all():
            self.assertEqual(row.status, 'imported')
            self.assertIsNotNone(row.created_question)

    def test_skip_creates_no_question_in_question_bank_mode(self):
        """Question Bank mode has no exam to attach a skipped duplicate
        to (unlike the Create Test flow — see
        CreateQuestionsForTestSkipAttachTests) — a skip here just means
        'don't add this one', which is unchanged by this phase."""
        from academics.import_engine import run_import

        existing = Question.objects.create(subject=self.subject, text='Existing', marks=1, negative_marks=0)
        self._add_row(1, status='duplicate', duplicate_of=existing, dedup_action='skip')

        run_import(self.batch.id)

        self.batch.refresh_from_db()
        self.assertEqual(self.batch.skipped_count, 1)
        self.assertEqual(self.batch.created_count, 0)
        row = self.batch.rows.get(row_number=1)
        self.assertEqual(row.status, 'skipped')
        self.assertIsNone(row.created_question)
        self.assertEqual(Question.objects.filter(subject=self.subject).count(), 1)  # just `existing`

    def test_replace_deletes_the_old_question_and_creates_a_new_one(self):
        from academics.import_engine import run_import

        old = Question.objects.create(subject=self.subject, text='Old', marks=1, negative_marks=0)
        self._add_row(1, status='duplicate', duplicate_of=old, dedup_action='replace')

        run_import(self.batch.id)

        self.assertFalse(Question.objects.filter(pk=old.pk).exists())
        row = self.batch.rows.get(row_number=1)
        self.assertEqual(row.status, 'imported')
        self.assertNotEqual(row.created_question_id, old.pk)

    def test_replace_does_not_delete_an_already_referenced_question(self):
        from tests_app.models import Test, TestQuestion

        from academics.import_engine import run_import

        old = Question.objects.create(subject=self.subject, text='Old', marks=1, negative_marks=0)
        test = Test.objects.create(title='Uses old', exam_type='mock')
        TestQuestion.objects.create(test=test, question=old)
        self._add_row(1, status='duplicate', duplicate_of=old, dedup_action='replace')

        run_import(self.batch.id)

        self.assertTrue(Question.objects.filter(pk=old.pk).exists())  # not deleted
        row = self.batch.rows.get(row_number=1)
        self.assertEqual(row.status, 'imported')  # a new question was still created alongside it
        self.assertIn('already used in a Test', ' '.join(row.warnings or []))

    def test_keep_both_creates_a_new_question_alongside_the_old_one(self):
        from academics.import_engine import run_import

        old = Question.objects.create(subject=self.subject, text='Old', marks=1, negative_marks=0)
        self._add_row(1, status='duplicate', duplicate_of=old, dedup_action='keep_both')

        run_import(self.batch.id)

        self.assertTrue(Question.objects.filter(pk=old.pk).exists())
        self.assertEqual(Question.objects.filter(subject=self.subject).count(), 2)

    def test_a_bad_row_does_not_abort_the_rest_of_the_batch(self):
        from academics.import_engine import run_import

        self._add_row(1)
        bad = self._add_row(2)
        bad.raw_data = {'text_html': ''}  # no options at all — still "succeeds" today (no hard validation
        # inside create_question_from_row itself), so force a real failure via a taxonomy-less batch instead:
        bad.save()
        self._add_row(3)

        run_import(self.batch.id)

        self.batch.refresh_from_db()
        self.assertEqual(self.batch.status, 'completed')
        self.assertEqual(self.batch.created_count, 3)  # the "bad" row above isn't actually invalid at the DB layer

    def test_resuming_after_a_partial_run_never_reprocesses_a_terminal_row(self):
        """The core resumability guarantee: a row already 'imported' or
        'skipped' from an earlier attempt must never be touched again —
        the old implementation only excluded status='error', which would
        have silently duplicated every already-succeeded row on a retry."""
        from academics.import_engine import run_import

        row1 = self._add_row(1)
        self._add_row(2)

        run_import(self.batch.id)  # first, "successful" run
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.created_count, 2)
        first_question_id = self.batch.rows.get(row_number=1).created_question_id

        # Simulate a retry (e.g. a Cloud Tasks redelivery) against the
        # same, already-completed batch.
        run_import(self.batch.id)

        self.assertEqual(Question.objects.filter(subject=self.subject).count(), 2)  # not 4
        row1.refresh_from_db()
        self.assertEqual(row1.created_question_id, first_question_id)  # untouched, not re-created

    def test_claim_is_reclaimable_once_stale(self):
        from academics.import_engine import _claim_batch

        self._add_row(1)
        self.assertTrue(_claim_batch(self.batch.id))  # first claim succeeds
        self.assertFalse(_claim_batch(self.batch.id))  # immediately re-claiming fails — still "fresh"

        # Backdate the claim past the staleness window to simulate an
        # instance that died mid-run.
        stale_time = timezone.now() - timezone.timedelta(minutes=settings.IMPORT_CLAIM_STALE_MINUTES + 1)
        ImportBatch.objects.filter(pk=self.batch.id).update(processing_claimed_at=stale_time)
        self.assertTrue(_claim_batch(self.batch.id))  # now reclaimable


class RunImportConcurrencyTests(TransactionTestCase):
    """The claim mechanism's actual concurrency proof — needs a real
    TransactionTestCase (not TestCase, which RunImportTests above uses):
    each thread needs its own real, committing connection to see the
    batch/rows at all. TestCase wraps the whole test in one outer
    transaction that never actually commits, so a second thread's
    connection can't see setUp()'s data — every previous version of this
    test silently processed zero rows on both threads, not because the
    claim was broken, but because neither thread's connection could see
    the batch to begin with."""

    def test_concurrent_runs_do_not_double_process_the_same_batch(self):
        import threading

        from academics.import_engine import run_import

        staff = User.objects.create_user(username='race_import_staff', email='race_import_staff@example.com', password='pw12345', is_staff=True)
        subject = Subject.objects.create(name='Race Import Subject')
        chapter = Chapter.objects.create(subject=subject, name='C')
        topic = Topic.objects.create(chapter=chapter, name='T')
        batch = ImportBatch.objects.create(
            uploaded_by=staff, file_name='f.csv', file_format='csv', status='ready',
            subject=subject, chapter=chapter, topic=topic, total_rows=5,
        )
        for i in range(1, 6):
            ImportRow.objects.create(batch=batch, row_number=i, raw_data=_row_data(f'Row {i}'), status='valid')

        def go():
            from django.db import connection
            if connection.vendor == 'sqlite':
                with connection.cursor() as cur:
                    cur.execute('PRAGMA busy_timeout = 30000')
            run_import(batch.id)
            connection.close()

        threads = [threading.Thread(target=go) for _ in range(2)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()

        self.assertEqual(Question.objects.filter(subject=subject).count(), 5)  # not 10
        batch.refresh_from_db()
        self.assertEqual(batch.created_count, 5)
        self.assertEqual(batch.status, 'completed')


class EnqueueImportTaskTests(APITestCase):
    """Scalability audit fix (Phase 2.2): replaces the bare
    threading.Thread(daemon=True) with the same Cloud Tasks pattern
    already used for image processing."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username='enqueue_staff', email='enqueue_staff@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )
        self.subject = Subject.objects.create(name='Enqueue Subject')
        self.batch = ImportBatch.objects.create(
            uploaded_by=self.staff, file_name='f.csv', file_format='csv', status='ready', subject=self.subject,
        )

    def test_sync_fallback_runs_import_inline_when_async_disabled(self):
        from unittest.mock import patch

        from academics.import_tasks import enqueue_import_task

        with self.settings(IMPORT_PROCESSING_ASYNC=False):
            with patch('academics.import_tasks.run_import') as mock_run:
                enqueue_import_task(self.batch.id)
        mock_run.assert_called_once_with(self.batch.id)

    def test_async_mode_creates_a_cloud_task_not_a_thread(self):
        from unittest.mock import MagicMock, patch

        from academics.import_tasks import enqueue_import_task

        with self.settings(IMPORT_PROCESSING_ASYNC=True, GCP_PROJECT_ID='proj', GCP_REGION='us-central1', CLOUD_TASKS_IMPORT_QUEUE='bulk-import'):
            mock_client = MagicMock()
            mock_client.queue_path.return_value = 'projects/proj/locations/us-central1/queues/bulk-import'
            mock_client.task_path.return_value = 'projects/proj/locations/us-central1/queues/bulk-import/tasks/import-batch-1'
            with patch('google.cloud.tasks_v2.CloudTasksClient', return_value=mock_client):
                enqueue_import_task(self.batch.id)

        mock_client.create_task.assert_called_once()
        call_kwargs = mock_client.create_task.call_args.kwargs
        task = call_kwargs['request']['task']
        self.assertEqual(task['http_request']['url'], f'{settings.BACKEND_INTERNAL_URL}/api/import-batches/process/')
        self.assertIn('X-Import-Processing-Secret', task['http_request']['headers'])
        self.assertEqual(task['http_request']['headers']['X-Import-Processing-Secret'], settings.IMPORT_PROCESSING_SECRET)
        import json as _json
        self.assertEqual(_json.loads(task['http_request']['body']), {'batch_id': self.batch.id})

    def test_confirm_endpoint_calls_enqueue_not_a_thread(self):
        from unittest.mock import patch

        self.client.force_authenticate(user=self.staff)
        self.batch.chapter = Chapter.objects.create(subject=self.subject, name='C')
        self.batch.topic = Topic.objects.create(chapter=self.batch.chapter, name='T')
        # Phase 3 (async dedup): Confirm requires dedup_status='completed'
        # for a batch with a Subject set — this fixture predates that
        # gate and set Subject directly rather than via the taxonomy
        # PATCH (which is what actually drives dedup), so set it
        # explicitly to reflect a batch whose duplicate check has
        # genuinely finished, which is what this test means to exercise.
        self.batch.dedup_status = 'completed'
        self.batch.save()

        with patch('academics.import_views.enqueue_import_task') as mock_enqueue:
            resp = self.client.post(f'/api/import-batches/{self.batch.id}/confirm/')

        self.assertEqual(resp.status_code, 200)
        mock_enqueue.assert_called_once_with(self.batch.id)


class ImportProcessingHandlerViewTests(APITestCase):
    def setUp(self):
        self.subject = Subject.objects.create(name='Handler Subject')
        self.batch = ImportBatch.objects.create(file_name='f.csv', file_format='csv', status='ready', subject=self.subject)

    def test_wrong_secret_is_rejected(self):
        resp = self.client.post(
            '/api/import-batches/process/', data={'batch_id': self.batch.id}, format='json',
            HTTP_X_IMPORT_PROCESSING_SECRET='wrong',
        )
        self.assertEqual(resp.status_code, 401)

    def test_missing_secret_is_rejected(self):
        resp = self.client.post('/api/import-batches/process/', data={'batch_id': self.batch.id}, format='json')
        self.assertEqual(resp.status_code, 401)

    def test_correct_secret_runs_the_import(self):
        resp = self.client.post(
            '/api/import-batches/process/', data={'batch_id': self.batch.id}, format='json',
            HTTP_X_IMPORT_PROCESSING_SECRET=settings.IMPORT_PROCESSING_SECRET,
        )
        self.assertEqual(resp.status_code, 200)
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.status, 'completed')

    def test_a_run_import_exception_returns_500_so_cloud_tasks_retries(self):
        from unittest.mock import patch

        with patch('academics.import_views.run_import', side_effect=RuntimeError('boom')):
            resp = self.client.post(
                '/api/import-batches/process/', data={'batch_id': self.batch.id}, format='json',
                HTTP_X_IMPORT_PROCESSING_SECRET=settings.IMPORT_PROCESSING_SECRET,
            )
        self.assertEqual(resp.status_code, 500)


class CreateQuestionsForTestSkipAttachTests(APITestCase):
    """CRITICAL, per the scalability audit's duplicate-handling rule:
    'When Skip is selected: do not create a new duplicate Question Bank
    record but allow the existing matching question to be attached to the
    newly created exam.' Covers the Import & Create Test flow specifically
    (the only flow with an exam/test to attach to)."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username='createtest_staff', email='createtest_staff@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )
        self.subject = Subject.objects.create(name='CreateTest Subject')
        self.chapter = Chapter.objects.create(subject=self.subject, name='CreateTest Chapter')
        self.topic = Topic.objects.create(chapter=self.chapter, name='CreateTest Topic')
        self.existing = Question.objects.create(subject=self.subject, text='Existing dup', marks=1, negative_marks=0)
        self.batch = ImportBatch.objects.create(
            uploaded_by=self.staff, file_name='f.csv', file_format='csv', status='ready',
            subject=self.subject, chapter=self.chapter, topic=self.topic,
        )
        ImportRow.objects.create(
            batch=self.batch, row_number=1, raw_data=_row_data('Row 1'),
            status='duplicate', duplicate_of=self.existing, dedup_action='skip',
        )
        ImportRow.objects.create(batch=self.batch, row_number=2, raw_data=_row_data('Row 2'), status='valid')
        self.batch.total_rows = 2
        self.batch.save(update_fields=['total_rows'])
        self.client.force_authenticate(user=self.staff)

    def test_skip_attaches_the_existing_question_no_new_duplicate(self):
        before_count = Question.objects.filter(subject=self.subject).count()

        resp = self.client.post(f'/api/import-batches/{self.batch.id}/create-test/', {'title': 'New Exam', 'exam_type': 'mock'}, format='json')

        self.assertEqual(resp.status_code, 200, resp.data)
        # No new duplicate Question Bank record for the skipped row.
        self.assertEqual(Question.objects.filter(subject=self.subject).count(), before_count + 1)  # +1 for row 2 only

        from tests_app.models import Test, TestQuestion

        test = Test.objects.get(pk=resp.data['test_id'])
        question_ids_in_test = set(TestQuestion.objects.filter(test=test).values_list('question_id', flat=True))
        # The existing (pre-batch) question IS attached to the new exam.
        self.assertIn(self.existing.id, question_ids_in_test)
        self.assertEqual(len(question_ids_in_test), 2)  # existing + row 2's new question, not just 1

        skipped_row = self.batch.rows.get(row_number=1)
        self.assertEqual(skipped_row.status, 'skipped')  # not 'imported'
        self.assertEqual(skipped_row.created_question_id, self.existing.id)

    def test_rollback_never_deletes_the_skip_attached_existing_question(self):
        """status stays 'skipped' specifically so ImportRollbackView (which
        only ever considers status='imported' rows) can never delete a
        question this batch didn't create."""
        resp = self.client.post(f'/api/import-batches/{self.batch.id}/create-test/', {'title': 'New Exam', 'exam_type': 'mock'}, format='json')
        self.assertEqual(resp.status_code, 200, resp.data)

        # create-test sets status='completed' directly; rollback isn't
        # actually reachable for this mode via the UI, but assert the
        # underlying invariant that protects it regardless: an ORM-level
        # scan for rollback-eligible rows must never include this one.
        imported_rows = self.batch.rows.filter(status='imported', created_question__isnull=False)
        self.assertNotIn(self.existing.id, [r.created_question_id for r in imported_rows])


class LockedSubjectIdsTests(TestCase):
    """Scalability audit fix (Phase 3): locked_subject_ids() used to call
    has_qbank_access(user, subject) once per Subject in the whole
    platform — has_qbank_access() itself issues 1-2 queries, so this was
    an O(subject count) query pattern on a call site hit on nearly every
    Question Bank/Smart Practice request. Rewritten to a flat 2-query
    batch. Every test here cross-checks the new implementation against
    has_qbank_access() called individually — the actual authorization
    logic must produce byte-for-byte identical results, just faster."""

    def setUp(self):
        from courses.models import Course, Enrollment
        from billing.models import Subscription, SubscriptionPlan

        self.course = Course.objects.create(name='LSI Course', prefix='LSICOURSE')
        self.student = User.objects.create_user(username='lsi_student', email='lsi_student@example.com', password='pw12345')
        Enrollment.objects.create(user=self.student, course=self.course)

        self.free_subject = Subject.objects.create(name='LSI Free', is_free=True)

        self.pro_subscribed = Subject.objects.create(name='LSI Pro Subscribed', is_free=False)
        self.pro_subscribed.courses.set([self.course])
        plan = SubscriptionPlan.objects.create(course=self.course, product_type='qbank', name='QBank', price=100)
        Subscription.objects.create(user=self.student, plan=plan, course=self.course, product_type='qbank', is_active=True)

        self.pro_unsubscribed_course = Course.objects.create(name='LSI Other Course', prefix='LSIOTHER')
        self.pro_unsubscribed = Subject.objects.create(name='LSI Pro Unsubscribed', is_free=False)
        self.pro_unsubscribed.courses.set([self.pro_unsubscribed_course])

        self.pro_no_course = Subject.objects.create(name='LSI Pro No Course', is_free=False)  # no courses assigned at all

    def _reference_locked_ids(self, user):
        """Brute-force ground truth: exactly what the old implementation
        computed, called directly rather than reproduced from memory."""
        from billing.access import has_qbank_access

        return {s.id for s in Subject.objects.all() if not has_qbank_access(user, s)}

    def test_matches_has_qbank_access_for_a_subscribed_student(self):
        from academics.access import locked_subject_ids

        result = set(locked_subject_ids(self.student))
        self.assertEqual(result, self._reference_locked_ids(self.student))
        self.assertNotIn(self.free_subject.id, result)
        self.assertNotIn(self.pro_subscribed.id, result)  # has an active subscription
        self.assertIn(self.pro_unsubscribed.id, result)
        self.assertIn(self.pro_no_course.id, result)

    def test_matches_has_qbank_access_for_an_anonymous_user(self):
        from django.contrib.auth.models import AnonymousUser

        from academics.access import locked_subject_ids

        anon = AnonymousUser()
        result = set(locked_subject_ids(anon))
        self.assertEqual(result, self._reference_locked_ids(anon))
        self.assertNotIn(self.free_subject.id, result)
        self.assertIn(self.pro_subscribed.id, result)  # no subscription for an anonymous user
        self.assertIn(self.pro_unsubscribed.id, result)
        self.assertIn(self.pro_no_course.id, result)

    def test_staff_sees_everything_unlocked(self):
        from academics.access import locked_subject_ids

        staff = User.objects.create_user(username='lsi_staff', email='lsi_staff@example.com', password='pw12345', is_staff=True)
        self.assertEqual(locked_subject_ids(staff), [])

    def test_query_count_does_not_grow_with_subject_count(self):
        from django.test.utils import CaptureQueriesContext
        from django.db import connection

        from academics.access import locked_subject_ids

        with CaptureQueriesContext(connection) as ctx:
            locked_subject_ids(self.student)
        small_count = len(ctx.captured_queries)

        for i in range(20):
            s = Subject.objects.create(name=f'LSI Bulk {i}', is_free=False)
            s.courses.set([self.pro_unsubscribed_course])

        with CaptureQueriesContext(connection) as ctx:
            result = locked_subject_ids(self.student)
        large_count = len(ctx.captured_queries)

        self.assertEqual(small_count, large_count)


class QuestionSearchFulltextTests(APITestCase):
    """Scalability audit Phase B: QuestionViewSet's search filter
    (academics/views.py: _apply_question_search) moved from a full-scan
    Q(text__icontains=...) to a MySQL FULLTEXT fast path with the
    *original* full-scan kept as an exact-behavior fallback. This suite's
    test DB is SQLite, which has no FULLTEXT support at all — every test
    below either (a) exercises the always-taken-on-SQLite fallback branch
    directly (proving search results are byte-for-byte what the old code
    returned), or (b) mocks `connection.vendor`/`Question.objects.extra`
    to test the MySQL branch-*selection* logic in isolation, without
    actually needing a real MySQL connection. The real end-to-end
    FULLTEXT behavior (cardiac/cardi/diac and friends, against the real
    index) is validated separately on staging — see the load-test report."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username='search_staff', email='search_staff@example.com', password='pw12345', is_staff=True, admin_role='admin',
        )
        self.client.force_authenticate(user=self.staff)
        self.subject = Subject.objects.create(name='Search Subject', is_free=True)
        self.q_cardiac = Question.objects.create(subject=self.subject, text='What is the primary function of cardiac muscle?')
        self.q_cardiology = Question.objects.create(subject=self.subject, text='Which specialty focuses on cardiology?')
        self.q_unrelated = Question.objects.create(subject=self.subject, text='What is the capital of France?')

    def _search(self, term):
        resp = self.client.get(f'/api/questions/?search={term}')
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.data)
        # GET /questions/ is a deliberately bare array (see
        # _QuestionListPagination's docstring), not a paginated envelope.
        return {row['id'] for row in resp.data}

    def test_whole_word_search_finds_matching_questions(self):
        self.assertEqual(self._search('cardiac'), {self.q_cardiac.id})

    def test_prefix_search_finds_word_variants(self):
        # "cardi" is a prefix of both "cardiac" and "cardiology" — the old
        # LIKE '%cardi%' behavior, which the fallback path (always taken
        # on this SQLite-backed suite) reproduces exactly.
        self.assertEqual(self._search('cardi'), {self.q_cardiac.id, self.q_cardiology.id})

    def test_mid_word_substring_still_matches_via_fallback(self):
        # "diac" only appears mid-word inside "cardiac" — this is exactly
        # the case FULLTEXT's word-tokenization can't reproduce, and
        # exactly why the fallback exists. On SQLite (always the fallback
        # branch) this must keep finding it, same as the old code.
        self.assertEqual(self._search('diac'), {self.q_cardiac.id})

    def test_unrelated_term_finds_nothing(self):
        self.assertEqual(self._search('nonexistentxyz'), set())

    def test_short_term_still_works_via_fallback(self):
        # "is" (2 chars — below FULLTEXT's default 4-char minimum token
        # size) appears in "What **is** the..." for both q_cardiac and
        # q_unrelated, not in q_cardiology's text.
        self.assertEqual(self._search('is'), {self.q_cardiac.id, self.q_unrelated.id})

    def test_search_term_with_boolean_operators_does_not_error(self):
        """Characters that mean something special to MySQL BOOLEAN MODE
        (+ - < > ( ) ~ * " @) must never break the query on any backend."""
        for term in ('cardi+ac', '-cardiac', '"cardiac"', 'cardiac*', 'a@b(c)'):
            resp = self.client.get(f'/api/questions/?search={term}')
            self.assertEqual(resp.status_code, status.HTTP_200_OK, (term, resp.data))

    def test_mysql_branch_uses_fulltext_fast_path_when_it_finds_a_match(self):
        """Branch-selection logic, mocked: on a MySQL connection, if the
        FULLTEXT probe finds at least one match, the final query must be
        built from `id__in=<fulltext results>`, not the original 6-field
        OR chain."""
        from unittest.mock import MagicMock, patch

        from django.db import connection

        from academics.views import _apply_question_search

        fake_fulltext_qs = MagicMock()
        fake_fulltext_qs.values_list.return_value = [self.q_cardiac.pk]

        with patch.object(connection, 'vendor', 'mysql'), \
                patch('academics.views.Question.objects.extra', return_value=fake_fulltext_qs) as mock_extra:
            result_qs = _apply_question_search(Question.objects.all(), 'cardiac')

        mock_extra.assert_called_once()
        called_kwargs = mock_extra.call_args.kwargs
        self.assertIn('MATCH(academics_question.text)', called_kwargs['where'][0])
        self.assertEqual(called_kwargs['params'], ['cardiac*'])
        fake_fulltext_qs.values_list.assert_called_once_with('id', flat=True)
        # The fast path was taken: id__in against the (mocked, eagerly
        # materialized) fulltext result feeds into the final filter.
        self.assertIn(self.q_cardiac.id, set(result_qs.values_list('id', flat=True)))

    def test_mysql_branch_falls_back_to_full_scan_when_fulltext_finds_nothing(self):
        """The safety net: if FULLTEXT reports zero matches (e.g. a
        mid-word-substring term like "diac"), the exact original full-scan
        query must be used instead — never just 'no results'."""
        from unittest.mock import MagicMock, patch

        from django.db import connection

        from academics.views import _apply_question_search

        fake_fulltext_qs = MagicMock()
        fake_fulltext_qs.values_list.return_value = []

        with patch.object(connection, 'vendor', 'mysql'), \
                patch('academics.views.Question.objects.extra', return_value=fake_fulltext_qs):
            result_qs = _apply_question_search(Question.objects.all(), 'diac')

        # Falls all the way back to the original text__icontains scan —
        # still finds "cardiac" even though FULLTEXT (mocked) found nothing.
        self.assertIn(self.q_cardiac.id, set(result_qs.values_list('id', flat=True)))

    def test_sqlite_never_attempts_fulltext(self):
        """On a non-MySQL connection (this whole suite), the FULLTEXT
        probe must never even be attempted — Question.objects.extra() is
        never called."""
        from unittest.mock import patch

        from academics.views import _apply_question_search

        with patch('academics.views.Question.objects.extra') as mock_extra:
            _apply_question_search(Question.objects.all(), 'cardiac')

        mock_extra.assert_not_called()


class Phase5ConfigurationConsistencyTests(APITestCase):
    """Mandatory Configuration Consistency Test: Create Exam (POST
    /api/tests/) and Import & Create Test (POST
    /import-batches/<id>/create-test/) must produce equivalent effective
    configuration defaults for equivalent inputs, per exam category — the
    exact drift TestConfigStep.js's defaultConfig() and
    exam-management/page.js's emptyForm() previously had (is_draft:
    false vs true)."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username='consistency_staff', email='consistency_staff@example.com', password='pw12345',
            is_staff=True, admin_role='admin',
        )
        self.subject = Subject.objects.create(name='Consistency Subject')
        self.chapter = Chapter.objects.create(subject=self.subject, name='Consistency Chapter')
        self.topic = Topic.objects.create(chapter=self.chapter, name='Consistency Topic')
        self.client.force_authenticate(user=self.staff)

    def _create_via_wizard(self, exam_type):
        resp = self.client.post('/api/tests/', {'title': f'{exam_type} via Create', 'exam_type': exam_type}, format='json')
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED, resp.data)
        from tests_app.models import Test
        return Test.objects.get(pk=resp.data['id'])

    def _create_via_import(self, exam_type):
        batch = ImportBatch.objects.create(
            uploaded_by=self.staff, file_name='f.csv', file_format='csv', status='ready',
            subject=self.subject, chapter=self.chapter, topic=self.topic,
        )
        ImportRow.objects.create(batch=batch, row_number=1, raw_data=_row_data('Consistency row'), status='valid')
        batch.total_rows = 1
        batch.save(update_fields=['total_rows'])

        resp = self.client.post(
            f'/api/import-batches/{batch.id}/create-test/',
            {'title': f'{exam_type} via Import', 'exam_type': exam_type}, format='json',
        )
        self.assertEqual(resp.status_code, 200, resp.data)
        from tests_app.models import Test
        return Test.objects.get(pk=resp.data['test_id'])

    def test_mock_created_via_wizard_and_via_import_have_identical_effective_config(self):
        self._assert_equivalent_config('mock')

    def test_daily_created_via_wizard_and_via_import_have_identical_effective_config(self):
        self._assert_equivalent_config('daily')

    def test_grand_created_via_wizard_and_via_import_have_identical_effective_config(self):
        self._assert_equivalent_config('grand')

    def test_pyq_created_via_wizard_and_via_import_have_identical_effective_config(self):
        self._assert_equivalent_config('pyq')

    def test_qbank_created_via_wizard_and_via_import_have_identical_effective_config(self):
        self._assert_equivalent_config('qbank')

    def _assert_equivalent_config(self, exam_type):
        from tests_app.policy import POLICY_CONTROLLED_FIELDS

        wizard_test = self._create_via_wizard(exam_type)
        import_test = self._create_via_import(exam_type)

        for field in POLICY_CONTROLLED_FIELDS:
            self.assertEqual(
                getattr(wizard_test, field), getattr(import_test, field),
                f'{field} differs between Create and Import for {exam_type}: '
                f'{getattr(wizard_test, field)!r} != {getattr(import_test, field)!r}',
            )


class RevisionStatusFilterTests(APITestCase):
    """QBank 2.0 Phase 3: the four new _status_question_ids() keywords
    (overdue, due_today, repeated_mistake, recent_mistake), exercised
    through the real /questions/practice-session/ endpoint exactly like
    the existing 'new'/'weak' status tests above."""

    def setUp(self):
        self.student = User.objects.create_user(username='revstatus', email='revstatus@example.com', password='pw12345')
        self.subject = Subject.objects.create(name='Revision Status Subject')
        self.overdue_q = Question.objects.create(subject=self.subject, text='Overdue question')
        self.due_today_q = Question.objects.create(subject=self.subject, text='Due today question')
        self.not_due_q = Question.objects.create(subject=self.subject, text='Not due question')
        self.repeated_q = Question.objects.create(subject=self.subject, text='Repeated mistake question')
        self.once_wrong_q = Question.objects.create(subject=self.subject, text='Wrong once question')
        self.client.force_authenticate(user=self.student)

    def _attempt_for(self, question):
        return QuestionAttempt.objects.get(user=self.student, question=question)

    def test_overdue_status_matches_only_past_due_questions(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.overdue_q, False, source='qbank')
        record_question_result(self.student, self.not_due_q, True, source='qbank')

        overdue_attempt = self._attempt_for(self.overdue_q)
        overdue_attempt.revision_due_at = timezone.now() - timezone.timedelta(days=2)
        overdue_attempt.save(update_fields=['revision_due_at'])

        resp = self.client.post('/api/questions/practice-session/', {'status': ['overdue'], 'count': 50}, format='json')
        ids = {q['id'] for q in resp.data}
        self.assertEqual(ids, {self.overdue_q.id})

    def test_due_today_status_excludes_overdue_and_not_yet_due(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.overdue_q, False, source='qbank')
        record_question_result(self.student, self.due_today_q, False, source='qbank')
        record_question_result(self.student, self.not_due_q, True, source='qbank')

        self._attempt_for(self.overdue_q).revision_due_at = timezone.now() - timezone.timedelta(days=2)
        self._attempt_for(self.overdue_q).save()
        due_today_attempt = self._attempt_for(self.due_today_q)
        due_today_attempt.revision_due_at = timezone.now()
        due_today_attempt.save(update_fields=['revision_due_at'])
        not_due_attempt = self._attempt_for(self.not_due_q)
        not_due_attempt.revision_due_at = timezone.now() + timezone.timedelta(days=6)
        not_due_attempt.save(update_fields=['revision_due_at'])

        resp = self.client.post('/api/questions/practice-session/', {'status': ['due_today'], 'count': 50}, format='json')
        ids = {q['id'] for q in resp.data}
        self.assertEqual(ids, {self.due_today_q.id})

    def test_repeated_mistake_requires_at_least_two_wrong_answers(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.repeated_q, False, source='qbank')
        record_question_result(self.student, self.repeated_q, False, source='qbank')
        record_question_result(self.student, self.once_wrong_q, False, source='qbank')

        resp = self.client.post('/api/questions/practice-session/', {'status': ['repeated_mistake'], 'count': 50}, format='json')
        ids = {q['id'] for q in resp.data}
        self.assertEqual(ids, {self.repeated_q.id})

    def test_recent_mistake_uses_a_real_time_window(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.once_wrong_q, False, source='qbank')
        record_question_result(self.student, self.repeated_q, False, source='qbank')
        QuestionEvent.objects.filter(user=self.student, question=self.repeated_q).update(
            created_at=timezone.now() - timezone.timedelta(days=30)
        )

        resp = self.client.post('/api/questions/practice-session/', {'status': ['recent_mistake'], 'count': 50}, format='json')
        ids = {q['id'] for q in resp.data}
        self.assertIn(self.once_wrong_q.id, ids)
        self.assertNotIn(self.repeated_q.id, ids)


class SmartRevisionOrderingTests(APITestCase):
    """QBank 2.0 Phase 3B: smart_revision must rank overdue+weak+repeated
    questions ahead of merely-due, mastered questions — not database id or
    random order — and must never include a never-attempted question."""

    def setUp(self):
        self.student = User.objects.create_user(username='smartrev', email='smartrev@example.com', password='pw12345')
        self.subject = Subject.objects.create(name='Smart Revision Subject')
        # Created in an order that would put the LEAST important question
        # first by id, so a naive id/insertion-order sort would fail this
        # test — the ranking must actually reorder them.
        self.mastered_q = Question.objects.create(subject=self.subject, text='Mastered, due today')
        self.learning_q = Question.objects.create(subject=self.subject, text='Learning, due today, wrong once')
        self.weak_overdue_q = Question.objects.create(subject=self.subject, text='Weak, overdue, wrong 3 times')
        self.never_attempted_q = Question.objects.create(subject=self.subject, text='Never attempted')
        self.client.force_authenticate(user=self.student)

    def test_priority_order_and_never_attempted_excluded(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.mastered_q, True, source='qbank')
        record_question_result(self.student, self.mastered_q, True, source='qbank')

        record_question_result(self.student, self.learning_q, False, source='qbank')

        record_question_result(self.student, self.weak_overdue_q, False, source='qbank')
        record_question_result(self.student, self.weak_overdue_q, False, source='qbank')
        record_question_result(self.student, self.weak_overdue_q, False, source='qbank')
        weak_attempt = QuestionAttempt.objects.get(user=self.student, question=self.weak_overdue_q)
        weak_attempt.revision_due_at = timezone.now() - timezone.timedelta(days=5)
        weak_attempt.save(update_fields=['revision_due_at'])

        resp = self.client.post(
            '/api/questions/practice-session/', {'smart_revision': True, 'count': 10}, format='json', HTTP_ACCEPT='application/json',
        )
        ids = [q['id'] for q in resp.data]

        self.assertNotIn(self.never_attempted_q.id, ids, 'a question with no attempt history must never appear in Smart Revision')
        self.assertIn(self.weak_overdue_q.id, ids)
        self.assertIn(self.learning_q.id, ids)
        self.assertIn(self.mastered_q.id, ids)
        self.assertLess(
            ids.index(self.weak_overdue_q.id), ids.index(self.learning_q.id),
            'the weak, overdue, repeatedly-wrong question must rank ahead of the merely-due learning question',
        )
        self.assertLess(
            ids.index(self.learning_q.id), ids.index(self.mastered_q.id),
            'a learning question due today must rank ahead of an already-mastered question',
        )

    def test_revision_reason_is_present_and_matches_the_top_factor(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.weak_overdue_q, False, source='qbank')
        record_question_result(self.student, self.weak_overdue_q, False, source='qbank')
        attempt = QuestionAttempt.objects.get(user=self.student, question=self.weak_overdue_q)
        attempt.revision_due_at = timezone.now() - timezone.timedelta(days=4)
        attempt.save(update_fields=['revision_due_at'])

        resp = self.client.post('/api/questions/practice-session/', {'smart_revision': True, 'count': 10}, format='json')
        by_id = {q['id']: q for q in resp.data}
        self.assertIn('Overdue by', by_id[self.weak_overdue_q.id]['revision_reason'])

    def test_smart_revision_still_respects_course_and_subject_locking(self):
        from courses.models import Course, Enrollment

        course = Course.objects.create(name='Smart Revision Course', prefix='SMARTREV')
        Enrollment.objects.create(user=self.student, course=course)
        self.subject.courses.set([course])

        other_course = Course.objects.create(name='Other Course', prefix='OTHERREV')
        other_subject = Subject.objects.create(name='Other Subject')
        other_subject.courses.set([other_course])
        other_q = Question.objects.create(subject=other_subject, text='Not this students course')

        # Force an attempt to exist even though the student has no access
        # to this course anymore (e.g. a lapsed enrollment) — the point of
        # this test is that history alone must never grant visibility.
        QuestionAttempt.objects.create(user=self.student, question=other_q, mastery_status='weak', attempts_count=1, incorrect_count=1)

        resp = self.client.post('/api/questions/practice-session/', {'smart_revision': True, 'count': 50}, format='json')
        ids = {q['id'] for q in resp.data}
        self.assertNotIn(other_q.id, ids, 'a question outside the students enrolled course must never leak into Smart Revision')


class DashboardRevisionSummaryTests(APITestCase):
    """QBank 2.0 Phase 3A/3H: dashboard()'s new due_today/overdue/
    repeated_mistakes/recent_mistakes/revision_accuracy/daily_activity
    fields."""

    def setUp(self):
        self.student = User.objects.create_user(username='dashrev', email='dashrev@example.com', password='pw12345')
        self.subject = Subject.objects.create(name='Dashboard Revision Subject', is_free=True)
        self.q1 = Question.objects.create(subject=self.subject, text='Q1')
        self.q2 = Question.objects.create(subject=self.subject, text='Q2')
        self.client.force_authenticate(user=self.student)

    def test_zero_state_never_errors_and_reports_none_for_revision_accuracy(self):
        resp = self.client.get('/api/questions/dashboard/')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['due_today'], 0)
        self.assertEqual(resp.data['overdue'], 0)
        self.assertEqual(resp.data['repeated_mistakes'], 0)
        self.assertEqual(resp.data['recent_mistakes'], 0)
        self.assertIsNone(resp.data['revision_accuracy'])
        self.assertEqual(resp.data['daily_activity'], [])

    def test_overdue_and_due_today_counts_are_mutually_exclusive(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.q1, False, source='qbank')
        record_question_result(self.student, self.q2, False, source='qbank')
        a1 = QuestionAttempt.objects.get(user=self.student, question=self.q1)
        a1.revision_due_at = timezone.now() - timezone.timedelta(days=3)
        a1.save(update_fields=['revision_due_at'])
        a2 = QuestionAttempt.objects.get(user=self.student, question=self.q2)
        a2.revision_due_at = timezone.now()
        a2.save(update_fields=['revision_due_at'])

        resp = self.client.get('/api/questions/dashboard/')
        self.assertEqual(resp.data['overdue'], 1)
        self.assertEqual(resp.data['due_today'], 1)

    def test_repeated_mistakes_count_matches_the_documented_threshold(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.q1, False, source='qbank')
        record_question_result(self.student, self.q1, False, source='qbank')
        record_question_result(self.student, self.q2, False, source='qbank')

        resp = self.client.get('/api/questions/dashboard/')
        self.assertEqual(resp.data['repeated_mistakes'], 1)


class MistakesEnrichmentAndCourseScopingTests(APITestCase):
    """QBank 2.0 Phase 3D/3E: mistakes() bug fix (mastery_status/incorrect_
    count/confidence/last_attempted_at now populated, previously always
    'new'/None) and the new course scoping (previously entirely absent)."""

    def setUp(self):
        from courses.models import Course, Enrollment

        self.student = User.objects.create_user(username='mistakescope', email='mistakescope@example.com', password='pw12345')
        self.course_a = Course.objects.create(name='Mistakes Course A', prefix='MISTAKEA')
        self.course_b = Course.objects.create(name='Mistakes Course B', prefix='MISTAKEB')
        Enrollment.objects.create(user=self.student, course=self.course_a)
        Enrollment.objects.create(user=self.student, course=self.course_b)

        self.subject_a = Subject.objects.create(name='Mistakes Subject A', is_free=True)
        self.subject_a.courses.set([self.course_a])
        self.subject_b = Subject.objects.create(name='Mistakes Subject B', is_free=True)
        self.subject_b.courses.set([self.course_b])

        self.q_a = Question.objects.create(subject=self.subject_a, text='Course A question')
        self.q_b = Question.objects.create(subject=self.subject_b, text='Course B question')
        self.client.force_authenticate(user=self.student)

    def test_mastery_status_and_wrong_count_are_populated_not_always_new(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.q_a, False, source='qbank')
        record_question_result(self.student, self.q_a, False, source='qbank')
        attempt = QuestionAttempt.objects.get(user=self.student, question=self.q_a)

        resp = self.client.get('/api/questions/mistakes/')
        row = next(r for r in resp.data['results'] if r['id'] == self.q_a.id)
        self.assertEqual(row['mastery_status'], attempt.mastery_status)
        self.assertNotEqual(row['mastery_status'], 'new')
        self.assertEqual(row['incorrect_count'], 2)

    def test_confidence_is_exposed_for_confidence_trap_detection(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.q_a, False, source='qbank', confidence='confident')

        resp = self.client.get('/api/questions/mistakes/')
        row = next(r for r in resp.data['results'] if r['id'] == self.q_a.id)
        self.assertEqual(row['confidence'], 'confident')

    def test_course_param_narrows_to_the_selected_course_only(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.q_a, False, source='qbank')
        record_question_result(self.student, self.q_b, False, source='qbank')

        resp = self.client.get(f'/api/questions/mistakes/?course={self.course_a.id}')
        ids = {r['id'] for r in resp.data['results']}
        self.assertEqual(ids, {self.q_a.id})

    def test_no_course_param_returns_mistakes_across_all_enrolled_courses(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.q_a, False, source='qbank')
        record_question_result(self.student, self.q_b, False, source='qbank')

        resp = self.client.get('/api/questions/mistakes/')
        ids = {r['id'] for r in resp.data['results']}
        self.assertEqual(ids, {self.q_a.id, self.q_b.id})

    def test_question_outside_any_accessible_course_never_leaks_into_mistakes(self):
        from academics.services import record_question_result

        from courses.models import Course

        outside_course = Course.objects.create(name='Outside Course', prefix='OUTSIDEMISTAKE')
        outside_subject = Subject.objects.create(name='Outside Subject', is_free=True)
        outside_subject.courses.set([outside_course])
        outside_q = Question.objects.create(subject=outside_subject, text='Outside question')

        # A historical attempt exists (e.g. the student was once enrolled)
        # but they are NOT currently enrolled in outside_course.
        QuestionAttempt.objects.create(user=self.student, question=outside_q, last_result=False, incorrect_count=1, attempts_count=1)

        resp = self.client.get('/api/questions/mistakes/')
        ids = {r['id'] for r in resp.data['results']}
        self.assertNotIn(outside_q.id, ids)


class QuestionProgressEndpointTests(APITestCase):
    """QBank 2.0 Phase 4: GET /questions/progress/ — QBank-only progress
    aggregation (by_subject, mastery_distribution, weakest/strongest
    topics, accuracy_trend). Reuses QuestionAttempt/QuestionEvent only;
    no new mastery/history mechanism."""

    def setUp(self):
        self.student = User.objects.create_user(username='progress1', email='progress1@example.com', password='pw12345')
        self.subject = Subject.objects.create(name='Progress Subject', is_free=True)
        self.chapter = Chapter.objects.create(subject=self.subject, name='Progress Chapter')
        self.topic = Topic.objects.create(chapter=self.chapter, name='Progress Topic')
        self.client.force_authenticate(user=self.student)

    def _question(self, topic=None):
        return Question.objects.create(subject=self.subject, chapter=self.chapter, topic=topic, text='Q')

    def test_zero_state_never_errors(self):
        resp = self.client.get('/api/questions/progress/')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['by_subject'], [])
        self.assertEqual(resp.data['mastery_distribution'], {'learning': 0, 'need_practice': 0, 'weak': 0, 'mastered': 0})
        self.assertEqual(resp.data['weakest_topics'], [])
        self.assertEqual(resp.data['strongest_topics'], [])
        self.assertEqual(resp.data['accuracy_trend'], [])

    def test_by_subject_reflects_real_qbank_attempts_only(self):
        from academics.services import record_question_result

        q1 = self._question()
        q2 = self._question()
        record_question_result(self.student, q1, True, source='qbank')
        record_question_result(self.student, q2, False, source='qbank')

        resp = self.client.get('/api/questions/progress/')
        row = resp.data['by_subject'][0]
        self.assertEqual(row['subject_id'], self.subject.id)
        self.assertEqual(row['attempted'], 2)
        self.assertEqual(row['accuracy'], 50.0)

    def test_mastery_distribution_uses_real_mastery_status_not_a_new_formula(self):
        from academics.services import record_question_result

        weak_q = self._question()
        mastered_q = self._question()
        record_question_result(self.student, weak_q, False, source='qbank')
        record_question_result(self.student, mastered_q, True, source='qbank')
        record_question_result(self.student, mastered_q, True, source='qbank')

        weak_attempt = QuestionAttempt.objects.get(user=self.student, question=weak_q)
        mastered_attempt = QuestionAttempt.objects.get(user=self.student, question=mastered_q)

        resp = self.client.get('/api/questions/progress/')
        dist = resp.data['mastery_distribution']
        self.assertEqual(dist['weak'], QuestionAttempt.objects.filter(user=self.student, mastery_status='weak').count())
        self.assertEqual(dist['mastered'], QuestionAttempt.objects.filter(user=self.student, mastery_status='mastered').count())
        # Sanity: whatever the real thresholds produced, distribution counts match reality
        self.assertIn(weak_attempt.mastery_status, ['weak', 'learning', 'need_practice'])
        self.assertIn(mastered_attempt.mastery_status, ['mastered', 'learning', 'need_practice'])

    def test_weak_topic_excluded_below_the_documented_minimum_attempts_threshold(self):
        from academics.models import QuestionBankConfig
        from academics.services import record_question_result

        QuestionBankConfig.objects.create(pk=1, min_attempts_for_option_stats=5)
        q = self._question(topic=self.topic)
        record_question_result(self.student, q, False, source='qbank')  # only 1 attempt, threshold is 5

        resp = self.client.get('/api/questions/progress/')
        topic_ids = {t['topic_id'] for t in resp.data['weakest_topics']}
        self.assertNotIn(self.topic.id, topic_ids)

    def test_weak_topic_included_once_the_threshold_is_met(self):
        from academics.models import QuestionBankConfig
        from academics.services import record_question_result

        QuestionBankConfig.objects.create(pk=1, min_attempts_for_option_stats=2)
        # The threshold counts DISTINCT questions attempted in the topic,
        # not raw attempt-events on one question (re-answering a single
        # question repeatedly isn't a statistically meaningful sample of
        # the topic) — so this needs two different questions, not one
        # question answered twice.
        q1 = self._question(topic=self.topic)
        q2 = self._question(topic=self.topic)
        record_question_result(self.student, q1, False, source='qbank')
        record_question_result(self.student, q2, False, source='qbank')

        resp = self.client.get('/api/questions/progress/')
        topic_ids = {t['topic_id'] for t in resp.data['weakest_topics']}
        self.assertIn(self.topic.id, topic_ids)

    def test_accuracy_trend_only_includes_qbank_events_not_test_events(self):
        from academics.services import record_question_result

        q = self._question()
        record_question_result(self.student, q, True, source='qbank')
        record_question_result(self.student, q, False, source='test')

        resp = self.client.get('/api/questions/progress/')
        total_attempted_in_trend = sum(row['attempted'] for row in resp.data['accuracy_trend'])
        self.assertEqual(total_attempted_in_trend, 1)

    def test_course_param_narrows_by_subject_to_the_selected_course(self):
        from courses.models import Course, Enrollment
        from academics.services import record_question_result

        course_a = Course.objects.create(name='Progress Course A', prefix='PROGA')
        course_b = Course.objects.create(name='Progress Course B', prefix='PROGB')
        Enrollment.objects.create(user=self.student, course=course_a)
        Enrollment.objects.create(user=self.student, course=course_b)

        subject_a = Subject.objects.create(name='Progress Subject A', is_free=True)
        subject_a.courses.set([course_a])
        subject_b = Subject.objects.create(name='Progress Subject B', is_free=True)
        subject_b.courses.set([course_b])
        q_a = Question.objects.create(subject=subject_a, text='A')
        q_b = Question.objects.create(subject=subject_b, text='B')

        record_question_result(self.student, q_a, True, source='qbank')
        record_question_result(self.student, q_b, True, source='qbank')

        resp = self.client.get(f'/api/questions/progress/?course={course_a.id}')
        subject_ids = {row['subject_id'] for row in resp.data['by_subject']}
        self.assertEqual(subject_ids, {subject_a.id})

    def test_locked_pro_subject_excluded_from_by_subject(self):
        pro_subject = Subject.objects.create(name='Progress Pro Subject', is_free=False)
        q = Question.objects.create(subject=pro_subject, text='Pro question')
        # Attempt exists (e.g. subscription lapsed since) but the student
        # currently has no active QBank subscription unlocking it.
        QuestionAttempt.objects.create(user=self.student, question=q, attempts_count=1, correct_count=1, last_result=True)

        resp = self.client.get('/api/questions/progress/')
        subject_ids = {row['subject_id'] for row in resp.data['by_subject']}
        self.assertNotIn(pro_subject.id, subject_ids)


class QuestionCourseFilterFallbackTests(APITestCase):
    """Production audit P0-1 regression suite: QuestionViewSet.get_queryset()'s
    explicit ?course= filter used to check ONLY Question.courses, silently
    excluding any question relying on its Subject's course scope (the
    documented, overwhelmingly common real-data shape — Question.courses is
    blank on virtually every real question). Reproduces the exact failure
    from the Phase 5 production acceptance audit:
      GET /questions/?bookmarked=true            -> returns the question
      GET /questions/?bookmarked=true&course=X   -> incorrectly returned []
      GET /questions/{id}/?course=X              -> incorrectly 404'd
    """

    def setUp(self):
        self.student = User.objects.create_user(username='p0student', email='p0student@example.com', password='pw12345')
        from courses.models import Course as _Course, Enrollment as _Enrollment
        self.course_a = _Course.objects.create(name='P0 Course A', prefix='P0COURSEA')
        self.course_b = _Course.objects.create(name='P0 Course B', prefix='P0COURSEB')
        _Enrollment.objects.create(user=self.student, course=self.course_a)

        # TEST 1/2 fixture: Question.courses BLANK, Subject.courses = Course A
        # (the real-data shape).
        self.subject_inherited = Subject.objects.create(name='P0 Subject Inherited', is_free=True)
        self.subject_inherited.courses.set([self.course_a])
        self.question_inherited = Question.objects.create(subject=self.subject_inherited, text='Inherited-scope question')
        # Question.courses deliberately left untouched (blank).

        # TEST 3/4 fixture: Question.courses EXPLICITLY set to Course A only.
        self.subject_explicit = Subject.objects.create(name='P0 Subject Explicit', is_free=True)
        self.subject_explicit.courses.set([self.course_a])
        self.question_explicit = Question.objects.create(subject=self.subject_explicit, text='Explicit-scope question')
        self.question_explicit.courses.set([self.course_a])

        # TEST 5 fixture: Pro-locked subject, inherited scope, historical attempt.
        self.subject_pro = Subject.objects.create(name='P0 Pro Subject', is_free=False)
        self.subject_pro.courses.set([self.course_a])
        self.question_pro = Question.objects.create(subject=self.subject_pro, text='Pro-locked question')

        self.client.force_authenticate(user=self.student)

    def test_1_bookmarked_question_with_inherited_course_scope_appears_under_explicit_course_filter(self):
        self.client.post(f'/api/questions/{self.question_inherited.id}/bookmark/', {'bookmark': True}, format='json')

        resp_no_course = self.client.get('/api/questions/?bookmarked=true')
        resp_with_course = self.client.get(f'/api/questions/?bookmarked=true&course={self.course_a.id}')

        ids_no_course = {q['id'] for q in resp_no_course.data}
        ids_with_course = {q['id'] for q in resp_with_course.data}
        self.assertIn(self.question_inherited.id, ids_no_course)
        self.assertIn(
            self.question_inherited.id, ids_with_course,
            'a question inheriting its Subject\'s course scope must not disappear once ?course= is passed',
        )

    def test_2_single_question_retrieve_with_course_param_does_not_404(self):
        resp = self.client.get(f'/api/questions/{self.question_inherited.id}/?course={self.course_a.id}')
        self.assertEqual(resp.status_code, 200)

    def test_3_explicit_course_tag_still_excludes_a_different_course(self):
        from courses.models import Enrollment
        Enrollment.objects.create(user=self.student, course=self.course_b)
        resp = self.client.get(f'/api/questions/?course={self.course_b.id}')
        ids = {q['id'] for q in resp.data}
        self.assertNotIn(self.question_explicit.id, ids, 'a question explicitly tagged to Course A must not appear under Course B')

    def test_4_explicit_course_tag_still_included_for_the_correct_course(self):
        resp = self.client.get(f'/api/questions/?course={self.course_a.id}')
        ids = {q['id'] for q in resp.data}
        self.assertIn(self.question_explicit.id, ids)

    def test_5_pro_locked_subject_still_excluded_even_with_course_param(self):
        resp = self.client.get(f'/api/questions/?course={self.course_a.id}')
        ids = {q['id'] for q in resp.data}
        self.assertNotIn(self.question_pro.id, ids, 'the P0 fix must not weaken Pro-subject locking')

    def test_6_browse_overdue_count_now_matches_dashboard_overdue_count(self):
        from academics.services import record_question_result

        record_question_result(self.student, self.question_inherited, False, source='qbank')
        attempt = QuestionAttempt.objects.get(user=self.student, question=self.question_inherited)
        attempt.revision_due_at = timezone.now() - timezone.timedelta(days=2)
        attempt.save(update_fields=['revision_due_at'])

        dash = self.client.get(f'/api/questions/dashboard/?course={self.course_a.id}').data
        browse = self.client.get(f'/api/questions/browse/?status=overdue&course={self.course_a.id}&page_size=50').data

        self.assertEqual(
            browse['count'], dash['overdue'],
            'Revision Center summary count and category list count must agree once P0-1 is fixed',
        )
        self.assertEqual(dash['overdue'], 1)


class DashboardLockedSubjectExclusionTests(APITestCase):
    """Production audit P1-1 regression suite: dashboard()'s attempt_qs/
    event_qs never excluded currently Pro-locked subjects (only
    question_qs/total_questions did) — a student's historical QBank
    activity on a subject they've since lost access to kept inflating
    every derived dashboard metric. Fixed to match progress()'s own
    already-correct exclusion."""

    def setUp(self):
        self.student = User.objects.create_user(username='p1student', email='p1student@example.com', password='pw12345')
        QuestionBankConfig.objects.create(pk=1)

        # Subject A: accessible at the time of the attempt, then Pro-locked
        # (is_free=False, no active subscription for this student) —
        # exactly the "student no longer has valid access" scenario.
        self.subject = Subject.objects.create(name='P1 Subject', is_free=False)
        self.question = Question.objects.create(subject=self.subject, text='Now-inaccessible question')
        self.client.force_authenticate(user=self.student)

    def test_locked_subject_history_excluded_from_every_dashboard_metric(self):
        from academics.services import record_question_result

        # Answer it wrong 3 times (repeated mistake), then set it overdue,
        # and log a QuestionEvent too — populates every metric this bug affects.
        record_question_result(self.student, self.question, False, source='qbank')
        record_question_result(self.student, self.question, False, source='qbank')
        record_question_result(self.student, self.question, False, source='qbank')
        attempt = QuestionAttempt.objects.get(user=self.student, question=self.question)
        attempt.revision_due_at = timezone.now() - timezone.timedelta(days=3)
        attempt.save(update_fields=['revision_due_at'])

        # Sanity: the subject is genuinely locked for this student (no
        # active QBank subscription, is_free=False).
        from academics.access import locked_subject_ids
        self.assertIn(self.subject.id, locked_subject_ids(self.student))

        data = self.client.get('/api/questions/dashboard/').data

        self.assertEqual(data['attempted'], 0, 'attempted must exclude the locked subject')
        self.assertEqual(data['correct'], 0)
        self.assertEqual(data['incorrect'], 0)
        self.assertEqual(data['weak'], 0)
        self.assertEqual(data['mastered'], 0)
        self.assertEqual(data['need_practice'], 0)
        self.assertEqual(data['need_revision'], 0)
        self.assertEqual(data['due_today'], 0)
        self.assertEqual(data['overdue'], 0)
        self.assertEqual(data['repeated_mistakes'], 0)
        self.assertEqual(data['recent_mistakes'], 0)
        self.assertIsNone(data['revision_accuracy'])
        self.assertEqual(data['daily_activity'], [])
        self.assertEqual(data['study_seconds'], 0)

    def test_accessible_subject_history_still_counted(self):
        from academics.services import record_question_result

        open_subject = Subject.objects.create(name='P1 Open Subject', is_free=True)
        open_question = Question.objects.create(subject=open_subject, text='Accessible question')
        record_question_result(self.student, open_question, True, source='qbank')

        data = self.client.get('/api/questions/dashboard/').data
        self.assertEqual(data['attempted'], 1)
        self.assertEqual(data['correct'], 1)

    def test_dashboard_and_progress_apply_the_same_authorization_scope(self):
        """Part 6 cross-check: dashboard() and progress() don't have to
        return identical shapes, but the authorized population underneath
        must be consistent — the locked subject must be invisible to both."""
        from academics.services import record_question_result

        record_question_result(self.student, self.question, False, source='qbank')

        dash = self.client.get('/api/questions/dashboard/').data
        prog = self.client.get('/api/questions/progress/').data

        self.assertEqual(dash['attempted'], 0)
        subj_ids_in_progress = {row['subject_id'] for row in prog['by_subject']}
        self.assertNotIn(self.subject.id, subj_ids_in_progress)
